"""
tech_compare.py — compare OUR Splice Report workbook against the tech's own
splice report and write a third workbook that highlights every difference.

Hub-side helper, like folder_intake.py: openpyxl only, no engine imports
(each engine ships its own sor_reader copy and the hub must never import one).

What it does
------------
* Reads the ribbon x splice grid out of both workbooks.  Ours is the
  "Splice Report" sheet write_xlsx() produces (two distance rows, a header
  row, one row per ribbon, every splice spanning a km+ft column pair).  The
  tech's is whatever they hand-build: one "Distance:" row, a "Ribbon /
  ILA / Splice N" header row and one row per ribbon.  Both layouts are
  auto-detected from the "Ribbon" header cell, so the row offsets do not
  have to match.
* Lines the columns up BY DISTANCE, not by index or by name: techs number
  splices from either end and add their own "bends" / "damage" / "HH"
  columns.  Both of our distance frames (A->B and B->A) are tried and the
  one that lines up more columns wins, so a tech who counts from the far
  end still gets a like-for-like comparison.
* Parses each cell into per-fiber entries ("49,50,60 .369" -> three fibers
  at 0.369 dB; "1-8 brok" -> eight broken fibers; "all" -> the whole
  ribbon) and compares fiber by fiber.
* Writes <site_a>_to_<site_b>_SpliceReport_vs_Tech.xlsx with three sheets:
  the grid with only the differing cells filled (colour = kind of
  difference), a flat list of every fiber-level difference, and a summary
  with the column line-up.

Differences reported (per fiber, per column)
--------------------------------------------
  Tech only      the tech flagged the fiber here and our report did not
  Ours only      our report flagged the fiber here and the tech did not
  Value differs  both flagged it with a loss and the losses differ by more
                 than the tolerance (0.010 dB — techs round to 2 decimals)
  Type differs   both flagged it but one has a loss and the other a
                 word (broke / bend / DZ ...), or the words differ
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

LOSS_TOL_DB = 0.010        # |ours - tech| above this = "Value differs"
COLUMN_MATCH_KM = 0.25     # a tech column within this of ours = the same column

KIND_TECH_ONLY = 'Tech only'
KIND_OURS_ONLY = 'Ours only'
KIND_VALUE = 'Value differs'
KIND_TYPE = 'Type differs'
KIND_ORDER = (KIND_TECH_ONLY, KIND_OURS_ONLY, KIND_TYPE, KIND_VALUE)

_FILLS = {
    KIND_TECH_ONLY: 'FFC7CE',   # pink-red: they saw something we did not print
    KIND_OURS_ONLY: 'BDD7EE',   # light blue: we printed something they did not
    KIND_VALUE:     'FFEB9C',   # yellow: same fiber, different number
    KIND_TYPE:      'F8CBAD',   # orange: number on one side, a word on the other
}
_UNMATCHED_HDR = 'BFBFBF'


# ──────────────────────────────────────────────────────────────────────────
#  Grid model
# ──────────────────────────────────────────────────────────────────────────
@dataclass
class Column:
    label: str
    km: float | None            # distance in the sheet's own frame
    km_alt: float | None = None  # our B->A reading (ours only)
    excel_col: int = 0
    kind: str = 'splice'        # splice | bend | damage | ref | ila_left | ila_right | other


@dataclass
class Grid:
    path: str
    columns: list = field(default_factory=list)
    ribbons: list = field(default_factory=list)    # [(idx, label, lo, hi)]
    cells: dict = field(default_factory=dict)      # (ribbon_idx, col_idx) -> text
    ribbon_size: int = 12
    site_left: str = ''
    site_right: str = ''


_KM_RE = re.compile(r'(-?\d+(?:\.\d+)?)\s*(?:km|k\b)', re.I)


def _km_from(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(',', '')
    m = _KM_RE.search(s)
    if m:
        return float(m.group(1))
    m = re.match(r'^-?\d+(?:\.\d+)?$', s)
    return float(m.group(0)) if m else None


def _col_kind(label: str) -> str:
    l = (label or '').strip().lower()
    if 'ila' in l:
        return 'ila'
    if l.startswith('splice') or l == 'entry' or l.startswith('hh'):
        return 'splice'
    if 'bend' in l:
        return 'bend'
    if 'damag' in l or 'brok' in l:
        return 'damage'
    if 'ref' in l or 'connector' in l:
        return 'ref'
    return 'other'


def _ribbon_bounds(label: str, order: int, ribbon_size: int):
    """(idx, lo, hi) from 'Fiber 13-24 (2) (A2)' / '793-804 (67)' / 'Ribbon 3'.
    Falls back to the row order when nothing parses."""
    s = str(label)
    m_rng = re.search(r'(\d+)\s*-\s*(\d+)', s)
    m_idx = re.search(r'\((\d+)\)', s)
    if m_idx:
        idx = int(m_idx.group(1)) - 1
    elif m_rng:
        idx = (int(m_rng.group(1)) - 1) // ribbon_size
    else:
        m_n = re.search(r'(\d+)', s)
        idx = int(m_n.group(1)) - 1 if m_n else order
    if m_rng:
        lo, hi = int(m_rng.group(1)), int(m_rng.group(2))
    else:
        lo, hi = idx * ribbon_size + 1, (idx + 1) * ribbon_size
    return idx, lo, hi


def read_grid(path: str, sheet: str | None = None) -> Grid:
    """Parse a splice-report grid (ours or the tech's) into a Grid."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    if sheet and sheet in wb.sheetnames:
        ws = wb[sheet]
    elif 'Splice Report' in wb.sheetnames:
        ws = wb['Splice Report']
    else:
        ws = None
        for cand in wb.worksheets:
            if _find_header_row(cand) is not None:
                ws = cand
                break
        if ws is None:
            raise ValueError('No sheet with a "Ribbon" header row found in '
                             + os.path.basename(path))
    hdr_row = _find_header_row(ws)
    if hdr_row is None:
        raise ValueError(f'No "Ribbon" header row in {os.path.basename(path)} '
                         f'sheet {ws.title!r}')
    g = Grid(path=path)

    # Distance rows: every row above the header holding km-ish values.  Ours
    # labels them "B→A:" / "A→B:" in column 2; the tech's says "Distance:".
    dist_rows = []
    for r in range(1, hdr_row):
        lab = str(ws.cell(r, 2).value or '').strip().lower()
        kms = [_km_from(ws.cell(r, c).value) for c in range(3, ws.max_column + 1)]
        if any(k is not None for k in kms):
            dist_rows.append((r, lab))
    row_ab = next((r for r, lab in dist_rows if 'a→b' in lab or 'a->b' in lab
                   or 'a-b' in lab), None)
    row_ba = next((r for r, lab in dist_rows if 'b→a' in lab or 'b->a' in lab
                   or 'b-a' in lab), None)
    if row_ab is None:
        row_ab = dist_rows[-1][0] if dist_rows else None

    # Columns: every header-row cell with a label from column 2 on.  Our
    # merged km+ft pairs leave the ft column's header empty, so it is
    # skipped naturally.
    for c in range(2, ws.max_column + 1):
        v = ws.cell(hdr_row, c).value
        if v is None or not str(v).strip():
            continue
        label = str(v).strip()
        kind = _col_kind(label)
        km = _km_from(ws.cell(row_ab, c).value) if row_ab else None
        km_alt = _km_from(ws.cell(row_ba, c).value) if row_ba else None
        g.columns.append(Column(label=label, km=km, km_alt=km_alt,
                                excel_col=c, kind=kind))
    ilas = [i for i, col in enumerate(g.columns) if col.kind == 'ila']
    if ilas:
        g.columns[ilas[0]].kind = 'ila_left'
        g.site_left = re.sub(r'^.*?ila\s*:?\s*', '', g.columns[ilas[0]].label,
                             flags=re.I).strip()
        if len(ilas) > 1:
            g.columns[ilas[-1]].kind = 'ila_right'
            g.site_right = re.sub(r'^.*?ila\s*:?\s*', '', g.columns[ilas[-1]].label,
                                  flags=re.I).strip()

    # Ribbon size from the first fiber range we can read.
    for r in range(hdr_row + 1, ws.max_row + 1):
        m = re.search(r'(\d+)\s*-\s*(\d+)', str(ws.cell(r, 1).value or ''))
        if m and int(m.group(2)) >= int(m.group(1)):
            g.ribbon_size = int(m.group(2)) - int(m.group(1)) + 1
            break

    order = 0
    for r in range(hdr_row + 1, ws.max_row + 1):
        lab = ws.cell(r, 1).value
        if lab is None or not str(lab).strip():
            continue
        if not re.search(r'\d', str(lab)):
            continue
        idx, lo, hi = _ribbon_bounds(str(lab), order, g.ribbon_size)
        order += 1
        g.ribbons.append((idx, str(lab).strip(), lo, hi))
        for ci, col in enumerate(g.columns):
            v = ws.cell(r, col.excel_col).value
            if v is None or not str(v).strip():
                continue
            g.cells[(idx, ci)] = str(v).strip()
    return g


def _find_header_row(ws) -> int | None:
    for r in range(1, min(ws.max_row, 30) + 1):
        v = ws.cell(r, 1).value
        if v is not None and str(v).strip().lower().startswith('ribbon'):
            return r
    return None


# ──────────────────────────────────────────────────────────────────────────
#  Cell parsing  → {fiber: Entry}
# ──────────────────────────────────────────────────────────────────────────
@dataclass
class Entry:
    loss: float | None = None
    tag: str = ''            # broke | break | bend | ref | dz | launch | damage | flag
    raw: str = ''


_FIBER_SPEC = re.compile(r'^(all|\d{1,4}(?:-\d{1,4})?(?:,\d{1,4}(?:-\d{1,4})?)*)$')
_NUM = re.compile(r'^[-+~]?(?:\d+\.\d+|\.\d+|\d+\.)(?:bd)?$')
_TAG_WORDS = (
    ('broke', 'broke'), ('brok', 'broke'), ('break', 'break'),
    ('bend', 'bend'), ('damag', 'damage'), ('refl', 'ref'), ('ref', 'ref'),
    ('dz', 'dz'), ('launch', 'launch'), ('bad_launch', 'launch'),
    ('bad_tailbox', 'launch'), ('no_events', 'launch'),
    ('high_launch', 'launch'), ('duration_mismatch', 'launch'),
    ('gain', 'gainer'),
)


def _norm_text(t: str) -> str:
    t = str(t)
    t = t.replace('→', ' ').replace('|', ' ')
    t = re.sub(r'(\d)-\s+(\d)', r'\1-\2', t)         # '179- 180' -> '179-180'
    t = re.sub(r'(\d)\s+-\s+(\d)', r'\1-\2', t)       # '179 - 180' (but not '176 -0.162')
    t = re.sub(r'(?<=\d)\s*,\s*(?=\d)', ',', t)        # '49, 50' -> '49,50'
    t = re.sub(r'\s*,\s+|\s+,\s*', ' ', t)           # a stray ' , ' between entries
    t = re.sub(r'(?<=\d)\(', ' (', t)                 # '.171(B-fill)' -> '.171 (B-fill)'
    t = re.sub(r'\bF(?=\d)', '', t)                   # 'F12' -> '12'
    return t


def _expand_fibers(spec: str, lo: int, hi: int) -> list:
    if spec == 'all':
        return list(range(lo, hi + 1))
    out = []
    for part in spec.split(','):
        if '-' in part:
            a, b = part.split('-', 1)
            a, b = int(a), int(b)
            if b < a:
                a, b = b, a
            if b - a > 999:
                continue
            out.extend(range(a, b + 1))
        else:
            out.append(int(part))
    return out


def parse_cell(text: str, lo: int = 1, hi: int = 12) -> dict:
    """'49,50,60 .369'  -> {49: Entry(.369), 50: ..., 60: ...}
    '1-8 brok'         -> eight Entry(tag='broke')
    '10 BEND .883 bidi 11 bend .127 bidi' -> {10: Entry(.883,'bend'), 11: ...}
    'all 145 .35'      -> whole ribbon flagged, F145 at .35 dB
    """
    if text is None:
        return {}
    tokens = _norm_text(text).split()
    segments = []          # [(fiberspec, [tokens...])]
    cur = None
    for tok in tokens:
        if _FIBER_SPEC.match(tok.lower()):
            cur = (tok.lower(), [])
            segments.append(cur)
        elif cur is not None:
            cur[1].append(tok)
    out = {}
    for spec, rest in segments:
        e = Entry(raw=(spec + ' ' + ' '.join(rest)).strip())
        for tok in rest:
            tl = tok.lower()
            if tl.startswith('(') or tl.endswith(')'):
                continue                       # '(B-fill)', '(refl', '-67dB)' — notes
            if e.loss is None and _NUM.match(tl) and not tl.startswith('~'):
                try:
                    e.loss = float(tl.rstrip('bd').lstrip('+'))
                except ValueError:
                    pass
                continue
            if not e.tag:
                for prefix, tag in _TAG_WORDS:
                    if tl.startswith(prefix):
                        e.tag = tag
                        break
        if e.loss is None and not e.tag:
            e.tag = 'flag'
        try:
            fibers = _expand_fibers(spec, lo, hi)
        except ValueError:
            continue
        for f in fibers:
            out[f] = Entry(loss=e.loss, tag=e.tag, raw=e.raw)
    return out


# ──────────────────────────────────────────────────────────────────────────
#  Column line-up
# ──────────────────────────────────────────────────────────────────────────
def _match_columns(ours: Grid, tech: Grid, use_alt: bool):
    """Greedy nearest-distance one-to-one pairing.  Returns
    {our_col_idx: tech_col_idx}."""
    pairs = []
    for oi, oc in enumerate(ours.columns):
        okm = oc.km_alt if use_alt else oc.km
        if okm is None or oc.kind in ('ila_left', 'ila_right'):
            continue
        for ti, tc in enumerate(tech.columns):
            if tc.km is None or tc.kind in ('ila_left', 'ila_right'):
                continue
            d = abs(okm - tc.km)
            if d <= COLUMN_MATCH_KM:
                pairs.append((d, oi, ti))
    pairs.sort()
    used_o, used_t, out = set(), set(), {}
    for d, oi, ti in pairs:
        if oi in used_o or ti in used_t:
            continue
        used_o.add(oi); used_t.add(ti); out[oi] = ti
    # The ILA end columns pair by physical end.  In the alt frame the tech's
    # left end is our right end.
    o_l = next((i for i, c in enumerate(ours.columns) if c.kind == 'ila_left'), None)
    o_r = next((i for i, c in enumerate(ours.columns) if c.kind == 'ila_right'), None)
    t_l = next((i for i, c in enumerate(tech.columns) if c.kind == 'ila_left'), None)
    t_r = next((i for i, c in enumerate(tech.columns) if c.kind == 'ila_right'), None)
    if use_alt:
        t_l, t_r = t_r, t_l
    if o_l is not None and t_l is not None:
        out[o_l] = t_l
    if o_r is not None and t_r is not None:
        out[o_r] = t_r
    return out


def line_up_columns(ours: Grid, tech: Grid):
    """Pick the frame (A->B or B->A) that lines up more columns."""
    m_ab = _match_columns(ours, tech, use_alt=False)
    n_ab = sum(1 for oi in m_ab if ours.columns[oi].kind not in ('ila_left', 'ila_right'))
    if any(c.km_alt is not None for c in ours.columns):
        m_ba = _match_columns(ours, tech, use_alt=True)
        n_ba = sum(1 for oi in m_ba
                   if ours.columns[oi].kind not in ('ila_left', 'ila_right'))
        if n_ba > n_ab:
            return m_ba, 'B→A'
    return m_ab, 'A→B'


# ──────────────────────────────────────────────────────────────────────────
#  Compare
# ──────────────────────────────────────────────────────────────────────────
@dataclass
class Diff:
    ribbon_idx: int
    ribbon_label: str
    our_col: int | None
    tech_col: int | None
    fiber: int
    ours: str
    tech: str
    kind: str


def _fmt(e: Entry | None) -> str:
    if e is None:
        return ''
    if e.loss is not None and e.tag and e.tag != 'flag':
        return f'{e.tag} {e.loss:.3f}'
    if e.loss is not None:
        return f'{e.loss:.3f}'
    return e.tag


def compare(ours: Grid, tech: Grid, tol_db: float = LOSS_TOL_DB):
    """Returns (diffs, colmap, frame).  colmap = {our_col_idx: tech_col_idx}."""
    colmap, frame = line_up_columns(ours, tech)
    rev = {t: o for o, t in colmap.items()}
    our_rib = {idx: (lab, lo, hi) for idx, lab, lo, hi in ours.ribbons}
    tech_rib = {idx: (lab, lo, hi) for idx, lab, lo, hi in tech.ribbons}
    diffs = []

    def bounds(ri):
        if ri in our_rib:
            return our_rib[ri][1:]
        if ri in tech_rib:
            return tech_rib[ri][1:]
        return ri * ours.ribbon_size + 1, (ri + 1) * ours.ribbon_size

    def label(ri):
        return (our_rib.get(ri) or tech_rib.get(ri) or (f'Ribbon {ri + 1}',))[0]

    keys = set()
    for (ri, oi) in ours.cells:
        keys.add((ri, oi, colmap.get(oi)))
    for (ri, ti) in tech.cells:
        keys.add((ri, rev.get(ti), ti))
    for ri, oi, ti in sorted(keys, key=lambda k: (k[0], k[1] if k[1] is not None else 10 ** 6, k[2] or 0)):
        lo, hi = bounds(ri)
        o_ent = parse_cell(ours.cells.get((ri, oi)), lo, hi) if oi is not None else {}
        t_ent = parse_cell(tech.cells.get((ri, ti)), lo, hi) if ti is not None else {}
        for f in sorted(set(o_ent) | set(t_ent)):
            o, t = o_ent.get(f), t_ent.get(f)
            kind = None
            if o is None:
                kind = KIND_TECH_ONLY
            elif t is None:
                kind = KIND_OURS_ONLY
            elif o.loss is not None and t.loss is not None:
                if abs(o.loss - t.loss) > tol_db + 1e-9:
                    kind = KIND_VALUE
            elif o.loss is None and t.loss is None:
                if o.tag != t.tag and 'flag' not in (o.tag, t.tag):
                    kind = KIND_TYPE
            else:
                kind = KIND_TYPE
            if kind:
                diffs.append(Diff(ri, label(ri), oi, ti, f, _fmt(o), _fmt(t), kind))
    return diffs, colmap, frame


# ──────────────────────────────────────────────────────────────────────────
#  Workbook
# ──────────────────────────────────────────────────────────────────────────
def _col_title(c: Column | None) -> str:
    if c is None:
        return '—'
    if c.km is None or '@' in c.label:
        return c.label
    return f'{c.label} @ {c.km:.2f} km'


def write_comparison(ours: Grid, tech: Grid, diffs, colmap, frame, out_path: str,
                     tol_db: float = LOSS_TOL_DB) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Differences'
    thin = Side(style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_font = Font(name='Calibri', bold=True, size=12, color='FFFFFF')
    hdr_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    body = Font(name='Calibri', size=11)
    wrap = Alignment(wrap_text=True, vertical='top')

    # Column order: ours left→right (with its tech partner), then any tech
    # columns that never lined up, in the tech's distance order.
    rev = {t: o for o, t in colmap.items()}
    order = [(oi, colmap.get(oi)) for oi in range(len(ours.columns))]
    order += [(None, ti) for ti in range(len(tech.columns)) if ti not in rev]

    ws.cell(1, 1, 'Ribbon').font = hdr_font
    ws.cell(1, 1).fill = hdr_fill
    ws.cell(2, 1, '').fill = hdr_fill
    for j, (oi, ti) in enumerate(order):
        c = j + 2
        oc = ours.columns[oi] if oi is not None else None
        tc = tech.columns[ti] if ti is not None else None
        top = ws.cell(1, c, 'Ours: ' + _col_title(oc))
        bot = ws.cell(2, c, 'Tech: ' + _col_title(tc))
        for cell in (top, bot):
            cell.font = hdr_font
            cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
            cell.fill = hdr_fill
        if oc is None or tc is None:
            fill = PatternFill(start_color=_UNMATCHED_HDR, end_color=_UNMATCHED_HDR,
                               fill_type='solid')
            top.fill = bot.fill = fill
            top.font = bot.font = Font(name='Calibri', bold=True, size=12, color='000000')
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 26
    ws.column_dimensions['A'].width = 24
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 32

    by_cell = {}
    for d in diffs:
        by_cell.setdefault((d.ribbon_idx, d.our_col, d.tech_col), []).append(d)
    ribbon_ids = sorted({idx for idx, *_ in ours.ribbons} | {idx for idx, *_ in tech.ribbons})
    our_lab = {idx: lab for idx, lab, *_ in ours.ribbons}
    tech_lab = {idx: lab for idx, lab, *_ in tech.ribbons}
    for i, ri in enumerate(ribbon_ids):
        r = i + 3
        ws.cell(r, 1, our_lab.get(ri) or tech_lab.get(ri)).font = body
        ws.cell(r, 1).border = border
        for j, (oi, ti) in enumerate(order):
            c = j + 2
            cell = ws.cell(r, c)
            cell.border = border
            cell.alignment = wrap
            cell.font = body
            ds = by_cell.get((ri, oi, ti))
            if not ds:
                continue
            o_txt = ours.cells.get((ri, oi), '') if oi is not None else ''
            t_txt = tech.cells.get((ri, ti), '') if ti is not None else ''
            kinds = {d.kind for d in ds}
            worst = next(k for k in KIND_ORDER if k in kinds)
            fibers = ', '.join(f'F{d.fiber}' for d in ds[:12]) + (' …' if len(ds) > 12 else '')
            cell.value = (f'Ours: {o_txt or "(blank)"}\nTech: {t_txt or "(blank)"}\n'
                          f'{" / ".join(k for k in KIND_ORDER if k in kinds)}: {fibers}')
            cell.fill = PatternFill(start_color=_FILLS[worst], end_color=_FILLS[worst],
                                    fill_type='solid')
    ws.freeze_panes = 'B3'

    # ── Difference list ──
    wl = wb.create_sheet('Difference list')
    heads = ['Ribbon', 'Fiber', 'Our column', 'Tech column', 'Ours', 'Tech', 'Difference']
    for c, h in enumerate(heads, 1):
        cell = wl.cell(1, c, h)
        cell.font = hdr_font; cell.fill = hdr_fill
    for r, d in enumerate(diffs, 2):
        oc = ours.columns[d.our_col] if d.our_col is not None else None
        tc = tech.columns[d.tech_col] if d.tech_col is not None else None
        vals = [d.ribbon_label, d.fiber, _col_title(oc), _col_title(tc),
                d.ours or '(blank)', d.tech or '(blank)', d.kind]
        for c, v in enumerate(vals, 1):
            cell = wl.cell(r, c, v)
            cell.font = body
            cell.fill = PatternFill(start_color=_FILLS[d.kind], end_color=_FILLS[d.kind],
                                    fill_type='solid')
    for col, w in zip('ABCDEFG', (24, 8, 26, 26, 22, 22, 16)):
        wl.column_dimensions[col].width = w
    wl.freeze_panes = 'A2'
    if diffs:
        wl.auto_filter.ref = f'A1:G{len(diffs) + 1}'

    # ── Summary ──
    wsum = wb.create_sheet('Summary', 0)
    wsum.column_dimensions['A'].width = 34
    wsum.column_dimensions['B'].width = 60
    wsum.column_dimensions['C'].width = 34
    rows = [
        ('Our report', os.path.basename(ours.path)),
        ('Tech report', os.path.basename(tech.path)),
        ('Distance frame used', f'{frame} (the frame that lined up more columns)'),
        ('Loss tolerance', f'{tol_db:.3f} dB'),
        ('Ribbons (ours / tech)', f'{len(ours.ribbons)} / {len(tech.ribbons)}'),
        ('Columns lined up', f'{sum(1 for oi in colmap if ours.columns[oi].kind not in ("ila_left", "ila_right"))} of '
                             f'{sum(1 for c in ours.columns if c.kind not in ("ila_left", "ila_right"))} ours, '
                             f'{sum(1 for c in tech.columns if c.kind not in ("ila_left", "ila_right"))} tech'),
        ('Fiber-level differences', len(diffs)),
    ]
    for k in KIND_ORDER:
        rows.append((f'   {k}', sum(1 for d in diffs if d.kind == k)))
    r = 1
    wsum.cell(r, 1, 'Splice report vs tech report').font = Font(bold=True, size=14)
    r += 2
    for k, v in rows:
        wsum.cell(r, 1, k).font = Font(bold=True)
        wsum.cell(r, 2, v)
        r += 1
    r += 1
    wsum.cell(r, 1, 'Colour key').font = Font(bold=True)
    r += 1
    for k in KIND_ORDER:
        c = wsum.cell(r, 1, k)
        c.fill = PatternFill(start_color=_FILLS[k], end_color=_FILLS[k], fill_type='solid')
        wsum.cell(r, 2, {
            KIND_TECH_ONLY: 'the tech flagged this fiber here; our report did not',
            KIND_OURS_ONLY: 'our report flagged this fiber here; the tech did not',
            KIND_VALUE: f'both flagged it; losses differ by more than {tol_db:.3f} dB',
            KIND_TYPE: 'both flagged it; a loss on one side and a word (broke, bend, DZ …) on the other',
        }[k])
        r += 1
    c = wsum.cell(r, 1, 'Grey column header')
    c.fill = PatternFill(start_color=_UNMATCHED_HDR, end_color=_UNMATCHED_HDR, fill_type='solid')
    wsum.cell(r, 2, f'a column only one report has (no column within {COLUMN_MATCH_KM * 1000:.0f} m in the other)')
    r += 2
    wsum.cell(r, 1, 'Column line-up').font = Font(bold=True)
    r += 1
    for c, h in enumerate(('Our column', 'Tech column', 'Distance gap'), 1):
        cell = wsum.cell(r, c, h)
        cell.font = hdr_font; cell.fill = hdr_fill
    r += 1
    for oi, ti in order:
        oc = ours.columns[oi] if oi is not None else None
        tc = tech.columns[ti] if ti is not None else None
        wsum.cell(r, 1, _col_title(oc))
        wsum.cell(r, 2, _col_title(tc))
        if oc is not None and tc is not None and tc.km is not None:
            okm = oc.km_alt if frame == 'B→A' else oc.km
            if okm is not None:
                wsum.cell(r, 3, f'{abs(okm - tc.km) * 1000:.0f} m')
        elif oc is None or tc is None:
            for c in (1, 2):
                wsum.cell(r, c).fill = PatternFill(start_color=_UNMATCHED_HDR,
                                                   end_color=_UNMATCHED_HDR, fill_type='solid')
        r += 1

    wb.save(out_path)
    return out_path


def compare_reports(ours_xlsx: str, tech_xlsx: str, out_path: str,
                    tol_db: float = LOSS_TOL_DB) -> dict:
    """One call for the hub: read both, compare, write, summarise."""
    ours = read_grid(ours_xlsx, 'Splice Report')
    tech = read_grid(tech_xlsx)
    diffs, colmap, frame = compare(ours, tech, tol_db)
    write_comparison(ours, tech, diffs, colmap, frame, out_path, tol_db)
    n_our_cols = sum(1 for c in ours.columns if c.kind not in ('ila_left', 'ila_right'))
    n_tech_cols = sum(1 for c in tech.columns if c.kind not in ('ila_left', 'ila_right'))
    n_matched = sum(1 for oi in colmap if ours.columns[oi].kind not in ('ila_left', 'ila_right'))
    return {
        'xlsx': out_path,
        'frame': frame,
        'n_diffs': len(diffs),
        'counts': {k: sum(1 for d in diffs if d.kind == k) for k in KIND_ORDER},
        'columns_matched': n_matched,
        'columns_ours': n_our_cols,
        'columns_tech': n_tech_cols,
        'ribbons_ours': len(ours.ribbons),
        'ribbons_tech': len(tech.ribbons),
        'diffs': [{'Ribbon': d.ribbon_label, 'Fiber': d.fiber,
                   'Our column': _col_title(ours.columns[d.our_col]) if d.our_col is not None else '—',
                   'Tech column': _col_title(tech.columns[d.tech_col]) if d.tech_col is not None else '—',
                   'Ours': d.ours or '(blank)', 'Tech': d.tech or '(blank)',
                   'Difference': d.kind} for d in diffs],
    }
