"""Project status: what the FQA package still needs, ticked off from files.

Robert, 2026-09-23: "we will want a project status area, where we see what
we need. as we collect photos, GPS coordinates, traces, etc those would be
removed from what we need. the basis for what we need is the completed FQA
form."  Option 2: the field's emailed files are dropped on the page and kept
in <project folder>/Field/.

The fixtures are built the way the real files are: the FQA workbook is the
blank customer form with cells filled and photos anchored under a per-location
caption on the Pictures tab (fieldcapture/web/fqa.js), the capture sheet has
Field Capture's 'Submissions' columns (fieldcapture/web/app.js buildWorkbook).
"""
from __future__ import annotations

import io
import json
import os
import re
import shutil
from pathlib import Path

import pytest

from conftest import (REPO_ROOT, FIXTURE_SPLICE_A_DIR, FIXTURE_SPLICE_B_DIR,
                      run_streamlit)

TEMPLATE = REPO_ROOT / "fqa" / "templates" / "FQA_Site_Survey_v1_1.xlsm"


@pytest.fixture
def hub():
    import app
    return app


@pytest.fixture
def settings_dir(tmp_path, monkeypatch):
    d = tmp_path / "settings"
    d.mkdir()
    monkeypatch.setenv("OTDR_SETTINGS_DIR", str(d))
    return d


@pytest.fixture
def span_dir(tmp_path):
    root = tmp_path / "SITEA-SITEB"
    shutil.copytree(FIXTURE_SPLICE_A_DIR, root / "A")
    shutil.copytree(FIXTURE_SPLICE_B_DIR, root / "B")
    return root


def _png():
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (40, 30), (200, 80, 40)).save(buf, "PNG")
    buf.seek(0)
    return buf


def field_fqa(path, cells=None, photos=None):
    """The blank form with Site Survey cells filled and Pictures bands:
    photos = {'A-Location: SITEA, photos by RC, 9/23/2026': 2, ...}."""
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    wb = openpyxl.load_workbook(TEMPLATE, keep_vba=True)
    ss = wb["Site Survey Data"]
    for ref, v in (cells or {}).items():
        sheet, _, ref = ref.rpartition("!")
        (wb[sheet] if sheet else ss)[ref] = v
    pics = wb["Pictures"]
    row = 1
    for caption, n in (photos or {}).items():
        pics[f"A{row}"] = caption
        for k in range(n):
            img = XLImage(_png())
            pics.add_image(img, f"B{row + 1 + 3 * k}")
        row += 2 + 3 * n
    wb.save(path)
    return path


def capture_sheet(path, rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Submissions"
    head = ["#", "Location", "Date", "Time", "Tester Initials", "Room", "Aisle",
            "Bay", "GPS Latitude", "GPS Longitude", "GPS Accuracy (m)",
            "GPS (lat, lon)", "GPS Source", "Photos", "Sheet"]
    ws.append(head)
    for i, r in enumerate(rows, 1):
        ws.append([i, r["loc"], "9/23/2026", "10:00", "RC", None, None, None,
                   r.get("lat"), r.get("lon"), r.get("acc"), None, "Phone GPS",
                   r.get("photos", 0), f"Site {i}"])
    wb.save(path)
    return path


def _by_item(items):
    return {i["item"]: i for i in items}


def _snap(span_dir):
    return {"spans": [{"mode": "two", "dir_a": str(span_dir / "A"),
                       "dir_b": str(span_dir / "B"), "folder": "",
                       "site_a": "SITEA", "site_b": "SITEB"}]}


def _sec(items, n):
    return {i["item"]: i for i in items if i["section"] == n}


def production_sheet(path):
    """The FQA Builder suite's Span-4-shaped production sheet."""
    import openpyxl
    import test_fqa_builder as tf
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    tf._write_termination_sheet(wb.create_sheet('Flagler ILA'), 'Flagler ILA',
                                '39.475609, -103.023762', 'RR 100.07', [432, 432, 288])
    for name, east, west in tf.SPAN4_MARKS:
        tf._write_splice_sheet(wb.create_sheet(name), name, f'GPS for {name}', east, west)
    tf._write_termination_sheet(wb.create_sheet('Bethune  ILA'), 'Bethune ILA',
                                '39.366775, -102.428434', 'RR 100.08', [432, 432, 288])
    wb.save(path)
    return str(path)


# ── the list mirrors the customer's form ────────────────────────────────
def test_checklist_cells_match_the_forms_own_submittal_checklist(hub):
    """Every "Entered?" row of the Submittal Checklist tests one Site Survey
    cell.  If the customer moves a cell in a form revision, this fails."""
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE, read_only=True)
    ws = wb["Submittal Checklist"]
    form = []
    for r in range(6, 21):
        f = ws.cell(r, 2).value
        m = re.search(r"'Site Survey Data'!([A-Z]+\d+)", str(f))
        form.append((f"{ws.cell(r, 3).value:.2f}", m.group(1)))
    wb.close()
    ours = [(no, ref) for no, _what, _sheet, ref in hub.FQA_CHECKLIST_CELLS]
    assert ours == form


def test_four_sections_and_a_blank_form_needs_everything(hub, span_dir, tmp_path):
    field = tmp_path / "Field"
    field.mkdir()
    field_fqa(field / "blank.xlsm")
    fqa, caps = hub.collect_field_files(str(field))
    items = hub.project_status(_snap(span_dir), fqa, caps)
    assert {i["section"] for i in items} == {1, 2, 3, 4}
    assert not _sec(items, 1)["1.01  Bldg 1 (A end) site address"]["ok"]
    assert not _sec(items, 2)["2.01  FAT table completed"]["ok"]
    assert not _sec(items, 3)["3.02-3.06  Events"]["ok"]
    s4 = _sec(items, 4)
    assert s4["4.01  A-direction traces"]["ok"] and s4["4.01  Every fiber shot both ways"]["ok"]
    assert not s4["4.02  Power meter files"]["ok"]
    assert not s4["4.03  Files named to the naming convention"]["ok"]


def test_field_files_tick_off_section_1(hub, span_dir, tmp_path):
    field = tmp_path / "Field"
    field.mkdir()
    field_fqa(field / "FQA field 1.2.xlsm",
              cells={"F51": "1", "H51": "COLO", "J51": "12", "L51": "4",
                     "F52": "RMU 30", "F56": "LC", "M56": "OSX", "J53": 144,
                     "E9": "123 Main St, Sitea", "E10": "SITEA CO"},
              photos={"A-Location: SITEA CO, photos by RC, 9/23/2026": 3,
                      "Z-Location, photos by RC, 9/23/2026": 1})
    capture_sheet(field / "OTDR_Field_RC.xlsx",
                  [{"loc": "Z-Location", "lat": 38.5, "lon": -98.1, "photos": 2}])
    fqa, caps = hub.collect_field_files(str(field))
    assert hub.fqa_photos_per_end(fqa[0][1]["pictures"]) == {"A": 3, "Z": 1, "unassigned": 0}
    s1 = _sec(hub.project_status(_snap(span_dir), fqa, caps), 1)
    assert s1["A end rack location (floor, room, aisle, bay)"]["ok"]
    assert s1["A end panel details (RMU, connector, panel type)"]["ok"]
    assert s1["A end photos"]["detail"] == "3 photos"
    assert s1["Z end photos"]["detail"] == "2 photos"      # the capture sheet's count
    assert not s1["Z end rack location (floor, room, aisle, bay)"]["ok"]
    assert s1["1.01  Bldg 1 (A end) site address"]["ok"]
    assert not s1["1.02  Bldg 2 (Z end) site address"]["ok"]


def test_form_prompts_do_not_count_as_entered(hub, tmp_path, span_dir):
    field = tmp_path / "Field"
    field.mkdir()
    field_fqa(field / "f.xlsm", cells={"F85": "<Select Package Type>", "F86": "<Select>",
                                        "F97": "  "})
    fqa, caps = hub.collect_field_files(str(field))
    s1 = _sec(hub.project_status(_snap(span_dir), fqa, caps), 1)
    assert not s1["1.11  Package type"]["ok"]
    assert not s1["1.10  Test revision"]["ok"]
    assert not s1["1.09  Number of fibers tested"]["ok"]


# ── Event Log: every event, and the GPS ─────────────────────────────────
@pytest.mark.parametrize("text,kind", [
    ("39 21 21.13 N 102 22 18.70 W", "gps"),
    ("39 21 07.27 N 102 21 12.17 W", "gps"),
    ("40 44 43 N 114 04 43 W", "gps"),
    ("39.475609, -103.023762", "gps"),
    ("32353 State Hwy 40 Bethune", "address"),
    ("US HWY 84 ON I40", "address"),
    ("SE 2ND St & E Island Rd", "address"),
    ("", "blank"),
    # Real mistakes from packages marked FINAL:
    ("39 18 59 52 N 102 10 06 54 W", "bad"),      # decimal points lost
    ("39 15 50.16N 101 5407.26 W", "bad"),        # a space lost
    ("32 47 06.0 N 114 47 55.0 E", "bad"),        # E for W
    ("39.475609, 103.023762", "bad"),             # lost the minus sign
])
def test_event_locations_are_read_and_bad_gps_is_caught(hub, text, kind):
    assert hub.parse_event_location(text)[0] == kind


def test_dms_reads_to_decimal_degrees(hub):
    kind, (lat, lon) = hub.parse_event_location("39 21 21.13 N 102 22 18.70 W")
    assert kind == "gps"
    assert abs(lat - 39.355869) < 1e-5 and abs(lon + 102.371861) < 1e-5


def _events(path, rows, length=42000):
    """A form with Event Log rows [(location, type, fiber, dist), ...] and
    Site Z after them, the way Excel saves one."""
    cells = {"Event Log!W8": length}
    for i, (loc, typ, fib, dist) in enumerate(rows):
        r = 18 + i
        cells.update({f"Event Log!B{r}": i + 1, f"Event Log!D{r}": loc,
                      f"Event Log!Q{r}": typ, f"Event Log!X{r}": fib,
                      f"Event Log!AB{r}": dist})
    cells[f"Event Log!B{18 + len(rows)}"] = "Site Z"
    return field_fqa(path, cells=cells)


def test_every_event_is_checked_not_just_the_first_two(hub, tmp_path, span_dir):
    """The customer's checklist reads D18:D19 only; event 3 blank passes it."""
    field = tmp_path / "Field"
    field.mkdir()
    ok = ("39 21 21.13 N 102 22 18.70 W", "New Field Splice", "SMF 28")
    _events(field / "f.xlsm", [
        (*ok, 60), (*ok, 5010), (None, "New Field Splice", "SMF 28", 7650),
        ("39 18 59 52 N 102 10 06 54 W", "New Field Splice", "SMF 28", 7000),
    ])
    fqa, caps = hub.collect_field_files(str(field))
    s3 = _sec(hub.project_status(_snap(span_dir), fqa, caps), 3)
    assert s3["3.01  Length entered"]["ok"]
    loc = s3["3.02  Location entered for each of the 4 events"]
    assert not loc["ok"] and loc["detail"] == "missing on event 3"
    gps = s3["GPS locations readable"]
    assert not gps["ok"] and "event 4" in gps["detail"] and "decimal point" in gps["detail"]
    assert s3["3.03  Splice / connection type for each event"]["ok"]
    back = s3["3.06  Distances increase along the span (no negatives)"]
    assert not back["ok"] and back["detail"] == "check event 4"


def test_a_package_built_by_the_fqa_builder_fills_sections_2_and_3(hub, tmp_path, span_dir):
    """End to end: production sheet -> the FQA Builder's own build() into
    the work folder's FQA folder -> the status reads the package, which has
    formulas and no computed values (nobody has opened it in Excel)."""
    from fqa.run_fqa import build
    work = tmp_path / "work"
    (work / "FQA").mkdir(parents=True)
    prod = production_sheet(tmp_path / "prod.xlsx")
    build(prod, str(work / "FQA" / "pkg.xlsm"),
          job_data={"fiber_count": 24}, span_length_m=42000.0)
    fqa, caps = hub.collect_field_files(str(work / "Field"), str(work / "FQA"))
    assert [n for n, _ in fqa] == ["pkg.xlsm"]
    items = hub.project_status(_snap(span_dir), fqa, caps, work=str(work))
    s2, s3 = _sec(items, 2), _sec(items, 3)
    assert s2["2.01  FAT table completed"]["ok"]
    assert s3["3.01  Length entered"]["ok"]
    n = next(k for k in s3 if k.startswith("3.02"))
    assert "12 events" in n and s3[n]["ok"]             # 12 splices on Span 4
    assert s3["3.05  Distance to each event"]["ok"]
    assert s3["3.06  Distances increase along the span (no negatives)"]["ok"]


# ── section 4: traces and the other data files ──────────────────────────
def test_missing_fibers_and_the_forms_fiber_count(hub, tmp_path, span_dir):
    b = sorted(p for p in (span_dir / "B").iterdir() if p.suffix.lower() == ".sor")
    import trace_server
    gone = sorted(trace_server.extract_fiber_num(p.name) for p in b[:2])
    for p in b[:2]:
        p.unlink()
    field = tmp_path / "Field"
    field.mkdir()
    field_fqa(field / "f.xlsm", cells={"F97": 48})
    fqa, caps = hub.collect_field_files(str(field))
    items = hub.project_status(_snap(span_dir), fqa, caps)
    s4 = _sec(items, 4)
    both = s4["4.01  Every fiber shot both ways"]
    assert not both["ok"] and both["detail"].startswith("B missing ")
    assert str(gone[0]) in both["detail"]
    cnt = s4["4.01  48 fibers, as the form says were tested"]
    assert not cnt["ok"] and cnt["detail"] == "22 have both directions"
    # F97 is date-formatted on the customer's form; the tech typed 48.
    assert _sec(items, 1)["1.09  Number of fibers tested"]["detail"] == "48"


def test_each_shoot_is_its_own_dated_folder_and_the_newest_is_final(hub, tmp_path, span_dir):
    work = tmp_path / "work"
    sid1, na, nb = hub.add_shoot(str(span_dir / "A"), str(span_dir / "B"), str(work))
    assert (sid1, na, nb) == ("2026-05-06", 24, 24)           # the .sor files' date
    keep = sorted((span_dir / "A").iterdir())[:3]
    small = tmp_path / "small"
    small.mkdir()
    for p in keep:
        shutil.copy2(p, small / p.name)
    sid2, na, nb = hub.add_shoot(str(small), str(span_dir / "B"), str(work),
                                 date="2026-06-01", label="reshoot")
    assert sid2 == "2026-06-01 reshoot" and na == 3
    # The first shoot is untouched; the newest is final by default.
    assert len(hub._trace_fibers(str(work / "Traces" / "2026-05-06" / "A"))) == 24
    hub.st.session_state.pop("project_final_shoot", None)
    assert hub.final_shoot(str(work))["id"] == "2026-06-01 reshoot"
    # os.path.join: the app builds native paths (backslashes on Windows).
    assert hub.work_trace_dirs(str(work))[0].endswith(
        os.path.join("2026-06-01 reshoot", "A"))
    # Choosing the older shoot makes it final.
    hub.st.session_state["project_final_shoot"] = "2026-05-06"
    assert hub.final_shoot(str(work))["id"] == "2026-05-06"
    hub.st.session_state.pop("project_final_shoot", None)


def test_same_day_shoots_get_distinct_folders(hub, tmp_path, span_dir):
    work = tmp_path / "w"
    s1, _, _ = hub.add_shoot(str(span_dir / "A"), str(span_dir / "B"), str(work))
    s2, _, _ = hub.add_shoot(str(span_dir / "A"), str(span_dir / "B"), str(work))
    assert (s1, s2) == ("2026-05-06", "2026-05-06 (2)")


def test_a_project_from_before_shoots_lists_its_traces_as_a_shoot(hub, tmp_path, span_dir):
    work = tmp_path / "old"
    shutil.copytree(span_dir / "A", work / "Traces" / "A")
    shutil.copytree(span_dir / "B", work / "Traces" / "B")
    shoots = hub.list_shoots(str(work))
    assert [s["id"] for s in shoots] == [hub.LEGACY_SHOOT]
    assert hub.shoot_info(shoots[0])[0] == "2026-05-06"


def test_power_meter_splice_logs_and_hand_ticks(hub, tmp_path, span_dir):
    work = tmp_path / "work"
    (work / "Power Meter").mkdir(parents=True)
    (work / "Power Meter" / "pm1.csv").write_text("x", encoding="utf-8")
    s4 = _sec(hub.project_status(_snap(span_dir), [], [], work=str(work),
                                 manual={"4.03": True, "4.04": "na"}), 4)
    assert s4["4.02  Power meter files"]["ok"] and s4["4.02  Power meter files"]["detail"] == "1 file"
    assert s4["4.03  Files named to the naming convention"]["ok"]
    assert s4["4.04  Splice logs"]["ok"] and "not needed" in s4["4.04  Splice logs"]["detail"]
    assert not s4["4.05  Splice logs and exception documents"]["ok"]


def test_adding_a_production_sheet_keeps_the_old_one(hub, tmp_path):
    work = tmp_path / "work"
    work.mkdir()
    one = production_sheet(tmp_path / "Span 4 v1.xlsx")
    two = production_sheet(tmp_path / "Span 4 v2.xlsx")
    hub._add_production_sheet(one, str(work))
    assert hub.project_production_sheet(str(work)).endswith("Span 4 v1.xlsx")
    hub._add_production_sheet(two, str(work))
    assert hub.project_production_sheet(str(work)).endswith("Span 4 v2.xlsx")
    assert (work / "Production" / "Superseded" / "Span 4 v1.xlsx").is_file()
    n_loc, n_spl, warns = hub._production_summary(hub.project_production_sheet(str(work)))
    assert (n_loc, n_spl) == (14, 12)


def test_dropped_file_is_kept_once_and_a_different_one_gets_a_new_name(hub, tmp_path):
    class Up:
        def __init__(self, name, data):
            self.name, self._d = name, data

        def getvalue(self):
            return self._d
    d = tmp_path / "Field"
    p1, new1 = hub._store_dropped(str(d), Up("x.xlsx", b"one"))
    p2, new2 = hub._store_dropped(str(d), Up("x.xlsx", b"one"))
    p3, new3 = hub._store_dropped(str(d), Up("x.xlsx", b"two"))
    assert (new1, new2, new3) == (True, False, True)
    assert p1 == p2 and p3.endswith("x (2).xlsx")


def test_another_spans_workbook_is_flagged(hub, tmp_path):
    field_fqa(tmp_path / "f.xlsm", cells={"E9": "1 Road, Sitec", "E10": "SITEC"})
    wb = hub.read_fqa_workbook(str(tmp_path / "f.xlsm"))
    assert hub._fqa_names_mismatch(wb, ("SITEA", "SITEB"))
    assert not hub._fqa_names_mismatch(wb, ("SITEC", "SITED"))
    assert not hub._fqa_names_mismatch(wb, ("A", "B"))


# ── the page ─────────────────────────────────────────────────────────────
def test_status_page_in_a_project_shows_the_four_sections(settings_dir, span_dir, monkeypatch):
    monkeypatch.setenv("OTDR_TEST_HOME", "1")
    field = span_dir / "Field"
    field.mkdir()
    capture_sheet(field / "cap.xlsx", [{"loc": "Z-Location", "lat": 1.0, "lon": -2.0, "photos": 2}])
    at = run_streamlit().run()
    next(b for b in at.button if b.label == "📂 Open Recent Project").click().run()
    at.text_input(key="home_folder").set_value(str(span_dir)).run()
    next(b for b in at.button if b.label == "Open This Folder").click().run()
    assert not at.exception, list(at.exception)
    text = " ".join(m.value for m in at.markdown)
    for title in ("1 · Site Survey Data", "2 · FAT", "3 · Event Log", "4 · Data Files"):
        assert title in text
    assert "✓ Z end photos · 2 photos" in text
    # No traces yet: no row, so nothing to run.
    assert not [b for b in at.button if (b.key or "").startswith("run_")]
    # Section 4's copy puts them there, and the tools open.
    at.text_input(key="ps_tr_a").set_value(str(span_dir / "A")).run()
    at.text_input(key="ps_tr_b").set_value(str(span_dir / "B")).run()
    next(b for b in at.button if b.key == "ps_tr_copy").click().run()
    assert not at.exception, list(at.exception)
    assert any("Copied 24 A and 24 B" in s.value for s in at.success)
    # A second shoot is added beside it, and the final choice switches the tools.
    at.text_input(key="ps_tr_label").set_value("reshoot").run()
    next(b for b in at.button if b.key == "ps_tr_copy").click().run()
    assert any("Added a shoot" in s.value for s in at.success)
    # One Final box per shoot row, newest first (the Final column picks the
    # final, 2026-09-27).
    assert {"final_cb_0", "final_cb_1"} <= {c.key for c in at.checkbox}
    # Adding a shoot does not move the final: still the first one.
    assert at.session_state["project_final_shoot"] == "2026-05-06"
    assert at.checkbox(key="final_cb_1").value and not at.checkbox(key="final_cb_0").value
    # The final is for the FQA side only; the tools run the shoot they are
    # opened on (2026-09-27).
    before = at.session_state["view_dir_a_input"]
    at.checkbox(key="final_cb_0").check().run()
    assert at.session_state["project_final_shoot"] == "2026-05-06 reshoot"
    assert at.session_state["view_dir_a_input"] == before
    at.button(key="run_viewer_2026-05-06").click().run()
    assert at.session_state["nav_radio"] == "Viewer"
    assert at.session_state["view_dir_a_input"] == str(span_dir / "Traces" / "2026-05-06" / "A")


def test_an_untouched_event_log_is_empty_not_events_missing_data(hub, tmp_path):
    """Two real 'FQA' files on disk had never had their Event Log filled:
    numbered rows, ' <Select>', a distance of 0 and the form's own blank
    formulas.  That is no events yet, not 101 events missing everything."""
    cells = {}
    for r in range(18, 30):
        cells.update({f"Event Log!B{r}": r - 17, f"Event Log!Q{r}": " <Select>",
                      f"Event Log!AB{r}": 0,
                      f"Event Log!D{r}": f'=IF(B{r}="Site Z","x","")',
                      f"Event Log!N{r}": f'=IF(B{r}="Site Z","NA","")'})
    field_fqa(tmp_path / "f.xlsm", cells=cells)
    wb = hub.read_fqa_workbook(str(tmp_path / "f.xlsm"))
    assert hub.fqa_event_rows(wb) == []


def test_photos_in_the_crews_side_by_side_layout(hub, tmp_path):
    """The packages crews submit put the two site names side by side (Span 4:
    'FLAGLER ILA' in D23, 'Bethune ILA' in K23) with each end's photos in its
    half, two per end: the rack label and the box on the rack."""
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    wb = openpyxl.load_workbook(TEMPLATE, keep_vba=True)
    pics = wb["Pictures"]
    pics["D23"], pics["K23"] = "FLAGLER ILA", "Bethune ILA"
    for anchor in ("A1", "D1", "H1", "L1"):          # where Span 4's four sit
        pics.add_image(XLImage(_png()), anchor)
    wb.save(tmp_path / "f.xlsm")
    got = hub.fqa_photos_per_end(hub.read_fqa_workbook(str(tmp_path / "f.xlsm"))["pictures"])
    assert got == {"A": 2, "Z": 2, "unassigned": 0}


def test_uncaptioned_photos_are_not_guessed_to_an_end(hub, span_dir, tmp_path):
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    wb = openpyxl.load_workbook(TEMPLATE, keep_vba=True)
    for anchor in ("B2", "H2", "O2"):
        wb["Pictures"].add_image(XLImage(_png()), anchor)
    field = tmp_path / "Field"
    field.mkdir()
    wb.save(field / "f.xlsm")
    fqa, caps = hub.collect_field_files(str(field))
    s1 = _sec(hub.project_status(_snap(span_dir), fqa, caps), 1)
    assert not s1["A end photos"]["ok"]
    assert "3 photos on the Pictures tab, but no site names" in s1["A end photos"]["detail"]


# ── the phone job: link out, capture package back ───────────────────────
def test_job_manifest_is_built_from_the_production_sheet(hub, tmp_path):
    prod = production_sheet(tmp_path / "p.xlsx")
    m = hub.job_manifest(prod, "ab12cd34", "Flagler-Bethune", {"fiber_count": 1152})
    assert m["id"] == "ab12cd34" and m["v"] == 1
    info = m["info"]
    assert info["fiberCount"] == 1152
    assert info["section12"]["A"]["aisle"] == "100" and info["section12"]["A"]["bay"] == "007"
    assert info["section12"]["Z"]["bay"] == "008"
    assert len(m["splices"]) == 12
    assert m["splices"][0][0] == 1 and m["splices"][0][1] == "ENTRY"
    # The link carries it whole, and comes back identical.
    link = hub.job_link("https://fc.example/app/", m)
    assert link.startswith("https://fc.example/app/#job=")
    assert hub.job_from_code(link.split("#job=")[1]) == json.loads(json.dumps(m))


def test_job_email_is_an_unsent_draft_with_the_link(hub, tmp_path):
    eml = hub.write_job_email(str(tmp_path), "Field Capture job x", "Tap:\n\nhttps://fc.example/#job=abc\n")
    raw = open(eml, "rb").read().decode()
    assert "X-Unsent: 1" in raw and "https://fc.example/#job=abc" in raw


def _package(path, job, sites=(), splices=()):
    import zipfile
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("capture.json", json.dumps({"format": "otdr-capture", "v": 1, "job": job,
                                               "sites": list(sites), "splices": list(splices)}))
    return path


def test_a_capture_package_ticks_the_boxes_and_the_splice_gps(hub, tmp_path, span_dir):
    field = tmp_path / "Field"
    field.mkdir()
    ok = lambda k: {"key": k, "status": "ok"}
    _package(field / "pkg.zip", "job1", sites=[
        {"site": "A", "floor": "001", "room": "0001", "aisle": "100", "bay": "007",
         "panel": {"rmu": "13", "termination": "LC", "panelType": "OSX"},
         "photos": ["photos/A-1-1.jpg", "photos/A-1-2.jpg"],
         "checks": [ok("rack"), ok("rmu"), ok("toward"), ok("fibers")]},
        {"site": "Z", "aisle": "100", "photos": ["photos/Z-1-1.jpg"],
         "checks": [ok("rack"), {"key": "toward", "status": "bad"}]},
    ], splices=[{"event": 1, "gps": {"lat": 39.4, "lon": -102.9}},
                {"event": 2, "gps": None}, {"event": 3, "gps": {"lat": 39.3, "lon": -102.8}}])
    _package(field / "other.zip", "someoneelse")
    all_pk = hub.collect_capture_packages(str(field))
    assert {n for n, _ in all_pk} == {"pkg.zip", "other.zip"}
    pkgs = [(n, p) for n, p in all_pk if p["job"] == "job1"]
    items = hub.project_status(_snap(span_dir), [], [], pkgs=pkgs, n_splices=3)
    s1, s3 = _sec(items, 1), _sec(items, 3)
    assert s1["A end rack location (floor, room, aisle, bay)"]["ok"]
    assert s1["A end panel details (RMU, connector, panel type)"]["ok"]
    assert s1["A end photos"]["detail"] == "2 photos"
    assert s1["A end labels match the job"]["ok"]
    z = s1["Z end labels match the job"]
    assert not z["ok"] and "far-end label" in z["detail"] and "RMU tags" in z["detail"]
    gps = s3["GPS at every splice point"]
    assert not gps["ok"] and gps["detail"] == "2 of 3 · missing on event 2"
    # A fix typed on the GPS tab fills the gap the phone left.
    items = hub.project_status(_snap(span_dir), [], [], pkgs=pkgs, n_splices=3,
                               gps_manual={"2": "39 21 0.00 N 102 51 0.00 W"})
    gps = _sec(items, 3)["GPS at every splice point"]
    assert gps["ok"] and gps["detail"] == "3 of 3 (1 typed on the GPS tab)"
    assert gps["source"] == "pkg.zip + GPS tab"


def test_not_a_capture_package(hub, tmp_path):
    import zipfile
    with zipfile.ZipFile(tmp_path / "x.zip", "w") as z:
        z.writestr("readme.txt", "hi")
    assert hub.read_capture_package(str(tmp_path / "x.zip")) is None


# ── splice GPS: in order, between the ends ──────────────────────────────
A_END = {"lat": 39.475609, "lon": -103.023762}      # Flagler ILA (production sheet)
Z_END = {"lat": 39.366775, "lon": -102.428434}      # Bethune ILA


def _along(frac, off_lat=0.0):
    return {"lat": A_END["lat"] + frac * (Z_END["lat"] - A_END["lat"]) + off_lat,
            "lon": A_END["lon"] + frac * (Z_END["lon"] - A_END["lon"])}


def test_splice_gps_in_order_between_the_ends_passes(hub):
    sp = [{"event": i + 1, "gps": _along(f)} for i, f in enumerate((0.05, 0.3, 0.6, 0.95))]
    assert hub.splice_gps_problems(sp, A_END, Z_END) == []


def test_splice_gps_out_of_order_is_caught(hub):
    sp = [{"event": 1, "gps": _along(0.2)}, {"event": 2, "gps": _along(0.6)},
          {"event": 3, "gps": _along(0.4)}]           # swapped with the next
    probs = hub.splice_gps_problems(sp, A_END, Z_END)
    assert [p[0] for p in probs] == [3] and "out of order" in probs[0][1]


def test_splice_gps_off_the_route_is_caught(hub):
    sp = [{"event": 1, "gps": _along(0.5, off_lat=0.25)}]   # ~28 km north of the line
    probs = hub.splice_gps_problems(sp, A_END, Z_END)
    assert probs and "not between the ends" in probs[0][1]
    # A route that wanders a few km is still between them.
    assert hub.splice_gps_problems([{"event": 1, "gps": _along(0.5, off_lat=0.03)}],
                                   A_END, Z_END) == []


def test_without_end_fixes_only_missing_gps_is_reported(hub):
    sp = [{"event": 1, "gps": _along(0.6)}, {"event": 2, "gps": None}]
    assert hub.splice_gps_problems(sp, None, None) == [(2, "no GPS fix")]


def test_status_shows_gps_order_and_the_override_reason(hub, tmp_path, span_dir):
    field = tmp_path / "Field"
    field.mkdir()
    _package(field / "p.zip", "j", sites=[{"site": "A", "gps": A_END}, {"site": "Z", "gps": Z_END}],
             splices=[{"event": 1, "gps": _along(0.6)}, {"event": 2, "gps": _along(0.3)}])
    pkgs = hub.collect_capture_packages(str(field))
    s3 = _sec(hub.project_status(_snap(span_dir), [], [], pkgs=pkgs, n_splices=2), 3)
    it = s3["Splice GPS in order and between the ends"]
    assert not it["ok"] and "event 2: out of order" in it["detail"]


def test_a_splice_past_the_z_end_is_not_between_the_ends(hub):
    probs = hub.splice_gps_problems([{"event": 1, "gps": _along(1.2)}], A_END, Z_END)
    assert probs and "not between the ends" in probs[0][1]


def test_job_details_count_before_any_package_is_built(hub, tmp_path, span_dir):
    """A project made from a production sheet starts with what the sheet
    answered, read through the FQA Builder's own cell map."""
    from fqa.job_facts import derive
    prod = hub._read_prod(production_sheet(tmp_path / "p.xlsx"))
    job = json.loads(derive(prod).to_json())
    jd = hub.job_details_workbook(job)
    s1 = _sec(hub.project_status(_snap(span_dir), [(hub.JOB_DETAILS_SOURCE, jd)], []), 1)
    assert s1["1.03  Bldg 1 alias"]["ok"] and s1["1.03  Bldg 1 alias"]["source"] == "job details"
    assert s1["1.09  Number of fibers tested"]["ok"]
    # A real workbook listed first still speaks first.
    field_fqa(tmp_path / "pkg.xlsm", cells={"E10": "FLAGLER ILA"})
    real = hub.read_fqa_workbook(str(tmp_path / "pkg.xlsm"))
    s1 = _sec(hub.project_status(_snap(span_dir), [("pkg.xlsm", real), (hub.JOB_DETAILS_SOURCE, jd)], []), 1)
    assert s1["1.03  Bldg 1 alias"]["source"] == "pkg.xlsm"


def test_the_builders_fiber_count_text_reads_as_a_number(hub):
    assert hub._xl_number("1152 Fibers") == 1152


def test_project_widgets_take_a_second_click(settings_dir, span_dir, monkeypatch):
    """A widget drawn with value=/index= from the project and no key changed
    identity the run after it was used, so the second click was dropped: the
    final shoot could be switched once but not back (2026-09-24).  Every
    project-driven widget is keyed and re-synced; each must take a second
    change."""
    monkeypatch.setenv("OTDR_TEST_HOME", "1")
    at = run_streamlit().run()
    next(b for b in at.button if b.label == "📂 Open Recent Project").click().run()
    at.text_input(key="home_folder").set_value(str(span_dir)).run()
    next(b for b in at.button if b.label == "Open This Folder").click().run()
    for lab in ("", "reshoot"):
        at.text_input(key="ps_tr_a").set_value(str(span_dir / "A")).run()
        at.text_input(key="ps_tr_b").set_value(str(span_dir / "B")).run()
        at.text_input(key="ps_tr_label").set_value(lab).run()
        next(b for b in at.button if b.key == "ps_tr_copy").click().run()
    box = {"2026-05-06 reshoot": "final_cb_0", "2026-05-06": "final_cb_1"}   # newest first
    for want in ("2026-05-06 reshoot", "2026-05-06", "2026-05-06 reshoot"):
        at.checkbox(key=box[want]).check().run()
        assert at.session_state["project_final_shoot"] == want
    # The final's own box cannot be unticked: there is always a final.
    at.checkbox(key="final_cb_0").uncheck().run()
    assert at.session_state["project_final_shoot"] == "2026-05-06 reshoot"
    assert at.checkbox(key="final_cb_0").value
    for want in (True, False, True):
        at.checkbox(key="ps_tick_403").set_value(want).run()
        assert (at.session_state["project_manual"].get("4.03") is True) == want


def test_a_phone_test_package_is_not_a_job_package(hub, tmp_path, span_dir):
    """Test Phone Connection's package (test: true, a photo and a GPS fix)
    is shown as a passed test, never counted as the job's field data and
    never flagged as another job's."""
    import zipfile
    field = tmp_path / "Field"
    field.mkdir()
    with zipfile.ZipFile(field / "OTDR_PhoneTest_test-ab12cd.zip", "w") as z:
        z.writestr("capture.json", json.dumps({"format": "otdr-capture", "v": 1, "test": True,
                                               "job": "test-ab12cd", "gps": {"lat": 39.4, "lon": -102.9}}))
        z.writestr("photos/test.jpg", b"\xff\xd8\xff")
    pk = hub.collect_capture_packages(str(field))
    assert pk and pk[0][1]["test"] is True
    src = (hub.REPO_ROOT if hasattr(hub, "REPO_ROOT") else None)
    app_src = open(hub.__file__, encoding="utf-8").read()
    assert "and not p.get('test')]" in app_src          # excluded from the job and from strays


def test_zfc_capture_package_is_read_and_old_zip_still_is(hub, tmp_path):
    """The phone now sends a .zfc (folder_intake share file); packages from
    before it were plain .zip and must keep counting."""
    import folder_intake as fi
    field = tmp_path / "Field"
    field.mkdir()
    cap = json.dumps({"format": "otdr-capture", "v": 1, "job": "j1", "sites": [], "splices": []}).encode()
    fi.share_write(field / "new", "field-capture", {"capture.json": cap, "photos/A-1-1.jpg": b"\xff\xd8"})
    _package(field / "old.zip", "j1")
    # a shared project is not a capture package, even if saved into Field/
    fi.share_write(field / "proj", "project", {"capture.json": cap})
    names = sorted(n for n, _p in hub.collect_capture_packages(str(field)))
    assert names == ["new.zfc", "old.zip"]
    assert "zfc" in [e.lstrip(".") for e in hub.capture_package_exts()]


def test_phone_page_writes_the_same_container_as_folder_intake():
    """No Node here to run app.js, so pin its .zfc constants to folder_intake's."""
    import re
    import folder_intake as fi
    js = (Path(__file__).resolve().parents[2] / "fieldcapture" / "web" / "app.js").read_text(encoding="utf-8")
    const = dict(re.findall(r"const (SHARE_FORMAT|ZFC_KIND|ZFC_VERSION|ZFC_EXT) = '?([^';]+)'?;", js))
    assert const == {"SHARE_FORMAT": fi.SHARE_FORMAT, "ZFC_KIND": "field-capture",
                     "ZFC_VERSION": str(fi.SHARE_KINDS["field-capture"]["version"]),
                     "ZFC_EXT": fi.share_save_extension("field-capture")}
    assert "'manifest.json'" in js and ".zip`" not in js


# ── unreadable packages are listed, other zips are ignored ────────────────
def test_an_unreadable_package_is_listed_not_dropped(hub, tmp_path):
    import zipfile
    field = tmp_path / "Field"
    field.mkdir()
    (field / "broken.zfc").write_bytes(b"not a zip at all")
    with zipfile.ZipFile(field / "bad.zip", "w") as z:          # ours, but garbled
        z.writestr("capture.json", "{not json")
    with zipfile.ZipFile(field / "photos.zip", "w") as z:       # somebody else's zip
        z.writestr("a.jpg", "x")
    (field / "junk.zip").write_bytes(b"plain bytes")             # not a zip: ignored
    _package(field / "good.zip", "job1")
    bad = []
    got = hub.collect_capture_packages(str(field), unreadable=bad)
    assert [n for n, _ in got] == ["good.zip"]
    names = dict(bad)
    assert set(names) == {"broken.zfc", "bad.zip"}
    assert all(names.values())                                   # a reason each
    why = []
    assert hub.read_capture_package(str(field / "photos.zip"), why) is None and why == []


def test_site_names_read_failure_is_reported_and_falls_back(hub, monkeypatch):
    calls = []
    monkeypatch.setattr(hub, "report_error", lambda where, exc, *a, **k: calls.append(where))
    monkeypatch.setattr(hub, "_engine_extras_from_profile",
                        lambda p: {"SITE_NAMES_FROM_IDENTIFIERS": True})

    class Boom:
        def span_site_names(self, a, b):
            raise ValueError("sidecar")
    monkeypatch.setattr(hub, "_splicereport_json_reader", lambda: Boom())
    monkeypatch.setattr(hub, "_derive_ila", lambda d: ("X" + d[-1], None))
    assert hub._site_names_for("/a", "/b", "p") == ("Xa", "Xb")
    assert calls == ["site names from identifiers"]


def test_a_package_that_will_not_open_is_reported(settings_dir, tmp_path, monkeypatch):
    import error_report
    calls = []
    monkeypatch.setattr(error_report, "report_error",
                        lambda where, exc, *a, **k: calls.append(where))
    monkeypatch.setenv("OTDR_TEST_HOME", "1")
    import json as _json
    import zipfile
    # Unpack into a tmp projects folder: with none saved this once landed
    # in the real ~/Documents/OTDR Projects as "y", "y (2)" ... (2026-09-27).
    root = tmp_path / "Projects"
    root.mkdir()
    (settings_dir / "settings.json").write_text(_json.dumps({"projects_root": str(root)}),
                                                encoding="utf-8")
    # A damaged download is the tech's to fix: the share message, no report.
    bad = tmp_path / "x.otdrproject"
    bad.write_bytes(b"garbage")
    at = run_streamlit().run()
    next(b for b in at.button if b.label == "📂 Open Recent Project").click().run()
    at.text_input(key="open_pkg_path").set_value(str(bad)).run()
    next(b for b in at.button if b.key == "open_pkg").click().run()
    assert any("is not an OTDR Suite file, or it is damaged" in e.value for e in at.error)
    assert calls == []
    # A package that reads but will not open (no project file) is reported.
    empty = tmp_path / "y.otdrproject"
    with zipfile.ZipFile(empty, "w") as z:
        z.writestr("otdrproject.json", _json.dumps(
            {"format": "otdr-suite-project-package", "name": "y"}))
        z.writestr("y/readme.txt", "hi")
    at.text_input(key="open_pkg_path").set_value(str(empty)).run()
    next(b for b in at.button if b.key == "open_pkg").click().run()
    assert any("Could not open y.otdrproject" in e.value for e in at.error)
    assert "project: package import" in calls
    assert list(root.iterdir()) == []              # the refused package left nothing
