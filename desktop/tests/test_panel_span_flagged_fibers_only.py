"""A tie-panel connector cell prints the fibers that FAILED, and only those.

SNARCAAH 1 East (Lumen, 2026-09-23), boss: "we are having issues".  The
span is launch reel -> panel connector -> 62 m tie -> panel connector ->
receive reel, so the grid is Connector / Section / Connector (#118).  On
ribbon 10 (fibers 109-120, the fixture here) the far connector read:

    110  .505  REFL-57.8   <- the only fail at the 0.500 connector gate
    113 -.075  REFL-64.6   <- ordinary A/B mismatch across a panel
    the other ten  .26-.39

and the grid printed ALL TWELVE fibers in one MINT "field gainer" cell:
the negative reading on 113 set is_gainer, the writer's gainer branch runs
before the #203 blank rule, and mint outranks pink.  So the one real fail
was hidden and eleven passing fibers read as flagged.  Across the whole
East report that was 6 whole-ribbon cells (72 fibers shown) for 3 fails.
"""
from __future__ import annotations

import openpyxl

from conftest import run_splicereport, FIXTURE_DIR

PANEL_DIR = FIXTURE_DIR / "panelmint"
MINT = "A5D6A7"


def _run(tmp_path):
    out = tmp_path / "panelmint.xlsx"
    rc, m, stderr = run_splicereport(PANEL_DIR / "A", PANEL_DIR / "B", out,
                                     "SNARCAAH 1 East", "SNARCAAH 1 East")
    assert rc == 0 and m and m.get("ok"), f"runner failed: {stderr[-1200:]}"
    return m, out


def _connector_cells(m, out):
    conn = {c["index"] for c in m["columns"] if c["kind"] == "connector"}
    assert conn, "no connector columns: the premise of this test is gone"
    ws = openpyxl.load_workbook(out)["Splice Report"]
    heads = {c: str(ws.cell(3, c).value or "") for c in range(1, ws.max_column + 1)}
    cols = [c for c, h in heads.items() if h.startswith("Connector @")]
    assert cols, heads
    return [ws.cell(r, c) for c in cols for r in range(4, ws.max_row + 1)
            if ws.cell(r, c).value is not None]


def test_the_real_fail_prints_alone(tmp_path):
    m, out = _run(tmp_path)
    flagged = sorted(c["fiber"] for c in m["cells"] if c["is_flagged"])
    assert flagged == [110], flagged
    cells = _connector_cells(m, out)
    assert [str(c.value) for c in cells] == ["110 .505 REFL-57.8dB"], \
        [str(c.value) for c in cells]


def test_a_connector_just_under_zero_is_not_a_field_gainer(tmp_path):
    m, out = _run(tmp_path)
    conn = {c["index"] for c in m["columns"] if c["kind"] == "connector"}
    f113 = [c for c in m["cells"] if c["fiber"] == 113 and c["splice"] in conn]
    assert any(c["loss"] is not None and c["loss"] < 0 for c in f113), \
        "premise: fiber 113 reads a negative connector loss"
    assert all(c["category"] != "gainer" for c in f113), f113
    for cell in _connector_cells(m, out):
        rgb = (cell.fill.fgColor.rgb or "") if cell.fill.fill_type else ""
        assert not rgb.endswith(MINT), f"mint connector cell: {cell.value}"


def test_every_measurement_still_reaches_the_viewer(tmp_path):
    """Only the GRID drops passing fibers; the manifest keeps all twelve."""
    m, _ = _run(tmp_path)
    conn = {c["index"] for c in m["columns"] if c["kind"] == "connector"}
    fibers = {c["fiber"] for c in m["cells"] if c["splice"] in conn}
    assert fibers == set(range(109, 121)), sorted(fibers)
