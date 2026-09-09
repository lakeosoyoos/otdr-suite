"""Unidirectional report: the declared span start and the Cable End column.

Field report (OGD->SLK, 1152 fibers, 1000 ns, 2026-09-09):

    "Looks like the Unidirectional report is not detecting the end event or
     is not displaying it on the report.  I found it as there is a cable cut
     2,600ish ft past the last splice event."

1148 of the 1152 files carry a GenParams user offset of 1.0044 km: the tech
set the span start on the launch connector, so the OTDR wrote the event table
RELATIVE TO THAT POINT while the DataPts samples still start at the port.
Nothing in the engine read that field.  The tail-box probe indexed the raw
trace at the tabled end (63.13 km) — live glass 1 km BEFORE the cut — found
"light through" on 99.7% of fibers, declared a receive reel, and reported the
cut as a mated connector ("Connector 2") at 62.12 km (the position of fiber 1,
one of 4 files shot WITHOUT a span start and therefore 1 km out of frame with
the rest).  The cable end itself appeared nowhere: the grid had no column for
it, so a cut cable and a healthy one of the same length made identical
workbooks.

These tests pin: the reader exposing the field; the tail-box probe reading
the trace in the raw frame; the declared span start acting as reel evidence
for the minority of untrimmed files; one launch column, not two; and the
Cable End column that shows the tech where the glass stops.
"""
import os
import sys

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, os.path.join(ROOT, 'splicereport'))

import splicereportmatchexfo as E  # noqa: E402
import sor_reader324802a as R      # noqa: E402

FIX = os.path.join(HERE, 'fixtures')
SP = 5e-08
M0 = SP * (299792458.0 / 1.468) / 2.0        # metres per sample


def _ev(km, loss, refl=-55.0, end=False, reflective=True, tot=None):
    t = ('1E' if reflective else '0E') if end else ('1F' if reflective else '0F')
    return {'dist_km': km, 'splice_loss': loss, 'reflection': refl,
            'is_end': end, 'is_reflective': reflective, 'type': t + '9999LS',
            'time_of_travel': (0 if tot is None else tot) if km == 0 else 1}


def _trace(n=30000, dead_from_km=None, noise=0.004, seed=3):
    rng = np.random.RandomState(seed)
    x = np.arange(n)
    tr = 47.0 + 0.19 * (x * M0 / 1000.0) + rng.normal(0, noise, n)
    if dead_from_km is not None:
        tr[int(dead_from_km * 1000 / M0):] = 64.0
    return tr


def _rec(events, trace, user_offset_km=0.0, fnum=1):
    return {'events': list(events), 'trace': trace,
            'exfo_sampling_period': SP, 'fxd_pulse_ns': 10.0,
            'num_points': len(trace), 'filename': f'f{fnum:04d}.sor',
            'gen_fiber_id': f'{fnum:04d}', 'user_offset_km': user_offset_km}


# ── 1. the reader exposes the declared span start ─────────────────────────

def test_reader_exposes_declared_span_start():
    pre = R.parse_sor_full(os.path.join(FIX, 'launchreel', 'BARTUL001_1550.sor'),
                           trim=False)
    assert abs(pre['user_offset_km'] - 1.0044) < 0.002
    raw = R.parse_sor_full(os.path.join(FIX, 'launchreel', 'SUIWSC0242.sor'),
                           trim=False)
    assert raw['user_offset_km'] == 0.0


# ── 2. the tail-box probe reads the trace where the glass actually is ─────

def _pretrimmed(dead_from_km, user_offset_km):
    """Table reel-relative (end tabled at 10.0 km); trace port-relative."""
    ev = [_ev(0.0, 0.2, tot=0), _ev(5.0, 0.05, reflective=False),
          _ev(10.0, 0.0, refl=-46.0, end=True)]
    return _rec(ev, _trace(dead_from_km=dead_from_km), user_offset_km)


def test_tail_box_probe_uses_the_declared_span_start():
    # The OGD->SLK shape: table end 10.0, trace dies 1.0 km later in the raw
    # frame because the first 1.0 km of samples are the launch reel.
    cut = _pretrimmed(dead_from_km=11.0, user_offset_km=1.0)
    present, frac = E.uni_detect_tail_box({1: cut})
    assert present is False and frac == 0.0
    # Same table, but the glass genuinely carries on past the end: a reel.
    reel = _pretrimmed(dead_from_km=None, user_offset_km=1.0)
    present, frac = E.uni_detect_tail_box({1: reel})
    assert present is True and frac == 1.0
    # Untrimmed control, unchanged: no offset, trace dies at the tabled end.
    bare = _pretrimmed(dead_from_km=10.0, user_offset_km=0.0)
    assert E.uni_detect_tail_box({1: bare})[0] is False


def test_raw_frame_offset_is_reel_plus_declared_start():
    r = _pretrimmed(dead_from_km=11.0, user_offset_km=1.0)
    assert E._uni_raw_frame_offset_km(r) == 1.0
    r['_trace_offset_km'] = 1.25
    assert E._uni_raw_frame_offset_km(r) == 1.25


# ── 3. a declared span start is reel evidence for the untrimmed minority ──

def _mixed_direction():
    fibers = {}
    for f in range(1, 10):                                   # 9 pre-trimmed
        fibers[f] = _rec([_ev(0.0, 0.1, tot=0),
                          _ev(5.0, 0.06, reflective=False),
                          _ev(10.0, 0.0, refl=-46.0, end=True)],
                         _trace(dead_from_km=11.0), 1.0, fnum=f)
    fibers[10] = _rec([_ev(0.0, 0.0, refl=-80.0, tot=0),      # OTDR port
                       _ev(1.0, 0.1),                          # launch connector
                       _ev(6.0, 0.06, reflective=False),       # the same splice
                       _ev(11.0, 0.0, refl=-46.0, end=True)],  # the same end
                      _trace(dead_from_km=11.0), 0.0, fnum=10)
    return fibers


def test_declared_span_start_normalizes_the_untrimmed_minority():
    fibers = _mixed_direction()
    assert E.launch_reel_consensus_km(list(fibers.values())) is None  # 1 of 10 tabled it
    E.uni_normalize_all(fibers)
    ten = fibers[10]
    assert [round(e['dist_km'], 3) for e in ten['events']] == [0.0, 5.0, 10.0]
    assert ten['_uni_event_offset_km'] == 1.0
    for f in range(1, 10):
        assert fibers[f]['_uni_event_offset_km'] == 0.0
        assert [round(e['dist_km'], 3) for e in fibers[f]['events']] == [0.0, 5.0, 10.0]
    # Every fiber maps table -> raw trace by the same 1.0 km.
    assert {round(r['_trace_offset_km'], 6) for r in fibers.values()} == {1.0}


def test_one_launch_column_across_trimmed_and_untrimmed_files():
    fibers = _mixed_direction()
    E.uni_normalize_all(fibers)
    span = E.uni_auto_detect_span(fibers)
    tb, _ = E.uni_detect_tail_box(fibers)
    assert tb is False
    conns = E.uni_find_connectors(fibers, span, tail_box=tb)
    assert len(conns) == 10 and all(c['is_launch'] for c in conns)
    cols = E.uni_cluster_connectors(conns)
    assert len(cols) == 1 and cols[0]['is_launch']
    # And the cut is not a connector: nothing dark, nothing at the far end.
    assert all(c['position_km'] == 0.0 for c in conns)
    assert not any(c['dark'] for c in conns)


# ── 4. the Cable End column ───────────────────────────────────────────────

def test_end_column_lists_fibers_that_reach_the_end():
    fibers = {
        1: _rec([_ev(0.0, 0.1, tot=0), _ev(10.00, 0.0, refl=-45.0, end=True)], _trace(), 1.0, 1),
        2: _rec([_ev(0.0, 0.1, tot=0), _ev(10.01, 0.0, refl=-50.0, end=True)], _trace(), 1.0, 2),
        3: _rec([_ev(0.0, 0.1, tot=0), _ev(4.00, 0.0, refl=-30.0, end=True)], _trace(), 1.0, 3),
        4: _rec([_ev(0.0, 0.1, tot=0), _ev(10.00, 0.0, refl=0.0, end=True, reflective=False)],
                _trace(), 1.0, 4),
    }
    cols = E.uni_end_column(fibers, span_km=10.0)
    assert len(cols) == 1
    col = cols[0]
    assert col['kind'] == 'end'
    assert col['position_km_display'] == 10.0
    assert set(col['end_members']) == {1, 2, 4}          # 3 broke: Break column's
    assert col['end_members'][1] == -45.0 and col['end_members'][4] is None
    # A receive reel marks the end as a connector already.
    assert E.uni_end_column(fibers, span_km=10.0, tail_box=True) == []
    assert E.uni_end_column({}, span_km=10.0) == []


def test_end_cell_text_is_reflective_vocabulary():
    assert E.uni_format_end_cell([(1, -45.0), (2, -50.0)], range(1, 3)) == "REFL-45.0dB"
    assert E.uni_format_end_cell([(1, -45.0), (2, -50.0)], range(1, 13)) == "F1,F2 REFL-45.0dB"
    assert E.uni_format_end_cell([(1, None), (2, None)], range(1, 3)) == "end"
    assert E.uni_format_end_cell([], range(1, 13)) == ""


def test_end_column_is_data_not_a_flag(tmp_path):
    fibers = {f: _rec([_ev(0.0, 0.1, tot=0), _ev(10.0, 0.0, refl=-45.0 - f, end=True)],
                      _trace(), 1.0, f) for f in range(1, 13)}
    cols = E.uni_end_column(fibers, span_km=10.0)
    columns = E.uni_build_columns([], cols, [])
    grid = E.uni_build_ribbon_grid(fibers, columns, 12)
    assert len(grid[(0, 0)]) == 12
    assert E.uni_flagged_event_rows(grid, columns) == []
    out = str(tmp_path / 'uni.xlsx')
    E.uni_write_xlsx(grid, columns, 12, 12, 10.0, out, site_a='OGD', site_b='SLK',
                     fibers=fibers)
    import openpyxl
    ws = openpyxl.load_workbook(out)['Unidir Events']
    hdr = {ws.cell(r, 2).value for r in range(1, 6)}
    assert "Cable End" in hdr and "10.00 km" in hdr
    cells = [ws.cell(r, 2).value for r in range(1, ws.max_row + 1)]
    assert "REFL-46.0dB" in cells                      # strongest of -46..-57
