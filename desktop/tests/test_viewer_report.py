"""The Viewer's Summary Report: the event panel, then a page per fibre, as a
PDF or an Excel workbook (Robert, 2026-09-28: "create a report that shows the
traces and the event panel from viewer ... an option to do it in pdf or excel
sheet"; 2026-09-30: "get rid of the page that has all the traces shown at
once ... start with the full event panel that has all traces and then move
right into the fiber by fiber").

The browser hands the server each fibre's chart as a PNG and the event table
as plain cells (text, colours resolved from the page's CSS, colspan /
rowspan); the server lays them out.  These tests pin the server half end to end and the
browser half's contract with the three event grids."""
import json
import os
import re
import struct
import sys
import threading
import zlib
from http.server import HTTPServer
from urllib.error import HTTPError
from urllib.request import Request, urlopen

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(ROOT, 'viewer'))
import trace_server as T  # noqa: E402
from test_viewer_foreign_post import ROUNDS, foreign_status  # noqa: E402

SRC = open(os.path.join(ROOT, 'viewer', 'viewer.html'), encoding='utf-8').read()


def _png(w=60, h=24):
    """A real, tiny PNG (solid red), built by hand so the test needs no Pillow."""
    raw = b''.join(b'\x00' + b'\xe7\x4c\x3c' * w for _ in range(h))

    def chunk(t, d):
        return struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return (b'\x89PNG\r\n\x1a\n'
            + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0))
            + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b''))


def _data_url(png):
    import base64
    return 'data:image/png;base64,' + base64.b64encode(png).decode('ascii')


# Styles as the browser sends them: one entry per distinct look.
HDR, PLAIN, FAIL, LAB = 0, 1, 2, 3
STYLES = [
    {'bg': '#e4edf7', 'fg': '#1f2d3d', 'b': True, 'al': 'c'},
    {'bg': None, 'fg': '#1f2d3d', 'b': False, 'al': 'c'},
    {'bg': '#e74c3c', 'fg': '#ffffff', 'b': True, 'al': 'c'},
    {'bg': '#eef3f8', 'fg': '#2e86de', 'b': True, 'al': 'l'},
]


def _suite_like_table(n_events=3, fibres=(64, 65)):
    """The shape of the Viewer's bidirectional grids: four identifier columns
    spanning a three-row header, then a Loss / Refl. pair per event."""
    lead = [{'t': t, 's': LAB if i == 0 else HDR, 'rs': 3}
            for i, t in enumerate(['Identifiers', 'P/F', 'λ (nm)', 'Dir.'])]
    h1 = lead + [{'t': f'Splice {i + 1} (2/2)', 's': HDR, 'cs': 2} for i in range(n_events)]
    h2 = [{'t': f"{i + 1}.0000km, {3281 * (i + 1):,}'", 's': HDR, 'cs': 2} for i in range(n_events)]
    h3 = [c for _ in range(n_events)
          for c in ({'t': 'Loss\n(dB)', 's': HDR}, {'t': 'Refl.\n(dB)', 's': HDR})]
    body = []
    for f in fibres:
        for d in ('A→B', 'B→A', 'Average'):
            row = [{'t': f'F{f}', 's': LAB, 'dot': '#ffffff'}, {'t': '✓', 's': PLAIN},
                   {'t': '1550', 's': PLAIN}, {'t': d, 's': PLAIN}]
            for i in range(n_events):
                bad = (f == 64 and d == 'Average' and i == 0)
                row.append({'t': '0.213' if bad else '0.041', 's': FAIL if bad else PLAIN,
                            **({'tip': 'Splice Report: 64 .213 (reburn)'} if bad else {})})
                row.append({'t': '---', 's': PLAIN})
            body.append(row)
    foot = [[{'t': lab, 's': LAB}, {'t': '', 's': PLAIN}, {'t': '', 's': PLAIN}, {'t': '', 's': PLAIN}]
            + [c for _ in range(n_events) for c in ({'t': '0.041', 's': PLAIN}, {'t': '', 's': PLAIN})]
            for lab in ('Minimum', 'Maximum', 'Average')]
    return {'title': 'Event Table', 'note': 'F64: OTDR Suite table', 'lead': 4,
            'head': [h1, h2, h3], 'body': body, 'foot': foot}


def _payload(fmt, **kw):
    p = {'format': fmt, 'dest': '', 'name': 'Viewer Report F64',
         'title': 'OTDR Viewer Report: TESTJOB', 'subtitle': 'Generated today',
         'meta': [['Mode', 'OTDR Suite'], ['Fibres', '64-65 (2 fibres, 4 traces)']],
         'styles': STYLES, 'tables': [_suite_like_table()]}
    p.update(kw)
    return p


@pytest.fixture
def downloads(tmp_path, monkeypatch):
    d = tmp_path / 'Downloads'
    d.mkdir()
    monkeypatch.setenv('OTDR_DOWNLOADS_DIR', str(d))
    monkeypatch.setitem(T.CONFIG, 'dir_a', None)
    monkeypatch.setitem(T.CONFIG, 'dir_b', None)
    return d


# ─── the server half ──────────────────────────────────────────────────────

def test_a_pdf_lands_in_downloads_and_never_over_an_earlier_one(downloads):
    one = T.write_viewer_report(_payload('pdf'))
    two = T.write_viewer_report(_payload('pdf'))
    assert one['path'] == str(downloads / 'Viewer Report F64.pdf')
    assert two['path'] == str(downloads / 'Viewer Report F64 (2).pdf')
    for p in (one['path'], two['path']):
        with open(p, 'rb') as f:
            assert f.read(5) == b'%PDF-'
    assert not [n for n in os.listdir(downloads) if n.endswith('.part')]


def test_a_wide_table_goes_over_several_pages_of_columns(downloads):
    """40 events is far wider than a landscape page: the PDF must still build
    (reportlab would otherwise run the table off the page edge)."""
    wide = _payload('pdf', tables=[_suite_like_table(n_events=40, fibres=range(1, 30))])
    out = T.write_viewer_report(wide)
    assert os.path.getsize(out['path']) > 5000


def test_the_workbook_keeps_numbers_colours_merges_and_a_frozen_header(downloads):
    from openpyxl import load_workbook
    out = T.write_viewer_report(_payload('xlsx'))
    assert out['name'] == 'Viewer Report F64.xlsx'
    wb = load_workbook(out['path'])
    assert wb.sheetnames == ['Report', 'Event Table']
    rep = wb['Report']
    assert rep['A1'].value == 'OTDR Viewer Report: TESTJOB'
    ws = wb['Event Table']
    # title row 1, note row 2, the table from row 4: three header rows, then
    # the body; frozen below the header and right of the four identifiers
    assert ws.freeze_panes == 'E7'
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert 'A4:A6' in merged                            # Identifiers spans the header
    assert 'E4:F4' in merged                            # an event's Loss + Refl. pair
    bad = ws['E9']                                      # F64 Average, first event
    assert bad.value == pytest.approx(0.213) and bad.number_format == '0.000'
    assert bad.fill.fgColor.rgb.endswith('E74C3C')
    assert bad.font.b and bad.font.color.rgb.endswith('FFFFFF')
    assert bad.comment is not None and '.213' in bad.comment.text
    assert ws['E7'].value == pytest.approx(0.041) and ws['E7'].fill.fill_type is None
    assert ws['F7'].value == '---'                      # text stays text
    assert ws['C7'].value == 1550 and ws['C7'].number_format == '0'
    assert ws['E6'].value == 'Loss\n(dB)' and ws['E6'].alignment.wrap_text


def test_numbers_are_numbers_and_everything_else_stays_text():
    assert T._xl_value('0.113') == (0.113, '0.000')
    assert T._xl_value('-53.9') == (-53.9, '0.0')
    assert T._xl_value('1550') == (1550.0, '0')
    assert T._xl_value("0.2600km, 853'") == ("0.2600km, 853'", None)
    assert T._xl_value('---') == ('---', None)
    assert T._xl_value('F64') == ('F64', None)


def test_the_grid_places_cells_under_rowspans():
    rows = [[{'t': 'a', 'rs': 2}, {'t': 'b', 'cs': 2}], [{'t': 'c'}, {'t': 'd'}]]
    anchors, ncols = T._report_grid(rows)
    assert ncols == 3
    assert [(r, c, rs, cs, x['t']) for r, c, rs, cs, x in anchors] == [
        (0, 0, 2, 1, 'a'), (0, 1, 1, 2, 'b'), (1, 1, 1, 1, 'c'), (1, 2, 1, 1, 'd')]


def test_column_blocks_never_cut_an_event_pair_and_repeat_the_lead():
    rows, _nh, _nb = T._report_rows(_suite_like_table(n_events=12))
    anchors, ncols = T._report_grid(rows)
    cut = T._report_cuts(anchors, ncols)
    widths = [40.0] * ncols
    blocks = T._report_col_blocks(widths, cut, 4, 400.0)
    assert blocks[0][0] == 4 and blocks[-1][1] == ncols
    for (s, e), (s2, _e2) in zip(blocks, blocks[1:]):
        assert e == s2                                  # contiguous, nothing lost
    for s, e in blocks:
        assert (s - 4) % 2 == 0 and (e - 4) % 2 == 0    # Loss + Refl. stay together
        assert 4 * 40 + (e - s) * 40 <= 400
    # one block wider than the page is kept whole rather than looping
    assert T._report_col_blocks([10.0] * 4 + [500.0, 500.0], [True] * 4 + [True, False, True], 4, 300.0) \
        == [(4, 6)]


def test_reports_stay_out_of_the_trace_folders(tmp_path, downloads, monkeypatch):
    a, b = tmp_path / 'A', tmp_path / 'B'
    a.mkdir()
    b.mkdir()
    monkeypatch.setitem(T.CONFIG, 'dir_a', str(a))
    monkeypatch.setitem(T.CONFIG, 'dir_b', str(b))
    for bad in (str(a), str(a / 'reports'), str(b)):
        with pytest.raises(ValueError, match='trace folder'):
            T.report_dest(bad)
    assert T.report_dest('') == str(downloads)
    assert T.report_dest('Span 7') == str(downloads / 'Span 7')
    assert T.report_dest(str(tmp_path / 'elsewhere')) == str(tmp_path / 'elsewhere')
    with pytest.raises(ValueError):
        T.report_dest('sub/dir')


def test_a_file_name_cannot_escape_the_folder(downloads):
    out = T.write_viewer_report(_payload('pdf', name='..\\..\\evil:name?.pdf'))
    assert os.path.dirname(out['path']) == str(downloads)
    assert not re.search(r'[\\/:?]', out['name']) and out['name'].endswith('evil_name_.pdf')
    assert not out['name'].startswith('.')             # no hidden file either


def test_open_only_opens_a_report_this_viewer_wrote(downloads, monkeypatch):
    calls = []
    monkeypatch.setattr(T.subprocess, 'Popen', lambda *a, **k: calls.append(a))
    if hasattr(T.os, 'startfile'):
        monkeypatch.setattr(T.os, 'startfile', lambda p: calls.append(p))
    with pytest.raises(ValueError, match='not a report'):
        T.open_report(str(downloads / 'something.pdf'))
    out = T.write_viewer_report(_payload('pdf'))
    T.open_report(out['path'])
    T.open_report(out['path'], reveal=True)
    assert len(calls) == 2


def _post(port, path, body, origin=None):
    req = Request(f'http://127.0.0.1:{port}{path}', data=json.dumps(body).encode('utf-8'),
                  headers={'Content-Type': 'application/json', **({'Origin': origin} if origin else {})})
    try:
        with urlopen(req, timeout=60) as r:
            return r.status, json.loads(r.read().decode('utf-8'))
    except HTTPError as e:
        raw = e.read().decode('utf-8', 'replace')
        try:
            return e.code, json.loads(raw)
        except ValueError:
            return e.code, {'raw': raw}


def test_the_route_writes_the_report_and_refuses_a_foreign_page(downloads):
    srv = HTTPServer(('127.0.0.1', 0), T.Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        port = srv.server_port
        with urlopen(f'http://127.0.0.1:{port}/api/report_defaults', timeout=30) as r:
            assert json.loads(r.read().decode('utf-8')) == {'dest': str(downloads)}
        st, j = _post(port, '/api/report', _payload('xlsx'))
        assert st == 200 and j['ok'] and os.path.isfile(j['path'])
        st, j = _post(port, '/api/report', _payload('docx'))
        assert st == 400 and 'format' in j['error']
        st, _j = _post(port, '/api/report', _payload('pdf'), origin='http://evil.example')
        assert st == 403
        st, j = _post(port, '/api/report_open', {'path': '/etc/hosts'})
        assert st == 400 and 'not a report' in j['error']
        st, j = _post(port, '/api/report_begin', {})
        assert st == 200 and j['token'] in T._REPORT_UPLOADS
        req = Request(f'http://127.0.0.1:{port}/api/report_image?token={j["token"]}&name=fibre-7',
                      data=_png(), headers={'Content-Type': 'image/png'})
        with urlopen(req, timeout=30) as r:
            assert json.loads(r.read().decode('utf-8'))['ok']
        st, k = _post(port, '/api/report', _payload('pdf', token=j['token'],
                                                    fibres=[_fibre(7, 'ref:fibre-7')]))
        assert st == 200 and os.path.isfile(k['path']) and j['token'] not in T._REPORT_UPLOADS
        st, _j = _post(port, '/api/report_begin', {}, origin='http://evil.example')
        assert st == 403
    finally:
        srv.shutdown()
        srv.server_close()


def test_a_foreign_page_leaves_nothing_behind_on_any_report_route(downloads, monkeypatch):
    """Each of the four is sent what the Viewer's own page sends, a body that
    would be acted on: a session opened, a chart stored, a report written, a
    report opened.  From a foreign page it is a 403 every time, read without
    a dropped connection (the last line of the test above lost its answer
    once in the Windows build), and nothing is left behind."""
    opened = []
    monkeypatch.setattr(T.subprocess, 'Popen', lambda *a, **k: opened.append(a))
    if hasattr(T.os, 'startfile'):
        monkeypatch.setattr(T.os, 'startfile', lambda p: opened.append(p))
    written = T.write_viewer_report(_payload('pdf'))['path']
    tok = T.report_begin()
    folder = T._REPORT_UPLOADS[tok][0]
    bodies = {
        '/api/report_begin': b'{}',
        f'/api/report_image?token={tok}&name=fibre-7': _png(),
        '/api/report': json.dumps(_payload('xlsx')).encode('utf-8'),
        '/api/report_open': json.dumps({'path': written}).encode('utf-8'),
    }
    before = (sorted(os.listdir(downloads)), sorted(T._REPORT_UPLOADS), sorted(T._REPORTS_WRITTEN))
    srv = HTTPServer(('127.0.0.1', 0), T.Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        port = srv.server_port
        for _ in range(ROUNDS):
            for path, body in bodies.items():
                assert foreign_status(port, path, body) == 403, path
        # more than the sockets hold: the old refusal closed on a page that
        # was still sending, and no machine saw the 403
        for path, body in bodies.items():
            assert foreign_status(port, path, body + b' ' * (4 * 1024 * 1024)) == 403, path
        assert (sorted(os.listdir(downloads)), sorted(T._REPORT_UPLOADS),
                sorted(T._REPORTS_WRITTEN)) == before
        assert os.listdir(folder) == [] and opened == []
        # and the same four from the Viewer's own page are acted on, so it
        # was the origin that refused them
        for path, body in bodies.items():
            assert foreign_status(port, path, body, origin=f'http://127.0.0.1:{port}') == 200, path
        assert len(T._REPORT_UPLOADS) == len(before[1]) + 1
        assert os.listdir(folder) == ['fibre-7.png'] and len(opened) == 1
        assert len(os.listdir(downloads)) == len(before[0]) + 1
    finally:
        srv.shutdown()
        srv.server_close()
        for t in set(T._REPORT_UPLOADS) - set(before[1]) | {tok}:
            T._report_end(t)


def _fibre_table():
    """A fibre page's table, as the browser turns the combined one on its side."""
    head = [[{'t': 'Event', 's': HDR, 'rs': 2}, {'t': 'Type', 's': HDR, 'rs': 2},
             {'t': 'Location / length', 's': HDR, 'rs': 2},
             {'t': 'A→B', 's': HDR, 'cs': 2}, {'t': 'B→A', 's': HDR, 'cs': 2}, {'t': 'Average', 's': HDR}],
            [{'t': 'Loss\n(dB)', 's': HDR}, {'t': 'Refl.\n(dB)', 's': HDR},
             {'t': 'Loss\n(dB)', 's': HDR}, {'t': 'Refl.\n(dB)', 's': HDR}, {'t': 'Loss\n(dB)', 's': HDR}]]
    body = [[{'t': f'Event {i}', 's': PLAIN}, {'t': 'Non-reflective', 's': PLAIN},
             {'t': f"{i}.0000km, {3281 * i:,}'", 's': PLAIN},
             {'t': '0.041', 's': PLAIN}, {'t': '---', 's': PLAIN},
             {'t': '0.213' if i == 2 else '0.052', 's': FAIL if i == 2 else PLAIN}, {'t': '---', 's': PLAIN},
             {'t': '0.047', 's': PLAIN}] for i in range(1, 19)]
    return {'title': 'Events', 'lead': 3, 'head': head, 'body': body, 'foot': []}


def _fibre(n, png):
    return {'title': f'F{n}', 'png': png, 'note': '',
            'meta': [['A→B file', f'A{n:04d}.sor'], ['B→A file', f'B{n:04d}.sor'], ['P/F', '✓']],
            'key': [{'label': f'F{n} A→B', 'color': '#1f77b4'}], 'table': _fibre_table()}


def _page_sizes(path):
    with open(path, 'rb') as f:
        raw = f.read()
    return [(float(w), float(h)) for w, h in
            re.findall(rb'/MediaBox \[ 0 0 ([\d.]+) ([\d.]+) \]', raw)]


def test_the_report_opens_on_the_event_table_with_no_chart_of_every_trace(downloads):
    """Robert, 2026-09-30: "get rid of the page that has all the traces shown
    at once ... start with the full event panel that has all traces".  The
    event table sits under the report's facts on page one, and a chart of
    every trace (as a page from before this change still sends) is drawn in
    neither format."""
    from openpyxl import load_workbook
    stale = {'images': [{'caption': 'Traces', 'png': _data_url(_png()), 'note': 'Distance 0 to 1 km'}],
             'key': [{'label': 'F64 A→B', 'color': '#1f77b4'}]}
    pdf = T.write_viewer_report(_payload('pdf', **stale))
    assert _page_sizes(pdf['path']) == [(792.0, 612.0)]       # facts + event table: one page
    wb = load_workbook(T.write_viewer_report(_payload('xlsx', **stale))['path'])
    assert wb.sheetnames == ['Report', 'Event Table']
    rep = wb['Report']
    assert len(rep._images) == 0
    assert [rep.cell(r, 1).value for r in range(1, rep.max_row + 1)] == \
        ['OTDR Viewer Report: TESTJOB', 'Generated today', None, 'Mode', 'Fibres']


def test_a_tall_event_table_still_starts_under_the_facts(downloads, monkeypatch):
    """The browser check, 2026-10-01: 24 fibres A+B left page one with the
    facts alone.  The heading kept with its table carried a table that fits a
    page, but not the rest of page one, over whole.  The table starts under
    the facts and carries on over the page; a table after it still gets a
    page of its own."""
    from reportlab.platypus import BaseDocTemplate, Paragraph
    drawn = []
    monkeypatch.setattr(BaseDocTemplate, 'afterFlowable', lambda self, f: drawn.append(
        (self.page, f.getPlainText())) if isinstance(f, Paragraph) else None)
    tall = _suite_like_table(fibres=range(1, 15))          # 42 rows: a page, not under the facts
    later = dict(_suite_like_table(), title='Traces', note='')
    pdf = T.write_viewer_report(_payload('pdf', tables=[tall, later]))
    page = {t: p for p, t in drawn}
    assert page['OTDR Viewer Report: TESTJOB'] == 1
    assert page['Event Table'] == 1
    assert page['Traces'] == 3                               # table over pages 1-2, then its own
    assert len(_page_sizes(pdf['path'])) == 3


def test_a_page_per_fibre_follows_the_event_table_in_portrait(downloads):
    """Robert, 2026-09-28: the combined table, then a page per fibre as
    FastReporter prints them; 2026-09-30: "move right into the fiber by
    fiber".  The event table stays landscape and the first fibre's portrait
    page comes straight after it: chart, files, events down the page."""
    tok = T.report_begin()
    T.report_put_image(tok, 'fibre-65', _png())
    out = T.write_viewer_report(_payload('pdf', token=tok,
                                         fibres=[_fibre(64, _data_url(_png())), _fibre(65, 'ref:fibre-65')]))
    sizes = _page_sizes(out['path'])
    assert sizes == [(792.0, 612.0), (612.0, 792.0), (612.0, 792.0)], sizes   # table, F64, F65
    assert tok not in T._REPORT_UPLOADS                        # the charts sent ahead are gone


def test_the_workbook_has_a_sheet_per_fibre(downloads):
    from openpyxl import load_workbook
    out = T.write_viewer_report(_payload('xlsx', fibres=[_fibre(64, _data_url(_png())), _fibre(65, '')]))
    wb = load_workbook(out['path'])
    assert wb.sheetnames == ['Report', 'Event Table', 'F64', 'F65']
    ws = wb['F64']
    assert ws['A1'].value == 'F64' and len(ws._images) == 1
    head = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'Event')
    assert [ws.cell(head, c).value for c in range(1, 5)] == ['Event', 'Type', 'Location / length', 'A→B']
    bad = ws.cell(head + 3, 6)                                 # Event 2, B→A loss
    assert bad.value == pytest.approx(0.213) and bad.fill.fgColor.rgb.endswith('E74C3C')
    assert wb['F65']['A1'].value == 'F65' and len(wb['F65']._images) == 0


def test_the_chart_upload_refuses_what_it_should(downloads):
    with pytest.raises(ValueError, match='session'):
        T.report_put_image('nope', 'fibre-1', _png())
    tok = T.report_begin()
    try:
        with pytest.raises(ValueError, match='name'):
            T.report_put_image(tok, '../evil', _png())
        with pytest.raises(ValueError, match='PNG'):
            T.report_put_image(tok, 'fibre-1', b'GIF89a...')
        folder = T._REPORT_UPLOADS[tok][0]
        T.report_put_image(tok, 'fibre-1', _png())
        assert os.listdir(folder) == ['fibre-1.png']
        # a ref that was never sent, or one that tries to leave the folder,
        # is no chart at all
        assert T._report_img('ref:fibre-2', folder) == (None, None)
        assert T._report_img('ref:../x', folder) == (None, None)
        src, size = T._report_img('ref:fibre-1', folder)
        assert src == os.path.join(folder, 'fibre-1.png') and size == (60, 24)
    finally:
        T._report_end(tok)
    assert tok not in T._REPORT_UPLOADS and not os.path.exists(folder)


# ─── the browser half ─────────────────────────────────────────────────────

def _fn(name):
    body = SRC.split(f'function {name}(', 1)[1]
    return body.split('\nfunction ', 1)[0]


def test_the_toolbar_button_makes_the_summary_report():
    """Robert, 2026-09-28: "we will call them Summary Report", made from the
    Viewer only."""
    # Its own toolbar group at the bar's right end, beside the gear (Robert 2026-10-01).
    group = SRC.split('<button id="btn-report"', 1)[0].rsplit('<div class="group" id="tb-right">', 1)[1]
    assert '</div>' not in group
    # One button for both reports since Robert 2026-10-02 (test_olts_reader):
    # Generate Report, the Summary Report when traces are what is loaded.
    assert '>Generate Report…</button>' in SRC
    assert "document.getElementById('btn-report').onclick = generateReport;" in SRC
    assert "generateReportKind() === 'olts' ? showOltsDialog(gOlts) : showReportDialog()" in SRC
    assert '`<h3>Summary Report</h3>`' in SRC and '`<h3>Summary Report Saved</h3>`' in SRC
    assert "return `Summary Report ${job" in _fn('reportDefaultName')
    assert "title: 'Summary Report' + (job" in _fn('reportPayload')
    assert 'Viewer Report' not in SRC


def test_an_unnamed_report_is_a_summary_report(downloads):
    p = _payload('pdf')
    del p['name'], p['title']
    out = T.write_viewer_report(p)
    assert out['name'] == 'Summary Report.pdf'


def test_every_grid_hands_the_report_every_row_not_the_slice_on_screen():
    """The panel is virtual: the DOM only holds the rows on screen.  Each of
    the three grids registers its FULL row list, and a fresh render clears the
    last one so a table that is still loading is never reported stale."""
    assert "gTableExport = { table, headExtra: statNameRow, rows: () => shown.map(rowHtml) };" \
        in _fn('renderFastReporterGrid')
    assert "gTableExport = { table, rows: () => descs.map(rowHtml) };" in _fn('paintFrBidiGrid')
    assert "gTableExport = { table, rows: () => descs.map(rowHtml) };" in _fn('paintSuiteBidiGrid')
    assert 'gTableExport = null;' in _fn('renderEventTable')
    snap = _fn('reportEventTable')
    assert 'ex.rows()' in snap and "classList.contains('fr-spacer')" in snap


def test_cell_colours_come_from_the_page_css():
    probe = _fn('reportStyleProbe')
    assert "table.className = 'fr-table';" in probe       # the panel's own class
    assert 'getComputedStyle(td)' in probe


def test_the_chart_is_drawn_at_print_scale_and_put_back():
    """Every chart in the report is the Viewer's own draw() at print size.
    Whatever the drawing does in between -- a hundred fibres drawn alone,
    awaits while each chart uploads -- the view, the picked trace, every
    trace's visibility and the canvas size are put back once, at the end."""
    fn = _fn('withPrintCanvas')
    body, fin = fn.split('finally', 1)
    assert 'return await fn(fit);' in body and 'gExportDpr = scale;' in body
    for restore in ('gView = keep.view;', 'gPickKey = keep.pick;',
                    'gTraces.forEach((t, i) => { t.visible = keep.vis[i]; });',
                    'gExportDpr = 0;', 'resizeCanvas();'):
        assert restore in fin, restore
    # each snapshot re-sizes the canvas for print (a resize in an await
    # would have put it back) and never carries the crosshair
    snap = _fn('snapChart')
    assert snap.index('fit();') < snap.index('draw();') and 'gMouse = null;' in snap
    assert 'withPrintCanvas(REPORT_FIBRE_W, REPORT_FIBRE_H, 1.5,' in _fn('reportFibreCharts')
    # plotRect and draw measure the canvas at the export scale while it runs
    assert 'canvas.width  / canvasDpr()' in _fn('plotRect')
    assert 'canvas.width / canvasDpr()' in _fn('draw')
    # declared with the other globals, long before the first draw can run
    assert SRC.index('let gExportDpr = 0;') < SRC.index("const canvas = document.getElementById('chart');")


def test_the_page_sends_the_event_table_first_and_no_chart_of_every_trace():
    """Robert, 2026-09-30: no page of every trace at once; the event panel
    first, then straight into the fibre pages.  The only chart the page draws
    is a fibre's own, and the zoomed-in "whole span too" box went with the
    chart it added to.  The trace list stays after the event table, fibre
    pages or not (Robert, 2026-10-01: "we can keep that table")."""
    pay = _fn('reportPayload')
    assert 'images' not in pay and 'reportChartPng' not in SRC and 'REPORT_CHART_W' not in SRC
    assert 'rpt-whole' not in SRC and 'reportWholeView' not in SRC and 'opt.whole' not in SRC
    ev, mk, tr = (pay.index('tables.push(ev ||'), pay.index('if (mk) tables.push(mk);'),
                  pay.index("    tables.push({\n      title: 'Traces',"))
    assert ev < mk < tr
    assert re.findall(r'withPrintCanvas\((\w+)', SRC) == ['w', 'REPORT_FIBRE_W']   # its definition, one caller
    dlg = _fn('showReportDialog')
    assert '<div class="sub">The event table as it is on screen, every row of it.</div>' in dlg
    assert 'After the event table, one page per fiber' in dlg


def test_each_fibre_is_drawn_alone_and_sent_ahead():
    fn = _fn('reportFibreCharts')
    assert 't.visible = vis[i] && t.fiber === f;' in fn          # that fibre's traces only
    assert 'const b = dataBounds();' in fn and 'snapChart(fit, b);' in fn   # fitted, whole span
    assert "canvas.toBlob(res, 'image/png')" in fn
    assert '/api/report_image?token=' in fn and "out.set(f, 'ref:' + name);" in fn


def test_the_fibre_table_is_the_combined_tables_cells_turned_on_their_side():
    """Nothing recomputed: each value cell is the combined table's own cell
    (text and style) for that fibre's row; a column the fibre has nothing in
    leaves; FastReporter-mode events are numbered 1..n per fibre."""
    fn = _fn('reportFibreTables')
    assert 'const v = { t: c.t, s: c.s };' in fn
    assert 'if (vals.every(v => reportBlank(v.t))) continue;' in fn
    assert "/^Event \\d+$/.test(g.title) ? `Event ${++n}` : g.title" in fn
    assert "k !== 'type'" in fn                                    # folded into one Type column
    assert "avgT ? avgT[1]" in fn                                  # the bidirectional event's type
    # the payload carries the pages after the combined table, and the
    # dialog offers them, on by default
    pay = _fn('reportPayload')
    assert 'reportFibreTables(ev, { H, B, C })' in pay and 'fibres, token,' in pay
    assert 'id="rpt-fibres" checked> Add Per Fiber Summary Pages</label>' in SRC   # Robert's wording
    # a Summary Report may hold one trace or many: no trace count in its words
    assert 'trace${vis.length' not in SRC
    # with a row or cell filter on, only the fibres the table still shows
    assert 'if (ev && (gFlaggedOnly || cellFilterOn())) list = list.filter(f => byTable.has(f));' in pay


def test_a_failing_column_average_is_red_in_the_pinned_strip():
    """The pinned Min / Max / Average strip's grey background outranked the red
    fill, so a failing column Average printed white on grey: invisible, on
    screen and in the report that copies the screen."""
    css = SRC.split('<style>', 1)[1].split('</style>', 1)[0]
    rule = re.search(r'([^{}]*)\{\s*background:\s*#e74c3c;\s*\}', css.split('tr.fr-pick td.fr-rowlab', 1)[1])
    assert rule and 'tfoot tr.fr-agg td.fr-hi' in rule.group(1) and 'tr.fr-pick td.fr-hi' in rule.group(1)
    assert re.search(r'tfoot tr\.fr-agg td\.fr-warn,\s*table\.fr-table tr\.fr-pick td\.fr-warn \{ background: #ffeb00; \}', css)
