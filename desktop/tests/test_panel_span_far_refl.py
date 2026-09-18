"""On a panel tie between reels the far-end reflectance is graded bare.

Redrock RDR4<->RDR5 (East), 2026-09-17, tech-confirmed OOS list: F1 .519
bidir at the RDR4 connector, and four reflectance fails, all on the B shot:
F67 -45.4 and F69 -49.7 at RDR5 (B's launch end), F74 -49.8 and F126 -47.0
at RDR4 (B's FAR end).  Our report had the first three and missed the last
two.

The launch rule grades bare against LAUNCH_BAD_REFL_DB.  The far-end
("tailbox") rule adds a population outlier test on top — a fiber must read
TAILBOX_OUTLIER_DB (7.5 dB) worse than its direction's median — which
exists for cables shot with no receive jumper, where every fiber shows the
same bare-glass end.  A panel tie has no bare glass: the far reading is the
other panel's connector, the same object the launch rule just graded.  On
East the B far median is -51.8 dB, so F74 sits +2.0 and F126 +4.8 over it,
and the bar hid both.  Below PANEL_SPAN_MAX_KM the far-end rule is bare.

Fixture: East fibers 2 (clean), 74 and 126, both directions.
"""
from __future__ import annotations

import re

import openpyxl

from conftest import run_splicereport, FIXTURE_DIR

EAST_A = FIXTURE_DIR / "panelreels_east" / "A"
EAST_B = FIXTURE_DIR / "panelreels_east" / "B"


def _ila(tmp_path):
    out = tmp_path / "panelreels_east.xlsx"
    rc, m, stderr = run_splicereport(EAST_A, EAST_B, out, "RDR4", "RDR5")
    assert rc == 0 and m and m.get("ok"), f"runner failed: {stderr[-1200:]}"
    assert 0 < float(m["span_km"]) < 0.1, m["span_km"]     # premise: panel tie
    ws = openpyxl.load_workbook(out)["Splice Report"]
    hdr = list(next(ws.iter_rows(min_row=3, max_row=3, values_only=True)))
    ila = {}
    for ci, h in enumerate(hdr, 1):
        if h and "ILA" in str(h):
            ila[str(h)] = " ".join(str(ws.cell(r, ci).value or "")
                                   for r in range(4, ws.max_row + 1))
    assert set(ila) == {"A-end ILA: RDR4", "B-end ILA: RDR5"}, hdr
    return ila


def test_far_end_reflectance_fails_are_reported_at_the_end_they_are_at(tmp_path):
    """THE regression.  F74 -49.8 and F126 -47.0 are the B shot's view of
    the RDR4 connector, so they belong in the RDR4 column."""
    ila = _ila(tmp_path)
    rdr4 = ila["A-end ILA: RDR4"]
    assert re.search(r"\b126 REFL-47\.0dB", rdr4), rdr4
    assert re.search(r"\b74 REFL-49\.8dB", rdr4), rdr4


def test_a_clean_fiber_and_the_other_end_stay_quiet(tmp_path):
    ila = _ila(tmp_path)
    assert not re.search(r"\b2 REFL", ila["A-end ILA: RDR4"]), ila
    assert "REFL" not in ila["B-end ILA: RDR5"], ila


def test_outlier_bar_still_applies_on_a_cable():
    """Scope guard: off a panel span the 7.5 dB population test is untouched.
    Eight 60 km fibers whose far 1F all read -49.0 dB (a bad-but-uniform
    receive end): nothing flags, exactly as before."""
    import subprocess
    import sys
    import textwrap
    from conftest import REPO_ROOT
    body = textwrap.dedent(f"""
        import sys
        sys.path.insert(0, {str(REPO_ROOT / 'splicereport')!r})
        import splicereportmatchexfo as E
        def rec(far_refl, span=60.0):
            return {{'events': [
                {{'dist_km': 0.0, 'splice_loss': 0.0, 'is_end': False, 'reflection': -55.0,
                  'is_reflective': True, 'time_of_travel': 0, 'type': '1F9999LS'}},
                {{'dist_km': 1.01, 'splice_loss': 0.1, 'is_end': False, 'reflection': -55.0,
                  'is_reflective': True, 'time_of_travel': 5000, 'type': '1F9999LS'}},
                {{'dist_km': span - 1.01, 'splice_loss': 0.1, 'is_end': False,
                  'reflection': far_refl, 'is_reflective': True, 'type': '1F9999LS'}},
                {{'dist_km': span, 'splice_loss': 0.0, 'is_end': True, 'reflection': -55.0,
                  'is_reflective': True, 'type': '1E9999LS'}},
            ]}}
        A = {{i: rec(-49.0) for i in range(1, 9)}}
        B = {{i: rec(-55.0) for i in range(1, 9)}}
        issues = E.detect_launch_issues(A, B)
        bad = [(f, i) for f, i in issues.items()
               if any('REFL' in t for t in (i.get('a_tags') or []) + (i.get('b_tags') or []))]
        assert not bad, bad
        print('OK')
    """)
    p = subprocess.run([sys.executable, "-c", body], capture_output=True, text=True)
    assert p.returncode == 0, f"{p.stdout}\n{p.stderr}"
    assert p.stdout.strip().splitlines()[-1] == "OK"
