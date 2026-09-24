"""A tie panel shot through sacrificial jumpers grades the panel, not the jumper.

FTH01<->FTH06 West Panel B (2026-03-07), fixture = ribbon 1, both directions:

    port -> 1.03 km launch reel -> 15 m jumper -> panel A -> 62 m tie
         -> panel B -> 15 m jumper -> 1.02 km receive reel

The tech set the span start on panel A and the end marker on panel B, so each
table opens with the reel-to-jumper joint at -0.015 km (the "first spike") and
carries the jumper-to-reel joint and the reel end past the end marker.

FastReporter 3.21, Lumen template (Reflectance Fail -50.0), read 2026-09-23:
Event 1 is the panel at 0.0000 km; the -0.0153 km spike (-49.9 / -50.0 dB)
and the 0.0773 km spike (-49.7 dB) are shown but not numbered and not red.
The report graded the first spike as each end's launch connector and printed
REFL-49.9dB on every fiber of this ribbon.
"""
from __future__ import annotations

import openpyxl

from conftest import run_splicereport, FIXTURE_DIR

JUMPER_DIR = FIXTURE_DIR / "paneljumper"


def _run(tmp_path):
    out = tmp_path / "paneljumper.xlsx"
    rc, m, stderr = run_splicereport(JUMPER_DIR / "A", JUMPER_DIR / "B", out,
                                     "FTH01", "FTH06")
    assert rc == 0 and m and m.get("ok"), f"runner failed: {stderr[-1200:]}"
    return m, out


def _grid(out):
    ws = openpyxl.load_workbook(out)["Splice Report"]
    return [str(ws.cell(r, c).value)
            for r in range(4, ws.max_row + 1) for c in range(2, ws.max_column + 1)
            if ws.cell(r, c).value is not None]


def test_the_jumper_joint_before_the_span_start_is_not_graded(tmp_path):
    _, out = _run(tmp_path)
    grid = _grid(out)
    assert not any("REFL" in t for t in grid), grid


def test_every_panel_connector_on_the_ribbon_passes(tmp_path):
    """Bidirectional panel losses on this ribbon are 0.053-0.352 dB and every
    panel reflectance is -51.1 dB or lower, so FR's verdict is a clean
    ribbon and the grid is blank."""
    m, out = _run(tmp_path)
    assert _grid(out) == [], _grid(out)
    assert not [c for c in m["cells"] if c["is_flagged"]]


def test_span_is_the_tie_between_the_panels(tmp_path):
    m, _ = _run(tmp_path)
    assert abs(m["span_km"] - 0.06) < 0.005, m["span_km"]
