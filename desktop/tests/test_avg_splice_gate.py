"""Per-fiber average splice loss — FastReporter's "Avg. Splice Loss", gated.

FR's definition, proven on the real WSC<->SUI exports (1152 fibers): for every
mid-span, non-reflective position EITHER direction recorded, score
(A->B + B->A) / 2; the fiber's average is the signed mean of those, taken on
the raw values and rounded once.  A position with only one reading is '---'
and is left out.  Reproduced by the engine to a median of 1 mdB, unbiased.

Three things have to hold:

  1. The arithmetic is FR's (union, signed, connectors out, round once).
  2. OFF is really off: the shipped report gains no sheet and no Legend row
     unless a profile or the panel sends a positive gate.
  3. The IIG profile turns it on at 0.08 dB and Default leaves it off, and
     the unticked value (0.0) survives run_splicereport's override guard.
"""
from __future__ import annotations

import importlib
import sys

import openpyxl
from conftest import (
    run_splicereport, FIXTURE_SPLICE_A_DIR, FIXTURE_SPLICE_B_DIR, SPLICEREPORT_DIR,
)

# The engine FIRST, with its own sor_reader copy, the way every other engine
# test does it.  The hub (`app`) is imported lazily inside the profile tests:
# importing it at module level would load the viewer's sor_reader324802a and
# poison the engine import for every test module collected after this one.
sys.path.insert(0, str(SPLICEREPORT_DIR))
import splicereportmatchexfo as E  # noqa: E402

SHEET = "Average splice loss"
IIG = "AWS / IIG MT.1085"


def _engine():
    return E


def _hub():
    return importlib.import_module('app')


def _rec(events, eof_km):
    """A minimal fiber record: stored events only, no trace (so a silent side
    cannot be measured and stays unpaired, exactly FR's '---')."""
    evs = [dict(dist_km=km, splice_loss=loss, type=typ, is_end=False)
           for km, loss, typ in events]
    evs.append(dict(dist_km=eof_km, splice_loss=0.0, type='2F9999', is_end=True))
    return {'events': evs, '_source': None}


# ── 1. The arithmetic is FR's ────────────────────────────────────────────
def test_union_signed_mean_rounded_once():
    E = _engine()
    # Three splices at 5, 10, 15 km on a 20 km fiber, both directions stored.
    # B positions are in B's own frame, so B's 15.0 is A's 5.0 and so on.
    # Per-splice averages: A5/B15 (0.05+0.07)/2=0.060, A10/B10 (0.10+0.00)/2
    # =0.050, A15/B5 (-0.03+0.01)/2=-0.010 (a gainer, kept SIGNED)
    # -> mean 0.033333...
    a = _rec([(5.0, 0.05, '0'), (10.0, 0.10, '0'), (15.0, -0.03, '0')], 20.0)
    b = _rec([(5.0, 0.01, '0'), (10.0, 0.00, '0'), (15.0, 0.07, '0')], 20.0)
    out = E.fiber_average_splice_loss({1: a}, {1: b})
    r = out[1]
    assert r['n_used'] == 3 and r['n_union'] == 3
    assert r['n_measured'] == 0 and r['n_unpaired'] == 0
    assert abs(r['avg'] - 0.1 / 3) < 1e-12
    # Rounded once, at the end: 0.033, not a mean of rounded terms.
    assert round(r['avg'], 3) == 0.033


def test_connectors_and_launch_are_left_out():
    E = _engine()
    # A reflective event mid-span (a connector) in A, and a launch-zone event.
    # B positions are in B's OWN frame (mirrored about B's 20 km end of
    # fiber): 15.0 in B is 5.0 in A, 8.0 in B is 12.0 in A.
    a = _rec([(0.005, 0.90, '0'), (5.0, 0.05, '0'), (12.0, 0.40, '1F9999')], 20.0)
    b = _rec([(15.0, 0.03, '0'), (8.0, 0.40, '1F9999')], 20.0)
    r = E.fiber_average_splice_loss({1: a}, {1: b})[1]
    assert r['n_union'] == 1                # the connector is FR's other column
    assert r['n_used'] == 1
    assert abs(r['avg'] - 0.04) < 1e-12


def test_one_sided_position_without_a_trace_is_unpaired():
    """B stored a splice A did not.  With no trace to measure A on, FR prints
    '---' and the position is left out of the mean — not scored as B alone."""
    E = _engine()
    a = _rec([(5.0, 0.05, '0')], 20.0)
    # B frame: 15.0 mirrors to A's 5.0 (paired); 10.0 mirrors to 10.0 (B-only).
    b = _rec([(15.0, 0.03, '0'), (10.0, 0.30, '0')], 20.0)
    r = E.fiber_average_splice_loss({1: a}, {1: b})[1]
    assert r['n_union'] == 2
    assert r['n_measured'] == 1            # the silent A side was attempted
    assert r['n_unpaired'] == 1            # ...and came back with no reading
    assert r['n_used'] == 1
    assert abs(r['avg'] - 0.04) < 1e-12


def test_verdict_is_on_the_rounded_value_and_0080_passes():
    E = _engine()
    old = E.AVG_SPLICE_LOSS_DB
    try:
        E.AVG_SPLICE_LOSS_DB = 0.080
        assert E.avg_splice_verdict(0.0800) == 'PASS'      # '<= 0.08' passes
        assert E.avg_splice_verdict(0.0804) == 'PASS'      # rounds to 0.080
        assert E.avg_splice_verdict(0.0805) == 'FAIL'      # rounds to 0.081
        assert E.avg_splice_verdict(-0.010) == 'PASS'
        assert E.avg_splice_verdict(None) is None
        E.AVG_SPLICE_LOSS_DB = 0.0
        assert E.avg_splice_verdict(0.5) is None           # gate off
    finally:
        E.AVG_SPLICE_LOSS_DB = old


# ── 2. OFF is really off ─────────────────────────────────────────────────
def _run(tmp_path, name, **kw):
    out = tmp_path / name
    rc, m, err = run_splicereport(FIXTURE_SPLICE_A_DIR, FIXTURE_SPLICE_B_DIR,
                                  out, **kw)
    assert rc == 0 and m and m.get("ok"), f"run failed: {err[-1200:]}"
    _run.manifest = m
    return openpyxl.load_workbook(out)


def _legend_mentions_average(wb):
    for ws in wb.worksheets:
        if ws.title == SHEET:
            continue
        for row in ws.iter_rows(values_only=True):
            if row and row[0] == "Average splice loss":
                return True
    return False


def test_default_run_has_no_average_sheet_and_no_legend_row(tmp_path):
    wb = _run(tmp_path, "off.xlsx")
    assert SHEET not in wb.sheetnames
    assert not _legend_mentions_average(wb)
    # ...and the manifest the hub reads is unchanged too.
    assert not any(k.startswith(("n_avg_splice", "avg_splice")) for k in _run.manifest)


def test_gate_on_adds_the_sheet_and_the_legend_row(tmp_path):
    wb = _run(tmp_path, "on.xlsx", overrides={"AVG_SPLICE_LOSS_DB": 0.08})
    assert SHEET in wb.sheetnames
    ws = wb[SHEET]
    hdr = [ws.cell(row=4, column=c).value for c in range(1, 7)]
    assert hdr == ["Fiber", "Splices averaged", "Measured one side",
                   "Left out (one reading)", "Avg. splice loss (dB)", "Verdict"]
    rows = [r[:6] for r in ws.iter_rows(min_row=5, values_only=True)
            if r[0] is not None and isinstance(r[0], int)]
    assert rows, "the fixture has fibers in both directions"
    for fib, n_used, n_meas, n_unp, avg, verdict in rows:
        if avg is None:
            assert verdict == "no paired splices" and n_used == 0
            continue
        assert n_used > 0
        assert verdict == ('FAIL' if round(avg, 3) > 0.08 else 'PASS')
    assert _legend_mentions_average(wb)
    # The manifest carries the gate and the counts, and the count agrees
    # with the sheet.  n_flagged is a CELL count and must not include them.
    m = _run.manifest
    assert m["avg_splice_gate_db"] == 0.08
    assert m["n_avg_splice_fibers"] == len(rows)
    assert m["n_avg_splice_fail"] == sum(1 for r in rows if r[5] == 'FAIL')


def test_unticked_zero_survives_the_override_guard(tmp_path):
    """The panel sends 0.0 for an unticked row; run_splicereport must apply
    it (it is finite and the key is not in the positive-only set), so OFF
    stays off rather than reverting to some baseline."""
    wb = _run(tmp_path, "zero.xlsx", overrides={"AVG_SPLICE_LOSS_DB": 0.0})
    assert SHEET not in wb.sheetnames


# ── 3. Profiles ──────────────────────────────────────────────────────────
def test_iig_average_splice_loss_gate():
    hub = _hub()
    ov = hub._overrides_from_settings(hub._otdr_settings_from_profile(IIG))
    assert ov["AVG_SPLICE_LOSS_DB"] == 0.080


def test_default_profile_leaves_the_gate_off():
    hub = _hub()
    for prof in hub.CUSTOMER_PROFILES:
        if prof == IIG:
            continue
        ov = hub._overrides_from_settings(hub._otdr_settings_from_profile(prof))
        assert ov.get("AVG_SPLICE_LOSS_DB") == 0.0, prof


def test_row_is_wired_and_supported():
    hub = _hub()
    row = next(r for r in hub.OTDR_ROWS if r[0] == "avg_splice_loss")
    assert row[4] is True
    assert hub._OTDR_KEY_TO_ENGINE_GLOBAL["avg_splice_loss"] == "AVG_SPLICE_LOSS_DB"
    assert hub._OTDR_KEY_DISABLE_VALUE["avg_splice_loss"] == 0.0
    assert "avg_splice_loss" not in hub.OTDR_DEFAULT_APPLY
