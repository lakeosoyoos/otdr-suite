"""
EXFO .olts in the Viewer (the boss, 2026-10-01: "we need to be able to
handle these in viewer ... pdf or excel so we can sort").

An .olts stores raw lock-in readings, not the loss / ORL / length EXFO
prints, so the reader computes them.  These tests pin:

  1. The reader's numbers equal EXFO's own library's on the scrubbed
     fixture, at full precision (fixtures/olts, see its README): loss both
     ways and their average exactly, ORL and length to rounding noise.
  2. The header facts EXFO prints: job, locations, reference, thresholds.
  3. Anything that is not a readable .olts is refused with a reason.
  4. trace_server writes the report: an Excel workbook whose Results sheet
     is the table alone with a filter row (so it sorts), and a PDF.
  5. The page sends a dropped .olts to /api/olts_load and never into the
     A / B drop, still loads the traces dropped with it, and lists the
     fibers in the Measurements tab, where a heading click sorts them.
"""
import importlib.util
import json
import os
import re
import subprocess
import urllib.error
import urllib.request

import pytest

from conftest import VIEWER_DIR, import_trace_server

HERE = os.path.dirname(os.path.abspath(__file__))
OLTS = os.path.join(HERE, 'fixtures', 'olts', 'OLTSFX_SAMPLE.olts')
EXFO = json.load(open(os.path.join(HERE, 'fixtures', 'olts', 'OLTSFX_SAMPLE_exfo.json'),
                     encoding='utf-8'))
VIEWER_HTML = os.path.join(str(VIEWER_DIR), 'viewer.html')
JSC = ('/System/Library/Frameworks/JavaScriptCore.framework/'
       'Versions/Current/Helpers/jsc')
needs_jsc = pytest.mark.skipif(not os.path.exists(JSC),
                               reason='JavaScriptCore shell (macOS) not present')


def _viewer_reader():
    spec = importlib.util.spec_from_file_location(
        'viewer_sor_reader_olts', str(VIEWER_DIR / 'sor_reader324802a.py'))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


V = _viewer_reader()


@pytest.fixture(scope='module')
def parsed():
    return V.parse_olts(OLTS)


# ── 1. EXFO's numbers ────────────────────────────────────────────────

def test_every_number_is_exfos(parsed):
    got = {f['id']: f for f in parsed['fibers']}
    assert sorted(got) == sorted(EXFO['fibers'])
    for fid, want in EXFO['fibers'].items():
        row, = got[fid]['rows']
        assert row['wl_nm'] == 1550
        # Loss and its (linear) average: the same arithmetic as EXFO's.
        for k in ('loss_ab', 'loss_ba', 'loss_avg'):
            assert row[k] == pytest.approx(want[k], abs=1e-12), (fid, k)
        for k in ('orl_a', 'orl_b'):
            assert row[k] == pytest.approx(want[k], abs=1e-9), (fid, k)
        assert got[fid]['length_m'] == pytest.approx(want['length_m'], abs=1e-6), fid


def test_the_average_is_linear_not_the_mean_of_the_db(parsed):
    """EXFO averages the two transmittances.  The mean of the two dB values
    differs past the second decimal on a tie (10 of 864 fibers of the job)."""
    for f in parsed['fibers']:
        r = f['rows'][0]
        assert r['loss_avg'] < (r['loss_ab'] + r['loss_ba']) / 2


def test_fiber_721_after_bs_phase_step_still_measures_the_span(parsed):
    """Unit B's phase stepped between fibers 720 and 721: one direction's
    'length' moves by kilometres, but the A+B delay the length is made of
    does not, and EXFO's own lengths stay within 0.4 km of each other."""
    lengths = [f['length_m'] for f in parsed['fibers']]
    assert max(lengths) - min(lengths) < 400


# ── 2. What EXFO prints around the table ─────────────────────────────

def test_header_facts(parsed):
    assert parsed['file'] == 'OLTSFX_SAMPLE.olts'
    assert (parsed['job'], parsed['customer'], parsed['company']) == \
        ('OLTS SAMPLE JOB 0001', 'CUSTA', 'CMP')
    a, b = parsed['units']['A'], parsed['units']['B']
    assert (a['operator'], a['model'], a['serial'], a['calibration']) == \
        ('OA', 'FTB-945-SM3-EA', '1852702', '2024-12-18')
    assert (b['operator'], b['serial'], b['calibration']) == ('OB', '1874614', '2024-12-28')
    assert parsed['wavelengths'] == [1550]


def test_reference_is_exfos_loopback(parsed):
    ref, = parsed['references']
    assert ref['method'] == 'Loopback'
    row, = ref['rows']
    assert row['ref_ab'] == pytest.approx(EXFO['reference']['ref_ab'], abs=1e-15)
    assert row['ref_ba'] == pytest.approx(EXFO['reference']['ref_ba'], abs=1e-15)
    # each unit's reference power, FastReporter's Ref. A->B / Ref. B->A
    assert row['power_ab'] == EXFO['reference']['power_ab']
    assert row['power_ba'] == EXFO['reference']['power_ba']
    assert ref['when'].isoformat().startswith('2026-09-28T16:22:54')


def test_link_orl_threshold(parsed):
    assert V.olts_orl_min(parsed, 1550) == 30.0
    assert V.olts_orl_min(parsed, 1310) == 30.0           # the set's 0 = every wavelength
    assert {(t['set'], t['kind'], t['fail']) for t in parsed['thresholds']} == \
        {('$ct', 3, 30.0), ('$mt', 3, 30.0)}


# ── 3. Refusals ──────────────────────────────────────────────────────

def test_not_an_olts_is_refused(tmp_path):
    p = tmp_path / 'bad.olts'
    p.write_bytes(b'not a compound file' * 100)
    with pytest.raises(ValueError, match='not an EXFO .olts'):
        V.parse_olts(str(p))


def test_a_truncated_olts_is_refused(tmp_path):
    data = open(OLTS, 'rb').read()
    p = tmp_path / 'cut.olts'
    p.write_bytes(data[:len(data) // 3])
    with pytest.raises(ValueError):
        V.parse_olts(str(p))


def test_the_compound_file_reader_reads_small_streams_too():
    """The previews sit in the mini stream; the measurements in full sectors."""
    streams = V._cfb_streams(open(OLTS, 'rb').read())
    assert sorted(streams) == sorted(['OltsMeasures/%d' % i for i in range(5)]
                                     + ['OltsMeasuresPreview/%d' % i for i in range(5)])
    prev = V._Nrbf(V._gunzip_capped(streams['OltsMeasuresPreview/0'])).parse()
    assert prev['__class'] == 'Metrino.Oltsx.OltsMeasurement'


# ── 4. The report ────────────────────────────────────────────────────

@pytest.fixture(scope='module')
def TS():
    return import_trace_server()


@pytest.fixture
def loaded(TS):
    return TS.olts_load('OLTSFX_SAMPLE.olts', open(OLTS, 'rb').read())


def test_load_answers_what_the_dialog_shows(loaded):
    assert loaded['file'] == 'OLTSFX_SAMPLE.olts'
    assert loaded['fibers'] == 5 and loaded['wavelengths'] == [1550]
    assert loaded['status'] == 'Pass' and loaded['orl_fails'] == 0
    assert loaded['orl_min'] == {'1550': 30.0}


def test_load_answers_the_measurements_rows(loaded):
    """A row per fiber for the Measurements tab, rounded as it prints, so a
    sort orders what the tech reads."""
    rows = {r['id']: r for r in loaded['rows']}
    assert sorted(rows) == sorted(EXFO['fibers'])
    r, want = rows['OLTSFX001'], EXFO['fibers']['OLTSFX001']
    assert (r['type'], r['dir'], r['pf']) == ('OLTS', 'Bidir', 'pass')
    w = r['wl']['1550']
    assert (w['loss_ab'], w['loss_ba'], w['loss_avg'], w['orl_a'], w['orl_b']) == (
        round(want['loss_ab'], 2), round(want['loss_ba'], 2), round(want['loss_avg'], 2),
        round(want['orl_a'], 2), round(want['orl_b'], 2))
    assert w['orl_a_fail'] is False and w['orl_b_fail'] is False
    # FR's FasTesT table prints 0.00 (never -0.00) and 1.62
    assert (w['ref_ab'], w['ref_ba']) == (0.0, 1.62) and str(w['ref_ab']) == '0.0'
    assert r['length_km'] == round(want['length_m'] / 1000, 3)
    assert r['when'].startswith('9/28/2026 ') and r['when_ts'] > 0


def test_the_workbook_sorts(TS, loaded, tmp_path):
    from openpyxl import load_workbook
    out = TS.write_olts_report({'token': loaded['token'], 'format': 'xlsx',
                                'dest': str(tmp_path), 'name': 'OLTS sample'})
    assert out['path'] == str(tmp_path / 'OLTS sample.xlsx') and os.path.isfile(out['path'])
    wb = load_workbook(out['path'])
    assert wb.sheetnames == ['Results', 'Job']
    ws = wb['Results']
    head = [c.value for c in ws[1]]
    assert head == ['Identifier', 'Wavelength (nm)', 'Loss Average (dB)', 'Loss Margin (dB)',
                    'Loss A->B (dB)', 'Loss B->A (dB)', 'ORL A (dB)', 'ORL B (dB)',
                    'Ref. A->B (dBm)', 'Ref. B->A (dBm)', 'Length (km)', 'Date/Time']
    assert ws.auto_filter.ref == 'A1:L6'
    assert ws.freeze_panes == 'B2'
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    want = EXFO['fibers']['OLTSFX001']
    r = rows['OLTSFX001']
    # Numbers, rounded as EXFO prints them, so a sort is numeric.
    assert r[1:11] == (1550, round(want['loss_avg'], 2), '---', round(want['loss_ab'], 2),
                       round(want['loss_ba'], 2), round(want['orl_a'], 2),
                       round(want['orl_b'], 2), 0.0, 1.62, round(want['length_m'] / 1000, 3))
    assert r[11].year == 2026
    assert ws['G2'].font.color.rgb.endswith('008000')        # ORL passes: green, as EXFO
    job = wb['Job']
    vals = [c for row in job.iter_rows(values_only=True) for c in row if c is not None]
    for v in ('OLTS SAMPLE JOB 0001', 'CUSTA', 'Loopback', 0.56, 0.04, 30.0, 'Pass'):
        assert v in vals, v


def test_the_pdf_is_written(TS, loaded, tmp_path):
    out = TS.write_olts_report({'token': loaded['token'], 'format': 'pdf',
                                'dest': str(tmp_path), 'name': ''})
    assert os.path.basename(out['path']) == 'OLTSFX_SAMPLE.pdf'
    data = open(out['path'], 'rb').read()
    assert data.startswith(b'%PDF') and len(re.findall(rb'/Type /Page\b', data)) == 1
    # The written file is one /api/report_open may open.
    assert out['path'] in TS._REPORTS_WRITTEN


def test_a_report_needs_a_loaded_file(TS, tmp_path):
    with pytest.raises(ValueError, match='no longer loaded'):
        TS.write_olts_report({'token': 'nope', 'format': 'xlsx', 'dest': str(tmp_path)})


def test_the_endpoints(TS, tmp_path):
    port = TS.start_in_thread(8795)

    def post(path, body, ctype, origin=None):
        h = {'Content-Type': ctype}
        if origin:
            h['Origin'] = origin
        req = urllib.request.Request(f'http://127.0.0.1:{port}{path}', data=body, headers=h)
        with urllib.request.urlopen(req, timeout=30) as r:
            return json.loads(r.read())

    j = post('/api/olts_load?name=OLTSFX_SAMPLE.olts', open(OLTS, 'rb').read(),
             'application/octet-stream')
    assert j['ok'] and j['fibers'] == 5
    k = post('/api/olts_report', json.dumps({'token': j['token'], 'format': 'xlsx',
                                              'dest': str(tmp_path)}).encode(),
             'application/json')
    assert k['ok'] and os.path.isfile(k['path'])
    with pytest.raises(urllib.error.HTTPError) as e:
        post('/api/olts_load?name=x.olts', b'junk', 'application/octet-stream')
    assert e.value.code == 400
    with pytest.raises(urllib.error.HTTPError) as e:
        post('/api/olts_load?name=x.olts', b'junk', 'application/octet-stream',
             origin='https://evil.example')
    assert e.value.code == 403


# ── 5. The page ──────────────────────────────────────────────────────

SRC = open(VIEWER_HTML, encoding='utf-8').read()


def _js_func(name):
    m = re.search(r'(?:async\s+)?function\s+' + re.escape(name) + r'\s*\([^)]*\)\s*\{', SRC)
    assert m, 'viewer.html no longer defines %s()' % name
    i, depth = m.end(), 1
    while i < len(SRC) and depth:
        depth += {'{': 1, '}': -1}.get(SRC[i], 0)
        i += 1
    return SRC[m.start():i]


def test_the_page_takes_olts_in_a_drop():
    exts = re.search(r'const DROP_EXTS = (\[[^\]]*\]);', SRC).group(1)
    assert "'.olts'" in exts
    assert 'drop .sor / .json / .trc / .olts files, a folder, or a .zip here' in SRC
    # FastReporter's Measurements tab, at the foot of the panel on the right
    assert '<button data-view="meas" hidden' in SRC


def test_a_right_click_on_a_heading_offers_both_sorts():
    """The boss, 2026-10-02: right-click any heading for sort high to low or
    low to high (words: A to Z, Z to A)."""
    menu = _js_func('showOltsSortMenu')
    for words in ("'Sort High to Low'", "'Sort Low to High'", "'Sort A to Z'", "'Sort Z to A'"):
        assert words in menu, words
    assert "gOltsSort = { col, desc: b.dataset.desc === '1' };" in menu
    hook = SRC.split("filesList.addEventListener('contextmenu', (ev) => {\n  if (gFilesView !== 'meas'", 1)[1]
    assert "showOltsSortMenu(ev.clientX, ev.clientY, th.dataset.col, th.dataset.kind," in hook.split('\n});', 1)[0]


_STUBS = r"""
const DROP_EXTS = ['.sor', '.json', '.trc', '.zip', '.olts'];
const DROP_BATCH_FILES = 32, DROP_BATCH_BYTES = 4 * 1024 * 1024;
var gDropFolder = new WeakMap(), gDropInFlight = false, gAutoFit = true;
var gRemovedFiles = new Set(), gTraces = [], gInfo = null, gLoadFailures = [];
var urls = [], dialogs = [], readout = null, signs = [], picked = null, rendered = 0;
var gOlts = null, gOltsSort = null, gOltsSel = null, gFilesView = 'files';
var window = globalThis;
function renderFilesPanel() { rendered++; }
var el = { textContent: '' };
var document = { getElementById: function () { return el; } };
function _packDropBatch(files) { return { names: files.map(function (f) { return f.name; }) }; }
function emptiedSides() { return ''; }
function forgetSide(d) {}
function renderChips() {} function fit() {} function draw() {} function renderEventTable() {}
function loadFailNote() { return ''; }
function reportDropFailure(e) { print('DROPFAIL ' + e); }
function setReadout(s, x) { readout = s; }
function showDropSign(state, small) { signs.push(state); }
function hideDropSign() { signs.push('hide'); }
async function selectFiles(want) { picked = Array.from(want).sort(); }
async function loadInfo() { gInfo = { dir_a: '/t/A', dir_b: '', fibers_a: [1], fibers_b: [] }; return true; }
function showOltsDialog(info) { dialogs.push(info); return Promise.resolve(); }
function fetch(url, opt) {
  urls.push([url, opt && opt.body && opt.body.names ? opt.body.names : (opt && opt.body && opt.body.name) || '']);
  var body = { ok: true };
  if (url.indexOf('/api/drop_begin') === 0) body.token = 't1';
  if (url.indexOf('/api/drop_end') === 0) body = { ok: true, dir_a: '/t/A', added: 'A', a_count: 1, new_keys: ['a-1'] };
  if (url.indexOf('/api/olts_load') === 0) body = { ok: true, token: 'o1', file: 'J.olts', fibers: 864 };
  return Promise.resolve({ ok: true, json: function () { return Promise.resolve(body); } });
}
"""

_CASES = r"""
(async function () {
  var out = {};
  await handleFilesDrop([{ name: 'J.olts', size: 9 }]);
  out.only = { urls: urls.slice(), dialogs: dialogs.length, signs: signs.slice(),
               view: gFilesView, listed: gOlts && gOlts.file, rendered: rendered };
  urls = []; dialogs = []; signs = []; gFilesView = 'files'; gOlts = null;
  await handleFilesDrop([{ name: 'J.olts', size: 9 }, { name: 'A1.sor', size: 1 }]);
  out.mixed = { urls: urls.slice(), dialogs: dialogs.length, picked: picked, listed: gOlts && gOlts.file };
  // The Measurements tab's sort: a heading click, then the same again.
  var info = { wavelengths: [1550], rows: [
    { id: 'F2', pf: 'pass', wl: { 1550: { loss_ab: 18.1 } }, when_ts: 2 },
    { id: 'F10', pf: 'fail', wl: { 1550: { loss_ab: 20.5 } }, when_ts: 3 },
    { id: 'F1', pf: 'pass', wl: { 1550: { loss_ab: null } }, when_ts: 1 },
    { id: 'F3', pf: 'pass', wl: { 1550: { loss_ab: 17.6 } }, when_ts: 4 }] };
  var cols = oltsCols(info);
  var ids = function (s) { return oltsSortRows(info.rows, cols, s).map(function (r) { return r.id; }); };
  out.sort = {
    high: ids({ col: 'loss_ab@1550', desc: true }), low: ids({ col: 'loss_ab@1550', desc: false }),
    az: ids({ col: 'id', desc: false }), pf: ids({ col: 'pf', desc: true }),
    none: ids(null), heads: cols.map(function (c) { return c.t; }) };
  print('OUT ' + JSON.stringify(out));
})().catch(function (e) { print('ERR ' + e + '\n' + e.stack); });
"""


@pytest.fixture(scope='module')
def page(tmp_path_factory):
    funcs = re.search(r'^const SRC_RE = [^\n]+;', SRC, re.M).group(0) + '\n' + '\n'.join(
        _js_func(n) for n in ('srcParts', 'sideOf', 'srcLabel', 'emptiedList',
                                            '_parentName', '_dropOk', '_dropBatches',
                                            'handleFilesDrop', 'isOltsFile', 'handleOltsDrop',
                                            'showOltsMeasurements', 'oltsCols', 'oltsValue',
                                            'oltsSortRows'))
    path = tmp_path_factory.mktemp('olts_drop') / 'drop.js'
    path.write_text(_STUBS + funcs + '\n' + _CASES, encoding='utf-8')
    r = subprocess.run([JSC, str(path)], capture_output=True, text=True, timeout=60)
    out = r.stdout + r.stderr
    assert r.returncode == 0 and 'OUT ' in out, out[-2000:]
    return json.loads(out.split('OUT ', 1)[1].strip())


@needs_jsc
def test_an_olts_alone_never_touches_the_a_b_drop(page):
    urls = [u for u, _ in page['only']['urls']]
    assert urls == ['/api/olts_load?name=J.olts']
    assert page['only']['signs'][-1] == 'hide'


@needs_jsc
def test_a_dropped_olts_opens_the_measurements_tab(page):
    """The boss, 2026-10-02: the values on screen, in the panel on the right.
    The report dialog waits for Save Report…."""
    assert page['only']['view'] == 'meas' and page['only']['listed'] == 'J.olts'
    assert page['only']['rendered'] >= 1
    assert page['only']['dialogs'] == 0


@needs_jsc
def test_a_heading_sorts_highest_first_then_lowest(page):
    s = page['sort']
    assert s['heads'] == ['P/F', 'Identifiers', 'Type', 'Direction', '1550 A->B Loss',
                          '1550 B->A Loss', '1550 Average Loss', '1550 ORL A', '1550 ORL B',
                          '1550 Ref. A->B', '1550 Ref. B->A', 'Length', 'Test Date/Time']
    assert s['high'] == ['F10', 'F2', 'F3', 'F1']           # a blank stays last
    assert s['low'] == ['F3', 'F2', 'F10', 'F1']
    assert s['az'] == ['F1', 'F2', 'F3', 'F10']             # F10 after F3, not after F1
    assert s['pf'] == ['F10', 'F2', 'F1', 'F3']             # the failures first, ties in file order
    assert s['none'] == ['F2', 'F10', 'F1', 'F3']           # unsorted: as the file has them


@needs_jsc
def test_traces_dropped_with_an_olts_still_load(page):
    urls = page['mixed']['urls']
    assert urls[0][0] == '/api/olts_load?name=J.olts'
    sent = [n for u, names in urls if u.startswith('/api/drop_file') for n in names]
    assert 'J.olts' not in sent and 'A1.sor' in sent
    assert page['mixed']['listed'] == 'J.olts'
    assert page['mixed']['picked'] == ['a-1']


def test_a_threshold_it_cannot_read_leaves_the_verdict_blank(TS, parsed):
    """A loss limit (any kind but Link ORL) is not read yet: no Pass for it."""
    other = dict(parsed, thresholds=parsed['thresholds'] + [
        {'set': '$ct', 'kind': 1, 'fiber_types': 16, 'wl_nm': 0, 'fail': 20.0, 'enabled': True}])
    assert TS._olts_status(parsed) == 'Pass'
    assert TS._olts_status(other) == ''
