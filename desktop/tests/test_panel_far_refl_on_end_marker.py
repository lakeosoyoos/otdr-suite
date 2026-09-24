"""A tie panel whose end marker sits ON the far panel grades that panel's reflectance.

LSC1<->LSC6 (Las Cruces, 2026-04-14), fixture = ribbon 11 (fibers 121-132),
both directions.  The tech set the span start on the near panel (0.0000 km)
and the end marker on the far panel (0.0311 km); nothing lies between them.

LSC6LSC10128 reads the LSC1 panel at -49.78 dB on its end marker.  FastReporter
(Lumen template, Reflectance Fail -50.0) shows -49.8 and fails it: Event 2,
with "Include span end values" on.  The report looked for the last 1F before
the end marker for the far reading, found the NEAR panel at 0.0000 km
(-52.0 dB) and passed the fiber.  Red Rock East F74 is the same finding read
from the B shot on a span with launch reels in the table (-49.8 dB, tech-
confirmed out of spec 2026-09-17), which the report already caught.

A fiber that breaks short of the far panel keeps its old reading: its end
marker is the break, not a connector (panelbreak fixture: RDR4<->RDR6 fibers
70, 73 and 74 break at a panel and must print no REFL).
"""
from __future__ import annotations

import openpyxl

from conftest import run_splicereport, FIXTURE_DIR


def _grid(tmp_path, name, site_a, site_b):
    out = tmp_path / f"{name}.xlsx"
    d = FIXTURE_DIR / name
    rc, m, stderr = run_splicereport(d / "A", d / "B", out, site_a, site_b)
    assert rc == 0 and m and m.get("ok"), f"runner failed: {stderr[-1200:]}"
    ws = openpyxl.load_workbook(out)["Splice Report"]
    return [str(ws.cell(r, c).value)
            for r in range(4, ws.max_row + 1) for c in range(2, ws.max_column + 1)
            if ws.cell(r, c).value is not None]


def test_far_panel_on_the_end_marker_is_graded(tmp_path):
    grid = _grid(tmp_path, "panelfarrefl", "LSC1", "LSC6")
    refl = [t for t in grid if "REFL" in t]
    assert refl == ["128 REFL-49.8dB"], grid


def test_a_break_is_not_read_as_the_far_panel(tmp_path):
    grid = _grid(tmp_path, "panelbreak", "RDR4", "RDR6")
    assert not any("REFL" in t for t in grid), grid
