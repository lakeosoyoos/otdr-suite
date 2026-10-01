"""A display-only reading never takes the report down.

The Closures and Port length tabs are display only: nothing on them changes
the duplicate likelihood.  So an error in either (a folder layout nobody
tested, or a bare % in the sheet text, which on 2026-09-27 stopped the whole
workbook being written) must cost only that tab, never the report.  Each
piece is guarded on its own: the reading, its summary for the manifest and
its sheet.

What has to hold:
  1. A reading that fails leaves no keys on the pairs and abstains out loud.
  2. A sheet that fails leaves a one-line note in its place.
  3. Through the runner, a failing reading leaves every other sheet, every
     pair's score and verdict, and the manifest exactly as without it.
  4. A failing manifest summary leaves the report and the sheet in place.

Namespace isolation rule: the engine and the runner are exercised only
through subprocesses.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys

import pytest

from conftest import FIXTURE_DIR, SECRETSAUCE_DIR, run_secretsauce

PANEL = FIXTURE_DIR / "paneljumper"

# The runner with the named engine functions replaced by one that raises.
_BROKEN_RUNNER = r"""
import runpy, sys
ss, folder, out, fmt = sys.argv[1:5]
sys.path.insert(0, ss)
import report_sor

def _raiser(name):
    def boom(*a, **k):
        raise RuntimeError('injected failure')
    boom.__name__ = name
    return boom

for name in sys.argv[5:]:
    setattr(report_sor, name, _raiser(name))
sys.argv = ['run_secretsauce.py', '--folder', folder, '--out-dir', out, '--format', fmt]
runpy.run_path(ss + '/run_secretsauce.py', run_name='__main__')
"""


def _run_broken(folder, out, fmt, *names):
    p = subprocess.run([sys.executable, "-c", _BROKEN_RUNNER, str(SECRETSAUCE_DIR),
                        str(folder), str(out), fmt, *names],
                       capture_output=True, text=True, timeout=600)
    manifest = None
    for line in reversed((p.stdout or "").strip().splitlines()):
        line = line.strip()
        if line.startswith("{") and line.endswith("}"):
            manifest = json.loads(line)
            break
    return p.returncode, manifest, p.stderr


def _engine(script: str):
    p = subprocess.run([sys.executable, "-c", script, str(SECRETSAUCE_DIR)],
                       capture_output=True, text=True, timeout=300)
    assert p.returncode == 0, p.stderr[-3000:]
    return json.loads(p.stdout.strip().splitlines()[-1])


def _content(path):
    """Every sheet's values, without the two display-only tabs, their Summary
    rows and the time the report was generated."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    out = {}
    for ws in wb.worksheets:
        if ws.title in ("Closures", "Port length"):
            continue
        rows = []
        for row in ws.iter_rows(values_only=True):
            row = list(row)
            while row and row[-1] is None:
                row.pop()
            if ws.title == "Summary" and row and (
                    row[0] in ("Closures", "Port length") or str(row[0]).startswith("Generated")):
                continue
            rows.append(tuple(row))
        out[ws.title] = rows
    return out


def _verdicts(pairs):
    return {(q["fileA"], q["fileB"]): (q["score"], q["p_dup"], q["verdict"])
            for q in pairs["pairs"]}


@pytest.fixture(scope="module")
def panel(tmp_path_factory):
    """The tie-panel fixture both ways in one folder, run as it ships."""
    folder = tmp_path_factory.mktemp("guard_panel")
    for side in ("A", "B"):
        for p in sorted((PANEL / side).glob("*.sor")):
            shutil.copy(p, folder / p.name)
    rc, xlsx, err = run_secretsauce(folder, tmp_path_factory.mktemp("guard_xlsx"), "xlsx")
    assert rc == 0 and xlsx and xlsx.get("ok") is True, err[-2000:]
    rc, pairs, err = run_secretsauce(folder, tmp_path_factory.mktemp("guard_pairs"), "pairs")
    assert rc == 0 and pairs and pairs.get("ok") is True, err[-2000:]
    # the readings this file breaks are really there when nothing is broken
    assert xlsx.get("port_length") and pairs.get("port_length")
    assert any(k.startswith("port_len") for q in pairs["pairs"] for k in q)
    return folder, xlsx, pairs


def test_a_failing_reading_leaves_no_keys_on_the_pairs():
    r = _engine(r"""
import sys, json
sys.path.insert(0, sys.argv[1])
import report_sor as R
pairs = [{'a': 'x', 'b': 'y', 'p_dup': 0.25}, {'a': 'x', 'b': 'z', 'p_dup': 0.5}]
def half_done(files, pairs):
    pairs[0]['port_len_diff_m'] = 0.01
    pairs[0]['port_len_left_in'] = True
    raise ValueError('half done')
out = R._display_only('Port length', 'port_len_', pairs, half_done, [], pairs)
print(json.dumps({'out': out, 'pairs': pairs}))
""")
    assert r["out"] == {"usable": False, "note": "Port length: NOT USABLE - error: half done"}
    assert r["pairs"] == [{"a": "x", "b": "y", "p_dup": 0.25}, {"a": "x", "b": "z", "p_dup": 0.5}]


def test_a_failing_sheet_leaves_a_note_in_its_place():
    r = _engine(r"""
import sys, json
sys.path.insert(0, sys.argv[1])
from openpyxl import Workbook
import report_sor as R
wb = Workbook()
wb.active.title = 'Summary'
def bad(ws, reading):
    ws.cell(row=1, column=1, value='half written')
    ws.cell(row=2, column=1, value=reading['summary'])
    raise ValueError('bad cell')
def good(ws, reading):
    ws.cell(row=1, column=1, value=reading['summary'])
R._display_only_sheet(wb, 'Closures', bad, {'summary': 'x'})
R._display_only_sheet(wb, 'Port length', good, {'summary': 'fine'})
print(json.dumps({'sheets': wb.sheetnames,
                  'closures': [list(r) for r in wb['Closures'].iter_rows(values_only=True)],
                  'port': [list(r) for r in wb['Port length'].iter_rows(values_only=True)]}))
""")
    assert r["sheets"] == ["Summary", "Closures", "Port length"]
    assert len(r["closures"]) == 1 and len(r["closures"][0]) == 1, r["closures"]
    assert "could not be written (error: bad cell)" in r["closures"][0][0]
    assert r["port"] == [["fine"]]


def test_a_failing_reading_costs_only_its_tab(panel, tmp_path):
    folder, xlsx, _ = panel
    rc, broken, err = _run_broken(folder, tmp_path, "xlsx", "_port_length", "_closure_fingerprint")
    assert rc == 0 and broken and broken.get("ok") is True, err[-2000:]
    assert "Port length: NOT USABLE - error: injected failure" in err
    assert "Closure fingerprint: NOT USABLE - error: injected failure" in err
    from openpyxl import load_workbook
    normal_path = xlsx["written"][0]["path"]
    broken_path = broken["written"][0]["path"]
    assert "Port length" in load_workbook(normal_path).sheetnames
    assert "Port length" not in load_workbook(broken_path).sheetnames
    assert _content(broken_path) == _content(normal_path)
    # the manifest is the one without the reading
    assert "port_length" not in broken and "closure_fp" not in broken
    same = {k: v for k, v in xlsx.items() if k not in ("port_length", "written", "folder")}
    assert {k: v for k, v in broken.items() if k not in ("written", "folder")} == same


def test_pairs_mode_keeps_every_verdict_when_a_reading_fails(panel, tmp_path):
    folder, _, pairs = panel
    rc, broken, err = _run_broken(folder, tmp_path, "pairs", "_port_length")
    assert rc == 0 and broken and broken.get("ok") is True, err[-2000:]
    assert "port_length" not in broken
    assert not [k for q in broken["pairs"] for k in q if k.startswith("port_len")]
    assert _verdicts(broken) == _verdicts(pairs)


def test_a_failing_manifest_summary_keeps_the_report(panel, tmp_path):
    folder, _, _ = panel
    rc, broken, err = _run_broken(folder, tmp_path / "x", "xlsx",
                                  "_port_length_meta", "_closure_meta")
    assert rc == 0 and broken and broken.get("ok") is True, err[-2000:]
    assert "_port_length_meta left out of the manifest - error: injected failure" in err
    assert "port_length" not in broken
    from openpyxl import load_workbook
    assert "Port length" in load_workbook(broken["written"][0]["path"]).sheetnames
    rc, broken, err = _run_broken(folder, tmp_path / "p", "pairs",
                                  "_port_length_meta", "_closure_meta")
    assert rc == 0 and broken and broken.get("ok") is True, err[-2000:]
    assert "port_length" not in broken
    # the reading itself still ran: its keys are on the pairs
    assert any(q.get("port_len_left_in") for q in broken["pairs"])
