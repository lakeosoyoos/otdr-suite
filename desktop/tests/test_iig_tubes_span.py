"""Two more lines from the AWS / IIG MT.1085 spec sheet, section 1.

  * The cable is 18 buffer tubes of 24 fibers.  The grid groups fibers by
    RIBBON_SIZE, which defaults to 12, so an IIG report showed 36 half-tubes
    with tube letters that do not match the cable.  The profile now sets 24
    through the whitelisted engine extras, and the runner's override guard
    (RIBBON_SIZE is a positive-int global) must let 24 through.
  * Spans on this job run 64.8 to 72.6 km.  The contract block gains a span
    length range and the acquisition audit checks every trace's stored span
    length against it -- a trace short of the range ended early or is the
    wrong span, one beyond it is the wrong span or a stretched group index.
    Reported, never corrected, like every other contract row.
"""
from __future__ import annotations

import json
import sys

from conftest import (
    run_splicereport, FIXTURE_SPLICE_A_DIR, FIXTURE_SPLICE_B_DIR, SPLICEREPORT_DIR,
)

sys.path.insert(0, str(SPLICEREPORT_DIR))
import acquisition_audit as AUD  # noqa: E402

import app as hub  # noqa: E402

IIG = "AWS / IIG MT.1085"


# ── Tubes of 24 ──────────────────────────────────────────────────────────
def test_iig_groups_the_grid_by_24():
    ex = hub._engine_extras_from_profile(IIG)
    assert ex["RIBBON_SIZE"] == 24
    assert "RIBBON_SIZE" in hub._PROFILE_ENGINE_KEYS


def test_other_profiles_keep_the_engine_ribbon_size():
    for prof in hub.CUSTOMER_PROFILES:
        if prof != IIG:
            assert "RIBBON_SIZE" not in hub._engine_extras_from_profile(prof), prof


def test_ribbon_size_24_survives_the_runner_and_reaches_the_manifest(tmp_path):
    """The extras travel as floats (24.0); the runner's int path must accept
    that and the manifest must say 24, because the hub draws the grid from
    the manifest's ribbon_size."""
    out = tmp_path / "r24.xlsx"
    rc, m, err = run_splicereport(FIXTURE_SPLICE_A_DIR, FIXTURE_SPLICE_B_DIR, out,
                                  overrides={"RIBBON_SIZE": 24.0})
    assert rc == 0 and m and m.get("ok"), err[-1200:]
    assert m["ribbon_size"] == 24
    out2 = tmp_path / "r12.xlsx"
    rc, m2, err = run_splicereport(FIXTURE_SPLICE_A_DIR, FIXTURE_SPLICE_B_DIR, out2)
    assert rc == 0 and m2 and m2.get("ok"), err[-1200:]
    assert m2["ribbon_size"] == 12


# ── Span length range ────────────────────────────────────────────────────
def _rec(span_m=None, end_km=None):
    r = {'events': []}
    if span_m is not None:
        r['exfo_spans_length'] = span_m
    if end_km is not None:
        r['events'] = [dict(dist_km=end_km, is_end=True)]
    return r


def _span_row(recs_a, recs_b, rng=(64.8, 72.6)):
    res = AUD.compute_contract_conformance(recs_a, recs_b,
                                           {"name": "t", "span_km_range": list(rng)})
    assert res is not None
    rows = [r for r in res["rows"] if r["name"] == "Span length"]
    assert len(rows) == 1
    return rows[0], res


def test_spans_inside_the_range_conform():
    row, res = _span_row([_rec(64850.0), _rec(72550.0)], [_rec(65000.0)])
    assert row["conforms"] is True
    assert row["expected"] == "64.8 to 72.6 km"
    assert "64.850 to 72.550 km over 3 trace(s)" == row["actual"]
    assert res["clean"] is True


def test_a_span_outside_the_range_is_a_finding_with_the_count():
    row, res = _span_row([_rec(64850.0), _rec(58200.0)], [_rec(74100.0)])
    assert row["conforms"] is False
    assert "2 of 3 trace(s)" in row["note"]
    assert "58.200 km" in row["note"]          # the farthest from the range
    assert "Reported, not corrected" in row["note"]
    assert res["clean"] is False


def test_the_range_carries_half_a_kilometre_of_tolerance():
    """The contract quotes 64.8 to 72.6; Span 29 measures 72.60-72.66 km on
    over half its traces and NCT's own median is 72.604.  Those conform.  A
    trace a full kilometre beyond the range does not."""
    row, _ = _span_row([_rec(72655.0), _rec(72604.0), _rec(64350.0)], [])
    assert row["conforms"] is True
    row, _ = _span_row([_rec(73650.0)], [])
    assert row["conforms"] is False


def test_end_event_is_the_fallback_when_no_span_is_stored():
    row, _ = _span_row([_rec(end_km=70.0)], [])
    assert row["conforms"] is True and "70.000 km" in row["actual"]


def test_no_span_anywhere_is_informational():
    row, _ = _span_row([_rec()], [_rec()])
    assert row["conforms"] is None and row["actual"] == "not stored"


def test_no_range_means_no_row():
    res = AUD.compute_contract_conformance([_rec(65000.0)], [], {"name": "t", "ior": 1.467})
    assert res is not None
    assert not any(r["name"] == "Span length" for r in res["rows"])


def test_runner_whitelists_the_range(tmp_path):
    """Two positive finite km figures, low then high; anything else is dropped
    and the run is unaffected."""
    import subprocess
    runner = SPLICEREPORT_DIR / "run_splicereport.py"

    def run(contract, name):
        out = tmp_path / name
        cmd = [sys.executable, str(runner), "--dir-a", str(FIXTURE_SPLICE_A_DIR),
               "--dir-b", str(FIXTURE_SPLICE_B_DIR), "--out", str(out),
               "--site-a", "A", "--site-b", "B", "--contract", json.dumps(contract)]
        p = subprocess.run(cmd, capture_output=True, text=True)
        assert p.returncode == 0, p.stderr[-1200:]
        import openpyxl
        ws = openpyxl.load_workbook(out)[openpyxl.load_workbook(out).sheetnames[0]]
        return "\n".join(" | ".join(str(c.value) for c in row if c.value is not None)
                         for row in ws.iter_rows())

    good = run({"name": "t", "span_km_range": [0.5, 200.0]}, "good.xlsx")
    assert "Span length" in good and "0.5 to 200.0 km" in good
    for bad in ([72.6, 64.8], [64.8], ["a", "b"], "64.8-72.6", [0, 70], [float("nan"), 70]):
        txt = run({"name": "t", "span_km_range": bad, "ior": 1.467}, "bad.xlsx")
        assert "Span length" not in txt, bad
