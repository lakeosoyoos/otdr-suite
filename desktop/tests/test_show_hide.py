"""Show / hide in report: the engine finds the same things, the switches
only decide which flagged cells reach the workbook, and a Display sheet
says what was left out.  Runs the real runner on the splice fixtures."""
import json
import subprocess
import sys

import openpyxl

from conftest import REPO_ROOT

RUNNER = REPO_ROOT / 'splicereport' / 'run_splicereport.py'
FIX = REPO_ROOT / 'desktop' / 'tests' / 'fixtures'


def _run(tmp_path, name, *args):
    out = tmp_path / f'{name}.xlsx'
    p = subprocess.run([sys.executable, str(RUNNER), *args, '--out', str(out)],
                       capture_output=True, text=True, cwd=str(RUNNER.parent))
    return json.loads(p.stdout.strip().splitlines()[-1]), out


def _bidir(tmp_path, name, show=None):
    args = ['--dir-a', str(FIX / 'splice_A'), '--dir-b', str(FIX / 'splice_B'),
            '--site-a', 'A', '--site-b', 'B']
    if show:
        args += ['--show', json.dumps(show)]
    return _run(tmp_path, name, *args)


def test_bidir_loss_hidden_keeps_bends(tmp_path):
    full, full_x = _bidir(tmp_path, 'full')
    cats = [c['category'] for c in full['cells']]
    assert 'reburn' in cats and 'bend' in cats
    assert 'Display' not in openpyxl.load_workbook(full_x).sheetnames
    hid, hid_x = _bidir(tmp_path, 'hid', {'loss': False})
    assert [c for c in hid['cells'] if c['category'] == 'reburn'] == []
    assert ([c for c in hid['cells'] if c['category'] == 'bend']
            == [c for c in full['cells'] if c['category'] == 'bend'])
    rows = [[c.value for c in r]
            for r in openpyxl.load_workbook(hid_x)['Display'].iter_rows(max_row=4)]
    assert rows == [['Category', 'Shown'], ['Splice loss', 'N'],
                    ['Bend/Damage', 'Y'], ['Breaks', 'Y']]


def test_uni_bend_hidden_drops_column_keeps_splices(tmp_path):
    base = ['--uni', '--dir-a', str(FIX / 'splice_A')]
    full, _ = _run(tmp_path, 'u_full', *base)
    hid, _ = _run(tmp_path, 'u_hid', *base, '--show', json.dumps({'bend': False}))
    kinds = lambda d: [c['kind'] for c in d['uni']['grid_columns']]
    assert 'bend_damage' in kinds(full) and 'bend_damage' not in kinds(hid)
    splice = lambda d: [(c['fiber'], c['loss']) for c in d['uni']['cells']
                        if c['kind'] == 'splice']
    assert splice(hid) == splice(full)
