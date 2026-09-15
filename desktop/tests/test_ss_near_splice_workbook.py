"""The two new sheets appear only where they apply, and render when they do.

'Near splice' needs a shared splice behind the panel; the repo's 5 km / 10 ns
WSC->SUI short shots carry one about 37 m past the port, so they are the
real-data case.  'Shot out of order' needs a fibre skipped and shot later; no
fixture has one, so the renderer is driven with a planted run.  Every folder
with neither keeps its exact sheet list and Summary rows.

Namespace isolation rule: the engine is only exercised through subprocesses.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys

from conftest import FIXTURE_A_DIR, FIXTURE_B_DIR, REPO_ROOT, SECRETSAUCE_DIR

CONTINUOUS = REPO_ROOT / "desktop" / "tests" / "fixtures" / "continuous"
NEAR_HDR = ["File", "Meter", "Shot at", "Splice loss (dB)", "Next fibre", "Time gap (s)",
            "Difference (dB)", "Difference (sd)",
            "Shots of one fibre that differ this much (%)"]

_SCRIPT = r"""
import sys, json
sys.path.insert(0, sys.argv[1])
import openpyxl
import report_sor as RS
folder, out_xlsx, plant = sys.argv[2], sys.argv[3], sys.argv[4] == '1'
if plant:
    def _planted(files):
        names = sorted(f['name'] for f in files)
        t = 1700000000
        return [{'names': [names[2]], 'first': 3, 'last': 3, 'shot_at': t + 7200,
                 'before': names[1], 'after': names[3], 'before_at': t,
                 'after_at': t + 25, 'minutes_later': 119.6}]
    RS._fill_ins = _planted
RS.build_xlsx_sor(folder, 'T', out_xlsx)
wb = openpyxl.load_workbook(out_xlsx)
summary = {str(r[0]): r[1] for r in wb['Summary'].iter_rows(values_only=True)
           if r and r[0] is not None}
out = {'sheets': wb.sheetnames, 'summary_near': summary.get('Near splice'),
       'summary_order': summary.get('Shot out of order')}
if 'Near splice' in wb.sheetnames:
    vals = list(wb['Near splice'].iter_rows(values_only=True))
    out['ns_first'] = vals[0][0]
    h = [i for i, r in enumerate(vals) if r and r[0] == 'File']
    out['ns_header'] = list(vals[h[0]][:9]) if h else None
    body = [r for r in vals[h[0] + 1:] if r and r[0]] if h else []
    out['ns_rows'] = len(body)
    out['ns_losses'] = sum(1 for r in body if isinstance(r[3], (int, float)))
if 'Shot out of order' in wb.sheetnames:
    out['order_row'] = list(next(wb['Shot out of order'].iter_rows(
        min_row=2, max_row=2, values_only=True)))
print(json.dumps(out, default=str))
"""


def _build(tmp_path, sources, plant=False):
    d = tmp_path / "span"
    d.mkdir()
    for src in sources:
        shutil.copy(src, d / src.name)
    p = subprocess.run([sys.executable, "-c", _SCRIPT, str(SECRETSAUCE_DIR), str(d),
                        str(tmp_path / "out.xlsx"), "1" if plant else "0"],
                       capture_output=True, text=True, timeout=600)
    assert p.returncode == 0, p.stderr[-3000:]
    return json.loads(p.stdout.strip().splitlines()[-1])


def test_near_splice_sheet_on_a_span_that_has_the_splice(tmp_path):
    r = _build(tmp_path, sorted(CONTINUOUS.glob("*.sor")))
    assert "Near splice" in r["sheets"]
    assert r["sheets"][-1] == "Near splice", "new sheets go last, indices stay stable"
    assert r["summary_near"].startswith("Splice ")
    assert r["ns_first"] == r["summary_near"]
    assert r["ns_header"] == NEAR_HDR
    assert r["ns_rows"] == 6
    assert r["ns_losses"] >= 5
    assert "Shot out of order" not in r["sheets"] and r["summary_order"] is None


def test_shot_out_of_order_sheet_renders(tmp_path):
    r = _build(tmp_path, sorted(CONTINUOUS.glob("*.sor")), plant=True)
    assert "Shot out of order" in r["sheets"]
    assert r["summary_order"].startswith("1 fibre(s) shot after both neighbouring fibres")
    row = r["order_row"]
    assert row[0] == "WSC_SUIsh_0018" and row[2] == "WSC_SUIsh_0017" and row[4] == "WSC_SUIsh_0019"
    assert row[6] == 119.6


def test_a_folder_with_neither_keeps_its_layout(tmp_path):
    r = _build(tmp_path, sorted(FIXTURE_A_DIR.glob("*.sor")) + sorted(FIXTURE_B_DIR.glob("*.sor")))
    assert "Near splice" not in r["sheets"]
    assert "Shot out of order" not in r["sheets"]
    assert r["summary_near"] is None and r["summary_order"] is None
