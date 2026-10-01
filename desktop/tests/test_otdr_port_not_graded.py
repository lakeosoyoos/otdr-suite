"""A shot with no launch reel: the OTDR's own port is judged against its peers.

The launch rule grades the first reflective event of each direction's table as
that end's connector.  Pass 0 strips the OTDR port whenever a launch reel
follows it, so on a reel shot the first event IS the panel connector.  With no
reel and no declared span start the port is still event 1, and the rule graded
the instrument.  A 66 km, 576-fibre cable shot straight from its panels
(2026-09-15) failed on every fibre at both ends: -37.7 to -38.9 dB where one
OTDR shot, about -45 dB where the other shot, at either end of the cable.
1152 end-column tags.  The tech's sheet from the same traces has none.

EXFO marks that row in the file: status bit 0x08, the launch LEVEL, in the
proprietary event list.  FastReporter prints the row as "Launch Level" and,
with a template that includes the span start (Reflectance Fail -50), paints
every one of those 1152 readings red.  A port far worse than the others is
still worth FR's red: one fibre at -24.7 dB against a -54 dB level is a bad
mate at the panel or on the test cord.  So the port is judged the way the far
end of a shot with no receive jumper already is: it must fail the gate AND
read PORT_OUTLIER_DB worse than the ports the same OTDR read in that
direction.  A profile that declares the 0 km event to be the panel connector
(PANEL_CONN_DIRECT) grades it bare, as before.

Fixtures, identifiers scrubbed:
  portnoreel   fibres 1, 2 and 500 of that 66 km cable, both directions.
  portoutlier  fibres 2-5 of a 48 km no-reel span, both directions; fibre 3's
               port reads -24.7 dB at end A, the rest -54 to -61 dB.

Engine tests run in a clean subprocess (3-engine sor_reader isolation).
"""
from __future__ import annotations

import re
import subprocess
import sys
import textwrap

import openpyxl

from conftest import run_splicereport, FIXTURE_DIR, REPO_ROOT

FX = FIXTURE_DIR / "portnoreel"
OUT_FX = FIXTURE_DIR / "portoutlier"
REEL_FX = FIXTURE_DIR / "launch_noreceive"
SPLICEREPORT_DIR = REPO_ROOT / "splicereport"


def _engine(body):
    header = ("import sys, os, glob\n"
              f"sys.path.insert(0, {str(SPLICEREPORT_DIR)!r})\n"
              "import sor_reader324802a as sr\n"
              "import splicereportmatchexfo as E\n"
              f"FX = {str(FX)!r}\n"
              f"REEL_FX = {str(REEL_FX)!r}\n")
    p = subprocess.run([sys.executable, "-c", header + textwrap.dedent(body)],
                       capture_output=True, text=True)
    assert p.returncode == 0, f"exit {p.returncode}\n{p.stdout}\n{p.stderr}"
    assert p.stdout.strip().splitlines()[-1] == "OK", p.stdout


def _report(tmp_path, fx, overrides=None):
    out = tmp_path / f"{fx.name}.xlsx"
    rc, m, stderr = run_splicereport(fx / "A", fx / "B", out, "SITEA", "SITEB",
                                     overrides)
    assert rc == 0 and m and m.get("ok"), f"runner failed: {stderr[-1200:]}"
    ws = openpyxl.load_workbook(out)["Splice Report"]
    hdr = list(next(ws.iter_rows(min_row=3, max_row=3, values_only=True)))
    ila = {}
    for ci, h in enumerate(hdr, 1):
        if h and "ILA" in str(h):
            ila[str(h).split(":")[0]] = " ".join(
                str(ws.cell(r, ci).value or "") for r in range(4, ws.max_row + 1)).strip()
    assert set(ila) == {"A-End ILA", "B-End ILA"}, hdr
    return m, ila


def test_premise_event_one_is_the_port_and_fails_the_gate():
    """What the old rule graded: event 1 of every file is EXFO's launch-level
    row at 0 km, with no reel after it and no span start, and it fails the
    -50.0 gate.  A record without the proprietary list is never matched."""
    _engine("""
        files = sorted(glob.glob(os.path.join(FX, '*', '*.sor')))
        assert len(files) == 6, files
        for f in files:
            r = sr.parse_sor_full(f, trim=False)
            e0, e1 = r['events'][0], r['events'][1]
            assert e0['dist_km'] == 0.0 and not r.get('user_offset_km'), f
            assert not (e1.get('is_reflective') and e1['dist_km'] < 3.0), f
            x0 = next(x for x in r['exfo_events'] if not x.get('_is_section'))
            assert x0['Status'] & 0x08 and x0['Position'] == 0.0, (f, x0)
            assert x0['Reflectance'] == e0['reflection'], f
            assert E.refl_fails(e0['reflection'], -50.0), (f, e0['reflection'])
            assert E._is_otdr_port(r, e0), f
            assert not E._is_otdr_port(r, e1), f
            bare = {k: v for k, v in r.items() if k != 'exfo_events'}
            assert not E._is_otdr_port(bare, e0), f
        print('OK')
    """)


def test_a_port_level_the_whole_instrument_shares_is_not_graded(tmp_path):
    """THE regression: every port fails the gate, none stands out from its
    own OTDR's readings, so no REFL at either end and none for the Viewer."""
    m, ila = _report(tmp_path, FX)
    assert m.get("end_refl") == [], m.get("end_refl")
    assert "REFL" not in ila["A-End ILA"] + ila["B-End ILA"], ila


def test_a_port_far_worse_than_its_peers_is_still_graded(tmp_path):
    """Fibre 3's -24.7 dB against a -54 dB level is kept, as FR keeps it;
    its neighbours and the whole B end pass."""
    m, ila = _report(tmp_path, OUT_FX)
    ports = [(v["fiber"], v["dir"], v["refl"]) for v in m["end_refl"]
             if (v["fiber"], v["dir"]) == (3, "A")]
    assert ports == [(3, "A", -24.7)], m["end_refl"]
    assert re.search(r"\b3 A→B REFL-24\.7dB", ila["A-End ILA"]), ila
    assert not re.search(r"\b[245] (?:A→B |B→A )?REFL", ila["A-End ILA"]), ila


def test_a_direct_panel_profile_still_grades_the_0km_event(tmp_path):
    """PANEL_CONN_DIRECT declares the 0 km event to be the panel connector
    (crews shooting straight from the panel), so that profile keeps grading
    it bare, exactly as before."""
    m, ila = _report(tmp_path, FX, {"PANEL_CONN_DIRECT": 1})
    got = {(v["fiber"], v["dir"]): v["refl"] for v in m["end_refl"]}
    assert got == {(1, "A"): -37.7, (2, "A"): -38.2, (500, "A"): -45.3,
                   (1, "B"): -45.0, (2, "B"): -44.6, (500, "B"): -38.5}, got
    assert re.search(r"\b1 A→B REFL-37\.7dB", ila["A-End ILA"]), ila
    assert re.search(r"\b500 B→A REFL-38\.5dB", ila["B-End ILA"]), ila


def test_a_launch_reel_connector_is_not_the_port():
    """On a reel shot Pass 0 moves the reel's connector onto 0 km.  It is the
    cable's first connector, not the port, and keeps its bare grade."""
    _engine("""
        for f in ('MONGRA0229_1550.sor', 'GRAMON1029_1550.sor'):
            r = sr.parse_sor_full(os.path.join(REEL_FX, f), trim=False)
            raw = r['events']
            assert E._is_otdr_port(r, raw[0]) and not E._is_otdr_port(r, raw[1]), f
            r['_raw_events'] = raw
            r['_trace_offset_km'] = E._trace_frame_offset_km(r, raw)
            assert 0.9 < r['_trace_offset_km'] < 1.1, (f, r['_trace_offset_km'])
            r['events'] = E._normalize_untrimmed_events(list(raw))
            le, _, _ = E._fiber_launch_info(r)
            assert le is not None and le['dist_km'] == 0.0, (f, le)
            assert le['reflection'] == raw[1]['reflection'], f
            assert not E._is_otdr_port(r, le), f
        print('OK')
    """)
