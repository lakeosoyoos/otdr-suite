"""An event column says what FastReporter calls it: Reflective,
Non-reflective, or Mixed (Robert 2026-10-01).

Under 80 fibres a job lists events, not splices or bends, labelled the way
FR's Event Table labels them: "Event 3 (22/24)" / "Non-reflective" /
"7.9295 km".  Per fibre an event is reflective when either direction's stored
event there is (SR-4731 type code 1 or 2, FR's bidirectional merge); per
column every fibre reflective is "Reflective", none is "Non-reflective", and
a disagreement is "Mixed", whose hover names each fibre's type.  The word is
stored once on the column ('event_kind', 'event_kind_tip'), and the Splice
Report sheet, the Uni sheet and the Viewer's OTDR Suite table all read it.

Engine tests run in a clean subprocess (3-engine sor_reader isolation).
"""
import json
import subprocess
import sys
import textwrap

import openpyxl

from conftest import FIXTURE_DIR, REPO_ROOT

SPLICEREPORT_DIR = REPO_ROOT / "splicereport"
RUNNER = SPLICEREPORT_DIR / "run_splicereport.py"
VIEWER = (REPO_ROOT / "viewer" / "viewer.html").read_text(encoding="utf-8")


def _engine(body, head=""):
    src = ("import sys\n"
           f"sys.path.insert(0, {str(SPLICEREPORT_DIR)!r})\n"
           "import splicereportmatchexfo as E\n" + head + textwrap.dedent(body))
    p = subprocess.run([sys.executable, "-c", src], capture_output=True, text=True)
    assert p.returncode == 0, f"{p.stdout}\n{p.stderr}"
    assert p.stdout.strip().splitlines()[-1] == "OK", p.stdout


def _runner(tmp_path, *args):
    tmp_path.mkdir(parents=True, exist_ok=True)
    out = tmp_path / "r.xlsx"
    p = subprocess.run([sys.executable, str(RUNNER), *map(str, args), "--out", str(out)],
                       capture_output=True, text=True)
    assert p.returncode == 0, p.stderr[-2000:]
    return json.loads(p.stdout.strip().splitlines()[-1]), out


_BODY = """
def ev(km, typ='0F9999LS', end=False):
    return {'dist_km': km, 'splice_loss': 0.05, 'type': typ, 'is_end': end,
            'is_reflective': typ[:1] in ('1', '2') and not end}
SPAN = 40.0
def fibre(evs):
    return {'events': [ev(0.0, '1F9999LS')] + evs + [ev(SPAN, '1E9999LS', True)]}
"""


def test_the_word_for_a_column():
    _engine("""
    assert E.event_kind({1: True, 2: True}) == ('Reflective', '')
    assert E.event_kind({1: False, 2: False}) == ('Non-reflective', '')
    assert E.event_kind({}) == ('Non-reflective', '')
    kind, tip = E.event_kind({2: True, 1: False})
    assert kind == 'Mixed' and tip == 'F1 Non-reflective, F2 Reflective', tip
    print('OK')
    """)


def test_either_direction_reflective_makes_the_fibre_reflective():
    """Fibre 1: A non-reflective, B reflective at 10 km -> reflective (FR's
    merge).  Fibre 2: non-reflective both ways.  So the 10 km column is
    Mixed and names both; the 25 km column (no reflective event) is
    Non-reflective, the 30 km one (all reflective) Reflective."""
    _engine(head=_BODY, body="""
    fa = {1: fibre([ev(10.0), ev(25.0), ev(30.0, '1F9999LS')]),
          2: fibre([ev(10.0), ev(25.0), ev(30.0, '2F9999LS')])}
    fb = {1: fibre([ev(SPAN - 10.0, '1F9999LS')]), 2: fibre([ev(SPAN - 10.0)])}
    cols = {round(c['position_km']): c for c in E.discover_event_columns(fa, fb)}
    assert cols[10]['event_kind'] == 'Mixed'
    assert cols[10]['event_kind_tip'] == 'F1 Reflective, F2 Non-reflective'
    assert cols[25]['event_kind'] == 'Non-reflective' and cols[25]['event_kind_tip'] == ''
    assert cols[30]['event_kind'] == 'Reflective'
    print('OK')
    """)


def test_a_column_holds_one_reading_per_fibre_per_direction():
    """Scattered small events chain at the closure gap across a 79-fibre
    job; a chain that would hold one fibre twice splits at its widest gap."""
    _engine(head=_BODY, body="""
    # fibre 1 at 10.00 and 10.40; fibres 2-3 bridge the gap at 10.20 / 10.22
    fa = {1: fibre([ev(10.00), ev(10.40)]), 2: fibre([ev(10.20)]),
          3: fibre([ev(10.22)])}
    cols = E.discover_event_columns(fa, None)
    assert len(cols) == 2, [c['position_km'] for c in cols]
    # the uni twin does the same
    with __import__('contextlib').redirect_stdout(__import__('io').StringIO()):
        u = E.uni_event_columns(fa)
    assert len(u) == 2 and all('event_kind' in c for c in u), u
    print('OK')
    """)


def test_splice_report_sheet_and_viewer_table_carry_the_word(tmp_path):
    """24 fibres: an event job.  The sheet's header reads "Event N" over the
    word, and the Viewer's table carries the same word on each event column
    (the end columns keep their own titles, with their own word)."""
    m, xlsx = _runner(tmp_path, "--dir-a", FIXTURE_DIR / "splice_A",
                      "--dir-b", FIXTURE_DIR / "splice_B",
                      "--viewer-table", tmp_path / "t.json")
    assert m["event_job"]
    kinds = [c["event_kind"] for c in m["columns"]]
    assert kinds and set(kinds) <= {"Reflective", "Non-reflective", "Mixed"}
    ws = openpyxl.load_workbook(xlsx)["Splice Report"]
    heads = [ws.cell(3, c).value for c in range(3, 3 + len(kinds))]
    assert heads == [f"Event {i + 1}\n{k}" for i, k in enumerate(kinds)], heads
    t = json.loads((tmp_path / "t.json").read_text(encoding="utf-8"))
    ev = [c for c in t["columns"] if c["title"].startswith("Event ")]
    assert [c["event_kind"] for c in ev] == kinds
    # the ends carry it too (Robert 2026-10-01), from the cable-end events
    assert t["columns"][0]["event_kind"] and t["columns"][-1]["event_kind"]


def test_a_mixed_header_has_a_hover_note(tmp_path):
    _engine(head=_BODY, body=f"""
    import openpyxl
    sp = [{{'position_km': 10.0, 'position_km_refined': 10.0, 'column_kind': 'splice',
           'is_event_column': True, 'splice_display_num': 1,
           'event_kind': 'Mixed', 'event_kind_tip': 'F1 Reflective, F2 Non-reflective'}}]
    cells, lca, lcb = E.build_ribbon_data({{}}, 2, 12, 1)
    E.write_xlsx(cells, sp, 2, 12, {str(tmp_path / 'm.xlsx')!r}, 'A', 'B', SPAN)
    c = openpyxl.load_workbook({str(tmp_path / 'm.xlsx')!r})['Splice Report'].cell(3, 3)
    assert c.value == 'Event 1\\nMixed', c.value
    assert c.comment is not None and 'F1 Reflective' in c.comment.text
    print('OK')
    """)


def test_uni_sheet_and_viewer_table_carry_the_word(tmp_path):
    m, xlsx = _runner(tmp_path, "--uni", "--dir-a", FIXTURE_DIR / "splice_A",
                      "--viewer-table", tmp_path / "u.json")
    ev = [c for c in m["uni"]["grid_columns"] if c["label"].startswith("Event ")]
    assert ev and all(c["event_kind"] in ("Reflective", "Non-reflective", "Mixed")
                      for c in ev)
    wb = openpyxl.load_workbook(xlsx)
    vals = [c.value for ws in wb.worksheets for row in ws.iter_rows(max_row=12)
            for c in row if isinstance(c.value, str) and c.value.startswith("Event ")]
    assert f"{ev[0]['label']}\n{ev[0]['event_kind']}" in vals, vals[:5]
    t = json.loads((tmp_path / "u.json").read_text(encoding="utf-8"))
    assert all(c.get("event_kind") for c in t["columns"] if c["title"].startswith("Event "))


def test_viewer_prints_the_word_before_the_distance():
    """The OTDR Suite table's second header row, as FR mode's: the word, then
    the distance; a Mixed word carries its hover."""
    i = VIEWER.index("function paintSuiteBidiGrid(")
    body = VIEWER[i:VIEWER.index("\n}\n", i)]
    assert "eventKind: c.event_kind ? String(c.event_kind) : ''," in body
    assert "eventTip: c.event_kind_tip ? String(c.event_kind_tip) : ''," in body
    assert ("? `<span class=\"fr-evkind\"${c.eventTip ? ` title=\"${esc(c.eventTip)}\"` : ''}>"
            "${esc(c.eventKind)}</span><br>`") in body
    assert ('h2 += `<th colspan="2" class="fr-evsub">${kindTxt}'
            '<span class="fr-km">${kmFt(c.km)}</span></th>`;') in body


def test_the_ends_read_fr_s_word_on_an_event_job(tmp_path):
    """Robert 2026-10-01: A-End / B-End read "Reflective" etc. like FR, on
    the sheet and in the Viewer table (an event job only); a reflective end
    is either direction's stored event there of type 1 or 2."""
    _engine("""
    def ev(km, typ, end=False):
        return {'dist_km': km, 'splice_loss': 0.0, 'type': typ, 'is_end': end}
    fa = {1: {'events': [ev(0.0, '1F9999LS'), ev(40.0, '0E9999LS', True)]},
          2: {'events': [ev(0.0, '0F9999LS'), ev(40.0, '1E9999LS', True)]}}
    fb = {1: {'events': [ev(0.0, '0F9999LS'), ev(40.0, '0E9999LS', True)]},
          2: {'events': [ev(0.0, '0F9999LS'), ev(40.0, '0E9999LS', True)]}}
    k = E.end_event_kinds(fa, fb)
    # A end: F1 A start reflective; F2 A start 0F, B end 0E -> Mixed
    assert k['A'] == ('Mixed', 'F1 Reflective, F2 Non-reflective'), k
    # B end: F1 A end 0E, B start 0F -> no; F2 A end 1E -> yes
    assert k['B'][0] == 'Mixed', k
    print('OK')
    """)
    m, xlsx = _runner(tmp_path, "--dir-a", FIXTURE_DIR / "splice_A",
                      "--dir-b", FIXTURE_DIR / "splice_B",
                      "--viewer-table", tmp_path / "t.json")
    ws = openpyxl.load_workbook(xlsx)["Splice Report"]
    t = json.loads((tmp_path / "t.json").read_text(encoding="utf-8"))
    a_kind, b_kind = t["columns"][0]["event_kind"], t["columns"][-1]["event_kind"]
    assert ws.cell(3, 2).value == f"A-End ILA: A\n{a_kind}"
    assert ws.cell(3, ws.max_column).value == f"B-End ILA: B\n{b_kind}"
    # a closure layout keeps its end headers as they were
    m2, xlsx2 = _runner(tmp_path / "c", "--dir-a", FIXTURE_DIR / "splice_A",
                        "--dir-b", FIXTURE_DIR / "splice_B",
                        "--overrides", json.dumps({"EVENT_JOB_MAX_FIBERS": 0}))
    assert openpyxl.load_workbook(xlsx2)["Splice Report"].cell(3, 2).value == "A-End ILA: A"


def test_uni_end_and_connector_columns_read_fr_s_word(tmp_path):
    m, _ = _runner(tmp_path, "--uni", "--dir-a", FIXTURE_DIR / "splice_A")
    cols = m["uni"]["grid_columns"]
    ends = [c for c in cols if c["kind"] in ("connector", "end")]
    assert ends and all(c.get("event_kind") in ("Reflective", "Non-reflective", "Mixed")
                        for c in ends), ends
