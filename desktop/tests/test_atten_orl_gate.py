"""Per-fiber span attenuation and ORL gates, from EXFO's stored span figures.

Every EXFO .sor carries the instrument's own span loss, span length and total
ORL in its proprietary block.  The span loss is the number FastReporter prints
as "Span Loss (dB)" (checked on the real WSC<->SUI exports: rounds to FR's
value on every one of 1152 fibers), so attenuation = span loss / span length
per direction, averaged, is graded on FR's own inputs.  ORL is the OTDR's
figure per direction and is graded as a FLOOR; it is not the OLTS ORL a
contract names, and the sheet says so.

Three things have to hold:

  1. The arithmetic: loss over length per direction, mean of the two; a
     missing figure is None, never zero; the ORL verdict fails on EITHER
     direction below the floor.
  2. OFF is really off: no sheet, no Legend rows, unchanged manifest unless
     a positive gate arrives.
  3. The IIG profile turns both on (0.250 dB/km, 30 dB); Default leaves both
     off; the unticked 0.0 survives the runner's guard.
"""
from __future__ import annotations

import importlib
import sys

import openpyxl
from conftest import (
    run_splicereport, FIXTURE_SPLICE_A_DIR, FIXTURE_SPLICE_B_DIR, SPLICEREPORT_DIR,
)

# Engine first (own sor_reader), hub lazily -- see test_avg_splice_gate.py.
sys.path.insert(0, str(SPLICEREPORT_DIR))
import splicereportmatchexfo as E  # noqa: E402

SHEET = "Span attenuation and ORL"
IIG = "AWS / IIG MT.1085"


def _hub():
    return importlib.import_module('app')


def _rec(loss=None, length_m=None, orl=None):
    r = {'events': [], '_source': 'sor'}
    if loss is not None:
        r['exfo_spans_loss'] = loss
    if length_m is not None:
        r['exfo_spans_length'] = length_m
    if orl is not None:
        r['exfo_total_orl'] = orl
    return r


# ── 1. Arithmetic ────────────────────────────────────────────────────────
def test_attenuation_is_loss_over_length_per_direction_then_averaged():
    a = _rec(loss=12.361, length_m=64042.4, orl=32.585)
    b = _rec(loss=12.352, length_m=64039.8, orl=32.954)
    s = E.fiber_span_attenuation_orl({1: a}, {1: b})[1]
    assert abs(s['att_a'] - 12.361 / 64.0424) < 1e-9
    assert abs(s['att_b'] - 12.352 / 64.0398) < 1e-9
    assert abs(s['att_avg'] - (s['att_a'] + s['att_b']) / 2) < 1e-12
    assert s['len_a_km'] == 64.0424 and abs(s['len_b_km'] - 64.0398) < 1e-9
    assert s['orl_a'] == 32.585 and s['orl_b'] == 32.954


def test_missing_figures_are_none_not_zero():
    a = _rec(loss=12.0, length_m=60000.0)          # no ORL stored
    b = _rec()                                       # nothing stored (JSON, old file)
    s = E.fiber_span_attenuation_orl({1: a}, {1: b})[1]
    assert s['att_a'] == 0.2 and s['att_b'] is None
    assert s['att_avg'] == 0.2                        # the one available direction
    assert s['orl_a'] is None and s['orl_b'] is None
    # A fiber present in only one direction still gets a row.
    s2 = E.fiber_span_attenuation_orl({2: a}, {})[2]
    assert s2['att_avg'] == 0.2 and s2['loss_b'] is None


def test_verdicts():
    old = (E.FIBER_ATTEN_DB_KM, E.SPAN_ORL_MIN_DB)
    try:
        E.FIBER_ATTEN_DB_KM, E.SPAN_ORL_MIN_DB = 0.250, 30.0
        assert E.atten_verdict(0.2500) == 'PASS'      # '<= 0.250' passes
        assert E.atten_verdict(0.2504) == 'PASS'      # rounds to 0.250
        assert E.atten_verdict(0.2505) == 'FAIL'
        assert E.atten_verdict(None) is None
        assert E.orl_verdict(32.5, 33.1) == 'PASS'
        assert E.orl_verdict(30.00, 31.0) == 'PASS'   # '>= 30' passes at 30
        assert E.orl_verdict(29.99, 33.0) == 'FAIL'   # either direction fails it
        assert E.orl_verdict(33.0, 29.5) == 'FAIL'
        assert E.orl_verdict(None, 31.0) == 'PASS'    # grade what exists
        assert E.orl_verdict(None, None) is None
        E.FIBER_ATTEN_DB_KM, E.SPAN_ORL_MIN_DB = 0.0, 0.0
        assert E.atten_verdict(0.9) is None and E.orl_verdict(1.0, 1.0) is None
    finally:
        E.FIBER_ATTEN_DB_KM, E.SPAN_ORL_MIN_DB = old


# ── 2. OFF is really off ─────────────────────────────────────────────────
def _run(tmp_path, name, **kw):
    out = tmp_path / name
    rc, m, err = run_splicereport(FIXTURE_SPLICE_A_DIR, FIXTURE_SPLICE_B_DIR,
                                  out, **kw)
    assert rc == 0 and m and m.get("ok"), f"run failed: {err[-1200:]}"
    _run.manifest = m
    return openpyxl.load_workbook(out)


def _legend_has(wb, label):
    for ws in wb.worksheets:
        if ws.title == SHEET:
            continue
        for row in ws.iter_rows(values_only=True):
            if row and row[0] == label:
                return True
    return False


def test_default_run_has_no_span_sheet_or_legend_rows(tmp_path):
    wb = _run(tmp_path, "off.xlsx")
    assert SHEET not in wb.sheetnames
    assert not _legend_has(wb, "Fiber attenuation")
    assert not _legend_has(wb, "ORL floor")
    assert not any(k in _run.manifest for k in
                   ("atten_gate_db_km", "orl_gate_db", "n_span_fibers",
                    "n_atten_fail", "n_orl_fail"))


def test_both_gates_on_add_the_sheet_legend_rows_and_manifest(tmp_path):
    wb = _run(tmp_path, "on.xlsx",
              overrides={"FIBER_ATTEN_DB_KM": 0.25, "SPAN_ORL_MIN_DB": 30.0})
    assert SHEET in wb.sheetnames
    ws = wb[SHEET]
    hdr = [ws.cell(row=4, column=c).value for c in range(1, 12)]
    assert hdr == ["Fiber", "Span loss A->B (dB)", "Span loss B->A (dB)",
                   "Length (km)", "Atten. A->B (dB/km)", "Atten. B->A (dB/km)",
                   "Atten. avg (dB/km)", "Atten. verdict", "ORL A->B (dB)",
                   "ORL B->A (dB)", "ORL verdict"]
    rows = [r[:11] for r in ws.iter_rows(min_row=5, values_only=True)
            if r[0] is not None and isinstance(r[0], int)]
    assert rows, "the fixture has fibers"
    n_att = n_orl = 0
    for r in rows:
        att_avg, att_v, orl_a, orl_b, orl_v = r[6], r[7], r[8], r[9], r[10]
        if att_avg is None:
            assert att_v == "not graded"
        else:
            assert att_v == ('FAIL' if round(att_avg, 3) > 0.25 else 'PASS')
            # The average really is the mean of the two directions shown.
            dirs = [v for v in (r[4], r[5]) if v is not None]
            assert abs(att_avg - sum(dirs) / len(dirs)) < 0.0015
        n_att += (att_v == 'FAIL')
        if orl_a is None and orl_b is None:
            assert orl_v == "not graded"
        else:
            worst = min(v for v in (orl_a, orl_b) if v is not None)
            assert orl_v == ('FAIL' if worst < 30.0 else 'PASS')
        n_orl += (orl_v == 'FAIL')
    assert _legend_has(wb, "Fiber attenuation") and _legend_has(wb, "ORL floor")
    m = _run.manifest
    assert m["atten_gate_db_km"] == 0.25 and m["orl_gate_db"] == 30.0
    assert m["n_span_fibers"] == len(rows)
    assert m["n_atten_fail"] == n_att and m["n_orl_fail"] == n_orl


def test_one_gate_on_leaves_the_other_not_graded(tmp_path):
    wb = _run(tmp_path, "att.xlsx", overrides={"FIBER_ATTEN_DB_KM": 0.25})
    ws = wb[SHEET]
    rows = [r for r in ws.iter_rows(min_row=5, values_only=True)
            if r[0] is not None and isinstance(r[0], int)]
    assert rows and all(r[10] == "not graded" for r in rows)
    assert _legend_has(wb, "Fiber attenuation") and not _legend_has(wb, "ORL floor")


def test_unticked_zero_survives_the_override_guard(tmp_path):
    wb = _run(tmp_path, "zero.xlsx",
              overrides={"FIBER_ATTEN_DB_KM": 0.0, "SPAN_ORL_MIN_DB": 0.0})
    assert SHEET not in wb.sheetnames


# ── 3. Profiles ──────────────────────────────────────────────────────────
def test_iig_turns_both_span_gates_on():
    hub = _hub()
    ov = hub._overrides_from_settings(hub._otdr_settings_from_profile(IIG))
    assert ov["FIBER_ATTEN_DB_KM"] == 0.250
    assert ov["SPAN_ORL_MIN_DB"] == 30.0


def test_other_profiles_leave_both_off():
    hub = _hub()
    for prof in hub.CUSTOMER_PROFILES:
        if prof == IIG:
            continue
        ov = hub._overrides_from_settings(hub._otdr_settings_from_profile(prof))
        assert ov.get("FIBER_ATTEN_DB_KM") == 0.0, prof
        assert ov.get("SPAN_ORL_MIN_DB") == 0.0, prof


def test_rows_are_wired_and_supported():
    hub = _hub()
    for key, glob in (("fiber_section_atten", "FIBER_ATTEN_DB_KM"),
                      ("span_orl", "SPAN_ORL_MIN_DB")):
        row = next(r for r in hub.OTDR_ROWS if r[0] == key)
        assert row[4] is True
        assert hub._OTDR_KEY_TO_ENGINE_GLOBAL[key] == glob
        assert hub._OTDR_KEY_DISABLE_VALUE[key] == 0.0
        assert key not in hub.OTDR_DEFAULT_APPLY
