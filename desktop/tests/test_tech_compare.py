"""Splice Report: compare our workbook against the tech's own and highlight
every difference in a third workbook, saved where the report went.

The boss's ask (2026-09-16): an upload box under the A/B folders for the
tech's Excel, then a cell-by-cell comparison that lands in the same folder
as the splice report.  Techs number splices from either end and add their
own bends/damage columns, so columns line up by DISTANCE, in whichever of
our two frames (A->B / B->A) fits better.
"""
from __future__ import annotations

import os

import sys
import types

import openpyxl
import pytest

from conftest import REPO_ROOT

SRC = open(os.path.join(REPO_ROOT, 'app.py'), encoding='utf-8').read()


def _load_block():
    """Exec the tech_compare block out of app.py in a bare module — it is
    engine-free and Streamlit-free by design, and lives inside app.py so no
    new file joins ENGINE_FILES (a new file freezes fleet hot-updates)."""
    start = SRC.index('# ─── tech_compare: begin ───')
    end = SRC.index('# ─── tech_compare: end ───')
    mod = types.ModuleType('tech_compare_block')
    sys.modules['tech_compare_block'] = mod       # dataclass() looks the module up
    exec(compile('from __future__ import annotations\nimport os, re\n' + SRC[start:end],
                 'app.py[tech_compare]', 'exec'), mod.__dict__)
    return mod


tc = _load_block()


# ── fixtures: a mini copy of our layout and a tech's hand-built sheet ──────
def _ours(path, site_a='LAN', site_b='KAN', span=50.0, cells=None):
    """Mimic write_xlsx: rows 1-2 = B->A / A->B distances (km and feet in one
    merged cell), row 3 merged headers, one row per ribbon with merged data
    cells."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Splice Report'
    splices = [('Splice 1', 7.83), ('Splice 2', 12.41), ('Bends @ 20.10km', 20.10),
               ('Splice 3', 30.25)]
    ws.cell(1, 2, 'B→A:'); ws.cell(2, 2, 'A→B:')
    for si, (lab, km) in enumerate(splices):
        kc, fc = 2 * si + 3, 2 * si + 4
        ws.cell(1, kc, f"{span - km:.2f}km, {(span - km) * 3280.84:,.0f}'")
        ws.cell(2, kc, f"{km:.2f}km, {km * 3280.84:,.0f}'")
        for r in (1, 2):
            ws.merge_cells(start_row=r, start_column=kc, end_row=r, end_column=fc)
        ws.cell(3, kc, lab); ws.merge_cells(start_row=3, start_column=kc, end_row=3, end_column=fc)
    end = 2 * len(splices) + 3
    ws.cell(1, end, "0.00km, 0'"); ws.cell(2, end, f"{span:.2f}km, {span * 3280.84:,.0f}'")
    ws.cell(3, 1, 'Ribbon'); ws.cell(3, 2, f'A-end ILA: {site_a}'); ws.cell(3, end, f'B-end ILA: {site_b}')
    for ri in range(3):
        r = ri + 4
        ws.cell(r, 1, f'Fiber {ri * 12 + 1}-{ri * 12 + 12} ({ri + 1}) (A{ri + 1})')
        for si in range(len(splices)):
            kc, fc = 2 * si + 3, 2 * si + 4
            ws.merge_cells(start_row=r, start_column=kc, end_row=r, end_column=fc)
    for (ri, col, text) in (cells or []):
        ws.cell(ri + 4, col, text)
    wb.save(path)
    return path


def _tech(path, reverse=False, cells=None, site_l='ILA:LAN', site_r='ILA: KAN'):
    """KANLAN-style tech sheet: Distance row, Ribbon/ILA/Splice header, ribbons."""
    wb = openpyxl.Workbook()
    ws = wb.active
    cols = [('Splice 1', '7.73km'), ('Splice 1A', '12.42KM'), ('bends', '20.05km'),
            ('Splice 2', '30.30km'), ('HH', '44.10km')]
    if reverse:   # tech counted from the far end
        cols = [(l, f'{50.0 - float(k[:-2]):.2f}km') for l, k in cols][::-1]
        site_l, site_r = site_r, site_l
    ws.cell(1, 2, 'Distance:')
    ws.cell(2, 1, 'Ribbon'); ws.cell(2, 2, site_l)
    for i, (lab, km) in enumerate(cols):
        ws.cell(1, 3 + i, km); ws.cell(2, 3 + i, lab)
    ws.cell(2, 3 + len(cols), site_r)
    for ri in range(3):
        ws.cell(ri + 3, 1, f'{ri * 12 + 1}-{ri * 12 + 12} ({ri + 1})')
    for (ri, col, text) in (cells or []):
        ws.cell(ri + 3, col, text)
    wb.save(path)
    return path


# ── cell parser: both vocabularies ────────────────────────────────────────
@pytest.mark.parametrize('text, want', [
    ('49,50,60 .369', {49: (0.369, ''), 50: (0.369, ''), 60: (0.369, '')}),
    ('196-198,200 .668', {196: (0.668, ''), 197: (0.668, ''), 198: (0.668, ''), 200: (0.668, '')}),
    ('1-8 brok', {f: (None, 'broke') for f in range(1, 9)}),
    ('179- 180 brok', {179: (None, 'broke'), 180: (None, 'broke')}),
    ('670 .706 near , 671 .723 far', {670: (0.706, ''), 671: (0.723, '')}),
    ('1 .207 12 .197', {1: (0.207, ''), 12: (0.197, '')}),
    ('169 .167 176 -0.162', {169: (0.167, ''), 176: (-0.162, '')}),
    ('10 BEND .883 bidi 11 bend .127 bidi', {10: (0.883, 'bend'), 11: (0.127, 'bend')}),
    ('1 broke@46.4k (B-fill OK) 2 broke@46.4k (B-fill OK)', {1: (None, 'broke'), 2: (None, 'broke')}),
    ('122 .171(B-fill) ~.086bd', {122: (0.171, '')}),
    ('1004 BAD_TAILBOX_REFL-27.5dB', {1004: (None, 'launch')}),
    ('112 ref 2.180 (refl -67dB)', {112: (2.18, 'ref')}),
    ('1 DZ 2 DZ 4 .183 (B-fill)', {1: (None, 'dz'), 2: (None, 'dz'), 4: (0.183, '')}),
])
def test_parse_cell_handles_tech_and_our_vocabularies(text, want):
    got = {f: (e.loss, e.tag) for f, e in tc.tc_parse_cell(text, 1, 12).items()}
    assert got == want


def test_all_expands_to_the_whole_ribbon():
    got = tc.tc_parse_cell('all 145 .35', 145, 156)
    assert set(got) == set(range(145, 157))
    assert got[145].loss == 0.35 and got[146].tag == 'flag'


# ── column line-up by distance, in whichever frame fits ───────────────────
def test_columns_line_up_by_distance_not_by_name(tmp_path):
    ours = tc.tc_read_grid(_ours(tmp_path / 'ours.xlsx'), 'Splice Report')
    tech = tc.tc_read_grid(_tech(tmp_path / 'tech.xlsx'))
    colmap, frame = tc.tc_line_up_columns(ours, tech)
    assert frame == 'A→B'
    pairs = {ours.columns[o].label: tech.columns[t].label for o, t in colmap.items()}
    assert pairs == {'A-end ILA: LAN': 'ILA:LAN', 'Splice 1': 'Splice 1',
                     'Splice 2': 'Splice 1A', 'Bends @ 20.10km': 'bends',
                     'Splice 3': 'Splice 2', 'B-end ILA: KAN': 'ILA: KAN'}
    # the tech's HH at 44.10 km has no partner within 250 m
    assert all(tech.columns[t].label != 'HH' for t in colmap.values())


def test_reverse_numbered_tech_sheet_uses_the_b_to_a_frame(tmp_path):
    ours = tc.tc_read_grid(_ours(tmp_path / 'ours.xlsx'), 'Splice Report')
    tech = tc.tc_read_grid(_tech(tmp_path / 'tech.xlsx', reverse=True))
    colmap, frame = tc.tc_line_up_columns(ours, tech)
    assert frame == 'B→A'
    pairs = {ours.columns[o].label: tech.columns[t].label for o, t in colmap.items()}
    assert pairs['Splice 1'] == 'Splice 1' and pairs['Splice 3'] == 'Splice 2'
    # ILA ends swap with the frame: our A end is the tech's right-hand column
    assert pairs['A-end ILA: LAN'] == 'ILA:LAN'
    assert pairs['B-end ILA: KAN'] == 'ILA: KAN'


# ── the comparison itself ─────────────────────────────────────────────────
def test_every_kind_of_difference_and_nothing_else(tmp_path):
    ours_p = _ours(tmp_path / 'ours.xlsx', cells=[
        (0, 3, '1 .207 12 .197'),          # Splice 1: F1 matches, F12 ours-only
        (1, 5, '13 .300 15 broke@20k'),    # Splice 2: F13 value differs, F15 type differs
        (2, 9, '25 .500'),                 # Splice 3: exact match
        (0, 2, '2 BAD_LAUNCH_REFL-40dB'),  # ILA A: ours-only launch issue
    ])
    tech_p = _tech(tmp_path / 'tech.xlsx', cells=[
        (0, 3, '1 .21'),                   # within the 0.010 tolerance
        (1, 4, '13 .350 14 .180 15 .400'), # F14 tech-only
        (2, 6, '25 .5'),
        (0, 7, '3 .900'),                  # HH column: no partner → tech-only
    ])
    out = tmp_path / 'diff.xlsx'
    r = tc.tc_compare_reports(str(ours_p), str(tech_p), str(out))
    kinds = {(d['Fiber'], d['Difference']) for d in r['diffs']}
    assert kinds == {(12, tc.TC_KIND_OURS_ONLY), (13, tc.TC_KIND_VALUE), (15, tc.TC_KIND_TYPE),
                     (14, tc.TC_KIND_TECH_ONLY), (2, tc.TC_KIND_OURS_ONLY), (3, tc.TC_KIND_TECH_ONLY)}
    assert r['n_diffs'] == 6 and r['xlsx'] == str(out) and out.exists()

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ['Summary', 'Differences', 'Difference list']
    ws = wb['Differences']
    # header pair names both sides; the unmatched HH column gets a grey header
    heads = [(ws.cell(1, c).value, ws.cell(2, c).value) for c in range(2, ws.max_column + 1)]
    assert ('Ours: Splice 2 @ 12.41 km', 'Tech: Splice 1A @ 12.42 km') in heads
    hh = next(c for c in range(2, ws.max_column + 1) if ws.cell(2, c).value == 'Tech: HH @ 44.10 km')
    assert ws.cell(1, hh).value == 'Ours: —'
    assert ws.cell(1, hh).fill.start_color.rgb.endswith('BFBFBF')
    # a matching cell stays blank; a differing one shows both texts and is coloured
    s3 = next(c for c in range(2, ws.max_column + 1) if ws.cell(1, c).value == 'Ours: Splice 3 @ 30.25 km')
    assert ws.cell(5, s3).value is None
    s2 = next(c for c in range(2, ws.max_column + 1) if ws.cell(1, c).value == 'Ours: Splice 2 @ 12.41 km')
    assert 'Ours: 13 .300 15 broke@20k' in ws.cell(4, s2).value
    assert 'Tech: 13 .350 14 .180 15 .400' in ws.cell(4, s2).value
    assert ws.cell(4, s2).fill.start_color.rgb.endswith('FFC7CE')   # worst kind = tech only
    wl = wb['Difference list']
    assert wl.max_row == 7 and [wl.cell(1, c).value for c in range(1, 8)][-1] == 'Difference'


def test_identical_reports_produce_no_differences(tmp_path):
    cells = [(0, 3, '1 .207 12 .197'), (2, 9, '25 .500')]
    ours_p = _ours(tmp_path / 'ours.xlsx', cells=cells)
    tech_p = _tech(tmp_path / 'tech.xlsx', cells=[(0, 3, '1 .207 12 .197'), (2, 6, '25 .500')])
    r = tc.tc_compare_reports(str(ours_p), str(tech_p), str(tmp_path / 'd.xlsx'))
    assert r['n_diffs'] == 0 and r['columns_matched'] == 4


def test_tech_sheet_without_a_ribbon_header_is_rejected_cleanly(tmp_path):
    wb = openpyxl.Workbook(); wb.active.cell(1, 1, 'nothing here'); wb.save(tmp_path / 't.xlsx')
    with pytest.raises(ValueError):
        tc.tc_read_grid(str(tmp_path / 't.xlsx'))


# ── hub wiring ────────────────────────────────────────────────────────────
def test_upload_box_sits_under_the_a_b_inputs_on_both_modes():
    # The A/B boxes + the upload live in _sr_span_inputs (one call per span
    # since the second-span option); the site names follow in _sr_site_inputs.
    body = SRC.split('def _sr_span_inputs(span):', 1)[1].split('\ndef ', 1)[0]
    up = body.index("key=k_tech")
    assert body.index("'sr_tech_xlsx'") < up                  # span 1's key
    assert body.index("st.text_input('B folder', key=k_b") < up   # two-folder mode
    assert body.index("key=k_zip") < up                       # one-folder/zip mode
    page = SRC.split('def page_splice_report(fr=False):', 1)[1]
    assert page.index('_sr_span_inputs(1)') < page.index('_sr_site_inputs(1, dir_a, dir_b)')


def test_comparison_is_written_to_the_report_folder_and_offered():
    assert "_render_tech_comparison(f'{_p}{sfx}', xp, tech_xlsx, dest," in SRC
    body = SRC.split('def _render_tech_comparison(', 1)[1].split('\ndef ', 1)[0]
    assert "_SpliceReport_vs_Tech.xlsx" in body
    assert 'out_path = os.path.join(dest_dir,' in body
    assert "st.download_button('⬇ Differences vs tech (Excel)'" in body
    assert "report_error('splice report — tech comparison'" in body


def test_no_new_engine_file_so_the_fleet_hot_updates():
    """The compare code lives inside app.py: a new shipped file changes the
    ENGINE_FILES set, which the installed launcher rejects until every tech
    runs a fresh installer."""
    assert not os.path.exists(os.path.join(REPO_ROOT, 'tech_compare.py'))
    launcher = open(os.path.join(REPO_ROOT, 'desktop/launcher.py'), encoding='utf-8').read()
    assert 'tech_compare' not in launcher
    assert 'def tc_compare_reports(' in SRC and 'def _render_tech_comparison(' in SRC


# ── the page renders with the box on both input modes ─────────────────────
def _splice_page(**state):
    from conftest import FIXTURE_SPLICE_A_DIR, FIXTURE_SPLICE_B_DIR, run_streamlit
    at = run_streamlit().run()
    at.session_state['view_dir_a_input'] = str(FIXTURE_SPLICE_A_DIR)
    at.session_state['view_dir_b_input'] = str(FIXTURE_SPLICE_B_DIR)
    for k, v in state.items():
        at.session_state[k] = v
    at.sidebar.radio[0].set_value('Splice Report').run()
    return at


def _uploaders(at):
    try:
        return [e for e in at.get('file_uploader')]
    except Exception:
        return None


@pytest.mark.parametrize('mode', ['Two folders (A + B)', 'One folder / zip (both directions)'])
def test_splice_report_page_renders_the_tech_upload_box(monkeypatch, tmp_path, mode):
    monkeypatch.setenv('HOME', str(tmp_path))
    monkeypatch.delenv('SS_ERROR_WEBHOOK', raising=False)
    at = _splice_page(sr_input_mode=mode)
    assert not at.exception, f'page raised: {list(at.exception)}'
    ups = _uploaders(at)
    if ups is not None:   # AppTest exposes uploaders on this Streamlit
        labels = [getattr(u, 'label', '') for u in ups]
        assert any('Tech' in (l or '') and 'compare' in (l or '') for l in labels), labels
